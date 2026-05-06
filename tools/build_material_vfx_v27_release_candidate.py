from pathlib import Path
import importlib.util
import math
import re
import subprocess

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image, ImageDraw, ImageFilter, ImageFont


ROOT = Path(__file__).resolve().parents[1]
V26_SCRIPT = ROOT / "tools" / "build_material_vfx_v26_polished_book.py"

spec = importlib.util.spec_from_file_location("material_vfx_v26", V26_SCRIPT)
v26 = importlib.util.module_from_spec(spec)
spec.loader.exec_module(v26)

base = v26.base
v24 = v26.v25.v24

OUT_DOCX = Path.home() / "Downloads" / "Realtime_Material_VFX_Field_Guide.docx"
OUT_PDF = Path.home() / "Downloads" / "Realtime_Material_VFX_Field_Guide.pdf"
base.OUT = OUT_DOCX

IMG_DIR = ROOT / "Saved" / "material_vfx_v5_images"
TITLE = "Realtime Material VFX Field Guide"
SUBTITLE = "Shader Thinking, Visual Math, and Production Diagnostics"
VERSION_LABEL = "First Edition"
WRITTEN_LABEL = "Written May 2026"


def _font(size=30, bold=False):
    candidates = [
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
        "C:/Windows/Fonts/segoeuib.ttf" if bold else "C:/Windows/Fonts/segoeui.ttf",
        "C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf",
    ]
    for path in candidates:
        if Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def _save(img, name):
    IMG_DIR.mkdir(parents=True, exist_ok=True)
    path = IMG_DIR / name
    img.save(path, quality=94)
    return path


def _add_field(paragraph, instr, placeholder="1"):
    run = paragraph.add_run()
    base.set_font(run, size=8.0, name=base.SANS, color=base.MUTED)

    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    run._r.append(fld_begin)

    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instr
    run._r.append(instr_text)

    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    run._r.append(fld_sep)

    text = OxmlElement("w:t")
    text.text = placeholder
    run._r.append(text)

    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_end)


def _add_toc_entry_field(paragraph, text, level=1, identifier="C"):
    safe_text = text.replace('"', "'")
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")

    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = f'TC "{safe_text}" \\f {identifier} \\l {level}'

    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr_text, fld_end])


def _set_footer_with_page_numbers(doc):
    for index, section in enumerate(doc.sections):
        section.different_first_page_header_footer = False
        section.footer.is_linked_to_previous = False
        section.first_page_footer.is_linked_to_previous = False

        footer = section.footer.paragraphs[0]
        footer.clear()
        first_footer = section.first_page_footer.paragraphs[0]
        first_footer.clear()

        if index == 0:
            continue

        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run(f"{TITLE} | Page ")
        base.set_font(run, size=8.0, name=base.SANS, color=base.MUTED)
        _add_field(footer, "PAGE")


def _apply_book_grid(doc):
    for section in doc.sections:
        section.left_margin = Cm(2.85)
        section.right_margin = Cm(2.85)


def code_box_v27(doc, lines, title=None):
    if title:
        base.p(doc, title, size=8.8, font=base.SANS, italic=True, color=base.MUTED, before=5, after=1)
    para = base.p(doc, "", size=8.05, font=base.CODE, before=0, after=7, color=base.INK)
    para.paragraph_format.left_indent = Cm(0.35)
    para.paragraph_format.right_indent = Cm(0.35)
    para.paragraph_format.line_spacing = 1.05
    base.shade(para, "E9EDF2")
    for i, line in enumerate(lines):
        if i:
            para.add_run().add_break()
        run = para.add_run(line)
        base.set_font(run, size=8.05, name=base.CODE, color=base.INK)


def _shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell(cell, text, bold=False, fill=None, color=None, size=8.7):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.line_spacing = 1.08
    run = p.add_run(text)
    base.set_font(run, size=size, name=base.SANS if bold else base.BODY, bold=bold, color=color or base.INK)
    if fill:
        _shade_cell(cell, fill)


def _set_cell_width(cell, width_cm):
    cell.width = Cm(width_cm)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def _comparison_table(doc, headers, rows):
    for row in rows:
        title = str(row[0])
        details = []
        for i, value in enumerate(row[1:], 1):
            label = headers[i] if i < len(headers) else f"Field {i + 1}"
            details.append(f"{label}: {value}")
        base.note(doc, title, " | ".join(details), False)
    return
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    if len(headers) == 2:
        widths = [4.4, 10.4]
    elif len(headers) == 3:
        widths = [4.0, 6.0, 4.8]
    else:
        widths = [14.8 / len(headers)] * len(headers)
    for i, header in enumerate(headers):
        _set_cell_width(table.rows[0].cells[i], widths[i])
        _set_cell(table.rows[0].cells[i], header, bold=True, fill="DCEBFA", color=base.BLUE, size=8.4)
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            _set_cell_width(cells[i], widths[i])
            _set_cell(cells[i], text, fill="F7F9FB", size=8.0)
    doc.add_paragraph().paragraph_format.space_after = Pt(5)


def _link_paragraph(doc, label, url, suffix):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.line_spacing = 1.12
    p.paragraph_format.left_indent = Cm(0.35)
    p.add_run("- ")
    _add_hyperlink(p, label, url)
    r = p.add_run(" - " + suffix)
    base.set_font(r, size=9.2, name=base.BODY, color=base.INK)


def cover_v27(doc):
    eyebrow = base.p(
        doc,
        VERSION_LABEL.upper(),
        size=10.5,
        font=base.SANS,
        bold=True,
        color=base.ORANGE,
        before=32,
        after=8,
    )
    title = base.p(
        doc,
        TITLE,
        size=29,
        font=base.HEAD,
        bold=True,
        color=base.BLUE,
        after=6,
    )
    title.paragraph_format.line_spacing = 1.0
    base.p(
        doc,
        SUBTITLE,
        size=12.2,
        font=base.BODY,
        italic=True,
        color=base.MUTED,
        after=14,
    )

    cover_art = IMG_DIR / "v27_cover_visual_system.png"
    if cover_art.exists():
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_before = Pt(6)
        para.paragraph_format.space_after = Pt(16)
        para.add_run().add_picture(str(cover_art), width=Cm(15.0))

    summary = base.p(
        doc,
        "A production-oriented guide for reading, building, debugging, and shipping realtime material VFX. Built around visual math, shader reasoning, defensive graph habits, profiling, and practical Unreal/Niagara workflows.",
        size=10.7,
        font=base.BODY,
        color=base.INK,
        before=2,
        after=14,
    )
    base.p(
        doc,
        "Space | Vector | Wave | Field | Sampling | Light | Energy | Debug | Cost",
        size=10.4,
        font=base.SANS,
        bold=True,
        color=base.BLUE,
        before=6,
        after=16,
    )
    meta = base.p(
        doc,
        f"{VERSION_LABEL} | {WRITTEN_LABEL} | Realtime VFX Brain",
        size=9.0,
        font=base.SANS,
        color=base.MUTED,
        before=2,
        after=0,
    )
    doc.add_section(WD_SECTION.NEW_PAGE)


def _canvas(title, subtitle, size=(1800, 980)):
    img = Image.new("RGB", size, "#F8F7F3")
    d = ImageDraw.Draw(img)
    d.rectangle([0, 0, size[0], 88], fill="#173B59")
    d.text((58, 24), title, fill="white", font=_font(30, True))
    d.text((58, 112), subtitle, fill="#42505C", font=_font(22))
    return img, d


def _arrow(d, p0, p1, fill="#2563EB", width=5):
    d.line([p0, p1], fill=fill, width=width)
    x0, y0 = p0
    x1, y1 = p1
    a = math.atan2(y1 - y0, x1 - x0)
    pts = []
    for da in (0, 2.45, -2.45):
        pts.append((x1 - 22 * math.cos(a + da), y1 - 22 * math.sin(a + da)))
    d.polygon(pts, fill=fill)


def _cover_visual():
    img = Image.new("RGB", (1800, 760), "#F6F4EF")
    d = ImageDraw.Draw(img)
    d.rounded_rectangle([30, 30, 1770, 730], radius=24, fill="#FFFFFF", outline="#D6DEE8", width=3)

    # Soft grid.
    for x in range(90, 1740, 90):
        d.line([x, 70, x, 690], fill="#EEF2F6", width=1)
    for y in range(90, 700, 70):
        d.line([70, y, 1730, y], fill="#EEF2F6", width=1)

    d.rectangle([30, 30, 1770, 106], fill="#173B59")
    d.text((72, 55), "Realtime material VFX as data flow", fill="#FFFFFF", font=_font(34, True))
    d.text((1180, 58), "Visual math / Debug / Cost", fill="#CDE5FF", font=_font(24, True))

    # Left: node graph lane.
    lane_y = [190, 285, 380, 475]
    labels = [
        ("Texture Alpha", "#EAF3FF"),
        ("Remap 0..1", "#ECFDF3"),
        ("Shape Curve", "#FFF4E8"),
        ("Output Mask", "#F3E8FF"),
    ]
    x0 = 110
    for i, (label, fill) in enumerate(labels):
        y = lane_y[i]
        d.rounded_rectangle([x0, y, x0 + 280, y + 56], radius=10, fill=fill, outline="#B8C7D9", width=2)
        d.text((x0 + 28, y + 16), label, fill="#173B59", font=_font(22, True))
        if i < len(labels) - 1:
            _arrow(d, (x0 + 280, y + 28), (x0 + 390, lane_y[i + 1] + 28), fill="#2563EB", width=4)

    # Center: code block.
    d.rounded_rectangle([650, 180, 1180, 540], radius=18, fill="#101820", outline="#334155", width=3)
    code_lines = [
        "float mask = saturate(alpha);",
        "mask = pow(mask, ShapePower);",
        "float ring = smoothstep(a,b,mask);",
        "float3 color = mask * Tint;",
        "return DebugView ? mask : color;",
    ]
    d.text((690, 220), "HLSL sketch", fill="#8BD3FF", font=_font(25, True))
    for i, line in enumerate(code_lines):
        d.text((690, 275 + i * 42), line, fill="#E8F1F8", font=_font(23))

    # Right: visual output and diagnostics.
    cx, cy = 1480, 330
    for r, color in [(160, "#FFE6A6"), (100, "#FFC64D"), (44, "#FFF2C7")]:
        d.ellipse([cx - r, cy - r, cx + r, cy + r], fill=color)
    d.ellipse([cx - 165, cy - 165, cx + 165, cy + 165], outline="#173B59", width=3)
    d.text((1360, 525), "Output preview", fill="#173B59", font=_font(24, True))

    # Bottom legend.
    chips = [("Node Graph", "#EAF3FF"), ("HLSL", "#E8F1F8"), ("Visual", "#FFF4E8"), ("Failure Atlas", "#FCE7E7")]
    x = 110
    for label, fill in chips:
        d.rounded_rectangle([x, 620, x + 260, 674], radius=10, fill=fill, outline="#CBD5E1", width=2)
        d.text((x + 28, 636), label, fill="#173B59", font=_font(22, True))
        x += 300

    _save(img, "v27_cover_visual_system.png")


def generate_v27_diagrams():
    _cover_visual()

    # Curve 0..1 / remap / shaping.
    img, d = _canvas("Curve 0..1 / Remap / Shaping", "Một range sạch giúp Power, SmoothStep và Lerp phản ứng có kiểm soát.")
    gx0, gy0, gx1, gy1 = 170, 760, 980, 240
    d.line([gx0, gy1, gx0, gy0, gx1, gy0], fill="#1F2937", width=4)
    for i in range(6):
        x = gx0 + i * (gx1 - gx0) / 5
        y = gy0 - i * (gy0 - gy1) / 5
        d.line([x, gy0, x, gy1], fill="#D1D5DB", width=1)
        d.line([gx0, y, gx1, y], fill="#D1D5DB", width=1)
    curves = [("linear", 1.0, "#111827"), ("soft", 0.5, "#2563EB"), ("core", 3.0, "#E66A2C")]
    for label, power, color in curves:
        pts = []
        for i in range(240):
            x = i / 239
            y = x ** power
            pts.append((gx0 + x * (gx1 - gx0), gy0 - y * (gy0 - gy1)))
        d.line(pts, fill=color, width=5)
    for idx, (label, _, color) in enumerate(curves):
        y = 300 + idx * 42
        d.line([1080, y, 1160, y], fill=color, width=7)
        d.text((1182, y - 15), label, fill=color, font=_font(23, True))
    d.rounded_rectangle([1060, 520, 1660, 660], radius=18, fill="white", outline="#CBD5E1", width=3)
    d.text((1095, 560), "t = saturate((x - inMin) / (inMax - inMin))", fill="#111827", font=_font(27, True))
    d.text((1095, 610), "value = pow(t, CorePower)", fill="#111827", font=_font(27, True))
    _save(img, "v27_curve_01.png")

    # SDF ring.
    img, d = _canvas("SDF Ring / Width / Softness", "Distance tạo shape sạch; SmoothStep tạo dải biên có kiểm soát.")
    cx, cy = 540, 540
    for r, color, w in [(260, "#CBD5E1", 5), (205, "#2563EB", 10), (150, "#CBD5E1", 5)]:
        d.ellipse([cx-r, cy-r, cx+r, cy+r], outline=color, width=w)
    d.ellipse([cx - 8, cy - 8, cx + 8, cy + 8], fill="#E66A2C")
    d.text((285, 825), "ring = 1 - smoothstep(width, width + softness, abs(r - radius))", fill="#111827", font=_font(24, True))
    d.line([990, 730, 1620, 730], fill="#1F2937", width=4)
    d.line([990, 260, 990, 730], fill="#1F2937", width=4)
    pts = []
    for i in range(560):
        x = i / 559
        v = max(0, 1 - abs(x - 0.5) / 0.14)
        v = min(1, v)
        pts.append((990 + i, 730 - v * 390))
    d.line(pts, fill="#2563EB", width=5)
    d.rectangle([1240, 300, 1370, 730], outline="#E66A2C", width=4)
    d.text((1230, 750), "edge band", fill="#E66A2C", font=_font(24, True))
    _save(img, "v27_sdf_ring.png")

    # Polar UV unwrap.
    img, d = _canvas("Polar UV Unwrap", "Cartesian UV đổi thành radius và angle; seam là giá phải quản lý.")
    d.rectangle([160, 220, 720, 780], outline="#1F2937", width=4, fill="#FFFFFF")
    for i in range(1, 8):
        x = 160 + i * 70
        y = 220 + i * 70
        d.line([x, 220, x, 780], fill="#E5E7EB", width=1)
        d.line([160, y, 720, y], fill="#E5E7EB", width=1)
    cx, cy = 440, 500
    for r in (80, 160, 240):
        d.ellipse([cx-r, cy-r, cx+r, cy+r], outline="#2563EB", width=3)
    d.line([cx, cy, cx + 250, cy], fill="#E66A2C", width=6)
    d.text((585, 470), "angle seam", fill="#E66A2C", font=_font(22, True))
    _arrow(d, (780, 500), (980, 500), "#173B59", 7)
    d.rectangle([1040, 220, 1640, 780], outline="#1F2937", width=4, fill="#FFFFFF")
    for i in range(8):
        y = 220 + i * 70
        d.line([1040, y, 1640, y], fill="#E5E7EB", width=1)
    for i in range(12):
        x = 1040 + i * 50
        d.line([x, 220, x, 780], fill="#EEF2F7", width=1)
    d.rectangle([1040, 220, 1085, 780], fill="#FFE4D6")
    d.text((1075, 805), "angle 0/1 seam", fill="#E66A2C", font=_font(24, True))
    d.text((1260, 805), "U = angle, V = radius", fill="#111827", font=_font(24, True))
    _save(img, "v27_polar_uv_unwrap.png")

    # Overdraw.
    img, d = _canvas("Overdraw / Alpha Card Area", "Một shader đơn giản vẫn đắt nếu nhiều translucent cards phủ cùng pixel.")
    for panel, x0, count, radius, label in [("A. Wide cards", 110, 34, 135, "area và layer count cao"), ("B. Cropped cards", 960, 13, 78, "crop alpha, ít overlap")]:
        d.rounded_rectangle([x0, 220, x0 + 720, 800], radius=20, fill="#111827", outline="#CBD5E1", width=3)
        d.text((x0 + 25, 235), panel, fill="white", font=_font(28, True))
        layer = Image.new("RGBA", (720, 580), (0, 0, 0, 0))
        ld = ImageDraw.Draw(layer, "RGBA")
        rng = __import__("random").Random(count)
        for _ in range(count):
            cx = rng.randint(140, 580)
            cy = rng.randint(170, 450)
            color = (240, 74, 50, 70) if count > 20 else (53, 190, 92, 95)
            ld.ellipse([cx-radius, cy-radius, cx+radius, cy+radius], fill=color)
        img.paste(Image.alpha_composite(Image.new("RGBA", (720, 580), (17, 24, 39, 255)), layer).convert("RGB"), (x0, 220))
        d.text((x0 + 35, 825), label, fill="#111827", font=_font(25, True))
    d.rectangle([1390, 145, 1470, 185], fill="#22C55E")
    d.rectangle([1470, 145, 1550, 185], fill="#EAB308")
    d.rectangle([1550, 145, 1630, 185], fill="#EF4444")
    d.rectangle([1630, 145, 1710, 185], fill="#F8FAFC", outline="#111827")
    d.text((1390, 195), "rẻ → đắt", fill="#42505C", font=_font(20, True))
    _save(img, "v27_overdraw_layers.png")

    # SubUV local UV.
    img, d = _canvas("SubUV Atlas UV vs Local Frame UV", "Renderer chọn frame; material cần local UV nếu tạo mask riêng trong từng frame.")
    x0, y0, cell = 130, 230, 125
    for r in range(4):
        for c in range(4):
            x, y = x0 + c * cell, y0 + r * cell
            fill = "#E0F2FE" if (r, c) == (1, 2) else "#FFFFFF"
            d.rectangle([x, y, x+cell, y+cell], fill=fill, outline="#1F2937", width=3)
            d.text((x + 48, y + 43), str(r * 4 + c), fill="#111827", font=_font(32, True))
    d.text((160, 760), "Atlas UV: 0..1 trên toàn sheet", fill="#111827", font=_font(24, True))
    _arrow(d, (720, 480), (930, 480), "#173B59", 7)
    d.rectangle([1010, 270, 1510, 770], fill="#E0F2FE", outline="#1F2937", width=4)
    for i in range(1, 5):
        d.line([1010 + i*100, 270, 1010 + i*100, 770], fill="#93C5FD", width=2)
        d.line([1010, 270 + i*100, 1510, 270 + i*100], fill="#93C5FD", width=2)
    d.text((1090, 495), "frame 6", fill="#111827", font=_font(52, True))
    d.text((1040, 810), "Local UV: 0..1 bên trong frame hiện tại", fill="#111827", font=_font(24, True))
    d.text((1020, 880), "Sai frame = renderer/SubUV. Sai gradient trong ô = local UV.", fill="#E66A2C", font=_font(24, True))
    _save(img, "v27_subuv_local_uv.png")

    # Blend/depth sorting.
    img, d = _canvas("Blend / Depth / Sorting", "Blend mode quyết định pixel trộn với scene và rủi ro sorting.")
    for i, (title, fill, text) in enumerate([
        ("Additive", "#111827", "shape dễ mất trên nền sáng"),
        ("Translucent", "#1F2937", "mềm nhưng sorting/overdraw"),
        ("DepthFade", "#0F172A", "làm mềm giao tuyến geometry"),
    ]):
        x = 120 + i * 555
        d.rounded_rectangle([x, 240, x+470, 760], radius=20, fill=fill, outline="#CBD5E1", width=3)
        d.text((x + 25, 265), title, fill="white", font=_font(30, True))
        d.polygon([(x+30, 650), (x+430, 565), (x+430, 720), (x+30, 720)], fill="#64748B")
        for r in range(150, 30, -35):
            alpha = int(210 * (r / 150))
            col = (250, 180, 70, alpha) if i == 0 else (160, 210, 250, alpha)
            layer = Image.new("RGBA", (470, 520), (0, 0, 0, 0))
            ld = ImageDraw.Draw(layer, "RGBA")
            if i == 2:
                ld.ellipse([235-r, 290-r, 235+r, 290+r], fill=(180, 210, 240, max(20, alpha - 70)))
            else:
                ld.ellipse([235-r, 250-r, 235+r, 250+r], fill=col)
            img.paste(Image.alpha_composite(img.crop((x, 240, x+470, 760)).convert("RGBA"), layer).convert("RGB"), (x, 240))
        d.text((x + 35, 790), text, fill="#111827", font=_font(22, True))
    _save(img, "v27_blend_depth_sorting.png")

    # Case study: M_Flare.
    img, d = _canvas("Case Study Diagram: M_Flare", "Diagram tái dựng từ workflow VFXSamplePack, không phải screenshot Unreal thật.", (2000, 1200))
    nodes = [
        ("Texture Alpha", 110, 260), ("Radial/Falloff", 420, 260), ("Particle Color", 110, 500),
        ("Intensity", 420, 500), ("Camera/Depth Fade", 730, 380), ("Emissive/Opacity", 1040, 380),
    ]
    for label, x, y in nodes:
        d.rounded_rectangle([x, y, x+230, y+82], radius=14, fill="white", outline="#CBD5E1", width=3)
        d.text((x+18, y+25), label, fill="#111827", font=_font(20, True))
    for p0, p1 in [((340, 301), (420, 301)), ((340, 541), (420, 541)), ((650, 301), (730, 410)), ((650, 541), (730, 430)), ((960, 421), (1040, 421))]:
        _arrow(d, p0, p1, "#2563EB", 4)
    preview = Image.new("RGB", (360, 360), "#101827")
    pp = preview.load()
    cx = cy = 180
    for y in range(360):
        for x in range(360):
            r = math.sqrt((x-cx)**2 + (y-cy)**2) / 175
            m = max(0, 1-r)
            core = m ** 5
            glow = m ** 0.55
            pp[x, y] = (min(255, int(core*255 + glow*90)), min(255, int(core*180 + glow*70)), min(255, int(core*45 + glow*18)))
    preview = preview.filter(ImageFilter.GaussianBlur(1.2))
    img.paste(preview, (1440, 245))
    d.text((1460, 625), "output preview", fill="#111827", font=_font(22, True))
    mask = Image.new("L", (220, 220))
    mp = mask.load()
    for y in range(220):
        for x in range(220):
            r = math.sqrt((x-110)**2 + (y-110)**2) / 105
            mp[x, y] = int(max(0, 1-r)**2.4 * 255)
    img.paste(mask.convert("RGB"), (1430, 720))
    img.paste(mask.point(lambda v: 255 if 80 < v < 170 else 0).convert("RGB"), (1685, 720))
    d.text((1435, 955), "mask debug", fill="#111827", font=_font(20, True))
    d.text((1690, 955), "edge debug", fill="#111827", font=_font(20, True))
    lanes = ["Shape", "UV", "Mask", "Color/Emissive", "Interaction", "Cost"]
    for i, lane in enumerate(lanes):
        x = 160 + (i % 3) * 390
        y = 760 + (i // 3) * 120
        d.rounded_rectangle([x, y, x+320, y+66], radius=12, fill="#EFF6FF", outline="#93C5FD", width=2)
        d.text((x+18, y+20), lane, fill="#173B59", font=_font(21, True))
    _save(img, "v27_case_m_flare.png")

    # Case study: SubUV.
    img, d = _canvas("Case Study Diagram: SubUV / NumberGrid", "Diagram tái dựng từ workflow VFXSamplePack, không phải screenshot Unreal thật.", (2000, 1150))
    for r in range(3):
        for c in range(3):
            x, y = 140 + c * 150, 245 + r * 150
            fill = "#F8FAFC" if (r, c) != (1, 1) else "#DCFCE7"
            d.rectangle([x, y, x+150, y+150], fill=fill, outline="#1F2937", width=3)
            d.text((x+58, y+48), str(r*3+c+1), fill="#111827", font=_font(42, True))
    d.text((160, 735), "NumberGrid atlas", fill="#111827", font=_font(24, True))
    d.rounded_rectangle([720, 260, 1120, 380], radius=16, fill="white", outline="#CBD5E1", width=3)
    d.text((755, 303), "Niagara frame index", fill="#111827", font=_font(24, True))
    d.rounded_rectangle([720, 500, 1120, 620], radius=16, fill="white", outline="#CBD5E1", width=3)
    d.text((765, 543), "Material local UV", fill="#111827", font=_font(24, True))
    _arrow(d, (600, 470), (720, 325), "#2563EB", 5)
    _arrow(d, (600, 470), (720, 560), "#E66A2C", 5)
    d.rectangle([1320, 260, 1760, 700], fill="#DCFCE7", outline="#1F2937", width=4)
    d.ellipse([1435, 375, 1645, 585], outline="#2563EB", width=12)
    d.text((1350, 735), "Đúng: frame đúng, local mask nằm trong ô", fill="#15803D", font=_font(24, True))
    d.text((160, 890), "Nếu số sai: kiểm renderer/SubUV grid hoặc frame curve.", fill="#DC2626", font=_font(25, True))
    d.text((160, 940), "Nếu số đúng nhưng mask trượt qua nhiều ô: kiểm local UV.", fill="#DC2626", font=_font(25, True))
    _save(img, "v27_case_subuv_numbergrid.png")

    # Niagara/material contract as a flat channel map.
    img, d = _canvas("Niagara -> Material Contract", "Data flow 2D: particle attributes đi vào material lanes, không phải ảnh render.")
    lanes = [
        ("ParticleColor.rgb", "Tint / hue", "#2563EB"),
        ("ParticleColor.a", "Opacity / fade", "#16A34A"),
        ("DynamicParam.r", "Dissolve", "#DC2626"),
        ("DynamicParam.g", "Distortion", "#7C3AED"),
        ("SubUV frame", "Atlas sample", "#EA580C"),
    ]
    for i, (src, dst, color) in enumerate(lanes):
        y = 230 + i * 120
        d.rounded_rectangle([120, y, 480, y + 72], radius=12, fill="white", outline=color, width=3)
        d.rounded_rectangle([1220, y, 1600, y + 72], radius=12, fill="white", outline=color, width=3)
        d.text((150, y + 22), src, fill="#111827", font=_font(22, True))
        d.text((1250, y + 22), dst, fill="#111827", font=_font(22, True))
        _arrow(d, (500, y + 36), (1200, y + 36), color, 5)
        d.rounded_rectangle([720, y - 8, 980, y + 80], radius=14, fill="#F8FAFC", outline="#CBD5E1", width=2)
        d.text((748, y + 10), "Material lane", fill="#173B59", font=_font(18, True))
        d.text((748, y + 42), ["Color", "Opacity", "Mask", "UV offset", "Texture"][i], fill="#111827", font=_font(19, True))
    d.text((130, 850), "Debug order: verify attribute -> isolate lane -> preview grayscale -> reconnect final output.", fill="#42505C", font=_font(25, True))
    _save(img, "v27_niagara_material_contract_2d.png")

    # Shader language as input/function/output blocks.
    img, d = _canvas("Shader Function Contract", "Một Custom Node tốt là function nhỏ với input/output và debug path rõ.")
    blocks = [
        ("Inputs", ["UV", "Mask", "Time", "Params"], 130),
        ("Function", ["Remap", "Shape", "Sample", "Clamp"], 620),
        ("Outputs", ["Emissive", "Opacity", "Debug"], 1120),
    ]
    for title, items, x in blocks:
        d.rounded_rectangle([x, 250, x + 360, 710], radius=18, fill="white", outline="#CBD5E1", width=3)
        d.text((x + 28, 275), title, fill="#173B59", font=_font(29, True))
        for idx, item in enumerate(items):
            y = 355 + idx * 78
            d.rounded_rectangle([x + 42, y, x + 318, y + 48], radius=10, fill="#EFF6FF", outline="#93C5FD", width=2)
            d.text((x + 65, y + 12), item, fill="#111827", font=_font(21, True))
    _arrow(d, (490, 480), (620, 480), "#2563EB", 6)
    _arrow(d, (980, 480), (1120, 480), "#2563EB", 6)
    d.rounded_rectangle([1540, 345, 1700, 610], radius=14, fill="#F1F5F9", outline="#CBD5E1", width=2)
    d.text((1562, 370), "Debug", fill="#173B59", font=_font(24, True))
    d.text((1562, 420), "return", fill="#111827", font=_font(20, True))
    d.text((1562, 455), "Mask.xxx", fill="#111827", font=_font(20, True))
    d.text((1562, 490), "UV.xy", fill="#111827", font=_font(20, True))
    d.text((1562, 525), "Edge.xxx", fill="#111827", font=_font(20, True))
    _save(img, "v27_shader_contract_blocks.png")

    # POM/height as side-view cross-section.
    img, d = _canvas("POM / Height: Side View", "POM raymarch giả lập chiều sâu trong UV; silhouette thật không đổi.")
    d.line([130, 705, 1650, 705], fill="#1F2937", width=4)
    pts = []
    for i in range(680):
        x = 220 + i * 1.8
        y = 520 - 80 * math.sin(i / 55) - 38 * math.sin(i / 19)
        pts.append((x, y))
    d.line(pts, fill="#2563EB", width=5)
    for x in range(280, 1480, 160):
        _arrow(d, (x, 260), (x - 70, 520), "#E66A2C", 4)
    d.text((185, 750), "mesh surface / silhouette", fill="#111827", font=_font(24, True))
    d.text((220, 470), "height field", fill="#2563EB", font=_font(24, True))
    d.text((780, 300), "view rays step through height", fill="#E66A2C", font=_font(24, True))
    d.rounded_rectangle([1190, 760, 1640, 875], radius=14, fill="#FFF7ED", outline="#FDBA74", width=3)
    d.text((1220, 792), "Risk: more samples, grazing angle cost,", fill="#111827", font=_font(20, True))
    d.text((1220, 825), "no true silhouette change.", fill="#111827", font=_font(20, True))
    _save(img, "v27_pom_height_cross_section.png")

    # Production story as decision matrix.
    img, d = _canvas("Production Strategy Matrix", "Style/platform quyết định layer nào giữ, bake, hoặc bỏ.")
    cols = ["Mobile", "Stylized", "iGaming", "AAA"]
    rows = ["Silhouette", "Runtime layers", "Texture budget", "Debug priority"]
    x0, y0 = 210, 235
    cw, rh = 330, 112
    for ci, col in enumerate([""] + cols):
        x = x0 + ci * cw
        d.rectangle([x, y0, x+cw, y0+rh], fill="#173B59" if ci else "#E5E7EB", outline="#CBD5E1", width=2)
        d.text((x+24, y0+38), col, fill="white" if ci else "#111827", font=_font(22, True))
    data = [
        ["large readable", "hard shape", "UI-safe sparkle", "context readable"],
        ["few, cheap", "medium", "mostly baked", "tiered"],
        ["packed atlas", "clean masks", "sequence atlas", "variant assets"],
        ["overdraw", "grayscale", "timing", "scene test"],
    ]
    for ri, row in enumerate(rows):
        y = y0 + (ri+1) * rh
        d.rectangle([x0, y, x0+cw, y+rh], fill="#F1F5F9", outline="#CBD5E1", width=2)
        d.text((x0+20, y+38), row, fill="#111827", font=_font(21, True))
        for ci, value in enumerate(data[ri]):
            x = x0 + (ci+1) * cw
            fill = ["#ECFDF5", "#EFF6FF", "#FFF7ED", "#F5F3FF"][ci]
            d.rectangle([x, y, x+cw, y+rh], fill=fill, outline="#CBD5E1", width=2)
            d.text((x+18, y+38), value, fill="#111827", font=_font(20, True))
    d.text((235, 850), "Không có một shader đúng cho mọi style; production chọn tradeoff theo context.", fill="#42505C", font=_font(25, True))
    _save(img, "v27_production_strategy_matrix.png")

    # Light/energy as simple falloff curves.
    img, d = _canvas("Light / Energy / Exposure", "Đọc emissive bằng hierarchy và falloff thay vì ảnh flare cinematic.")
    gx0, gy0, gx1, gy1 = 180, 750, 1020, 250
    d.line([gx0, gy0, gx1, gy0], fill="#1F2937", width=4)
    d.line([gx0, gy0, gx0, gy1], fill="#1F2937", width=4)
    curves = [("linear fade", lambda x: 1-x, "#64748B"), ("inverse square-ish", lambda x: 1/(1+8*x*x), "#E66A2C"), ("art-directed core", lambda x: max(0, (1-x)**4), "#2563EB")]
    for label, fn, color in curves:
        pts = []
        for i in range(280):
            x = i / 279
            y = max(0, min(1, fn(x)))
            pts.append((gx0 + x*(gx1-gx0), gy0 - y*(gy0-gy1)))
        d.line(pts, fill=color, width=5)
    for idx, (label, _, color) in enumerate(curves):
        y = 320 + idx * 52
        d.line([1120, y, 1205, y], fill=color, width=7)
        d.text((1225, y-17), label, fill=color, font=_font(23, True))
    d.rounded_rectangle([1120, 560, 1640, 720], radius=14, fill="#FFFFFF", outline="#CBD5E1", width=3)
    d.text((1150, 595), "core + mid + outer", fill="#111827", font=_font(26, True))
    d.text((1150, 642), "test with bloom OFF first", fill="#DC2626", font=_font(24, True))
    _save(img, "v27_light_energy_curves.png")


def _add_toc(doc):
    head = base.p(doc, "Mục lục", size=18.2, font=base.HEAD, bold=True, color=base.INK, after=10, style="Heading 1")
    head.paragraph_format.line_spacing = 1.04
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\f C \\h \\z'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "Mở file bằng Word và update field nếu mục lục chưa hiển thị."
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, text, fld_end])


def _set_metadata(doc):
    props = doc.core_properties
    props.title = TITLE
    props.subject = SUBTITLE
    props.author = "ForgePilot / VFXSamplePack"
    props.keywords = "Material VFX, Realtime VFX, Unreal Material, Niagara, Shader, HLSL, SubUV, SDF, Overdraw"
    props.comments = f"{VERSION_LABEL} generated from VFXSamplePack build pipeline."
    props.category = "Technical VFX Education"
    props.language = "vi-VN"


def _add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F5FBF")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([color, underline])
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _hyper_bullet(doc, label, url, suffix):
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Cm(0.6)
    para.paragraph_format.first_line_indent = Cm(-0.24)
    para.paragraph_format.space_after = Pt(4)
    para.paragraph_format.line_spacing = 1.16
    run = para.add_run("- ")
    base.set_font(run, size=10.0, name=base.BODY)
    _add_hyperlink(para, label, url)
    run = para.add_run(" - " + suffix)
    base.set_font(run, size=10.0, name=base.BODY)


def references_v27(doc):
    head = base.h1(doc, "Nguồn tham khảo và hướng đọc tiếp")
    _add_toc_entry_field(head, "Nguồn tham khảo và hướng đọc tiếp")
    base.body(doc, "Các nguồn dưới đây dùng để định hướng học và kiểm tra thuật ngữ. Khi đọc tiếp, hãy đọc theo vấn đề production đang gặp thay vì đọc dàn trải.")
    refs = [
        ("The Book of Shaders", "https://thebookofshaders.com/", "pixel shader, tọa độ, shaping function, pattern và noise."),
        ("Unity advanced VFX e-book", "https://unity.com/en/resources/creating-advanced-vfx-unity6", "Shader Graph, VFX Graph, render pipelines và workflow VFX hiện đại."),
        ("Unreal Material Editor docs", "https://dev.epicgames.com/documentation/en-us/unreal-engine/unreal-engine-material-editor-user-guide?application_version=5.6", "Material Editor, node graph và workflow material trong Unreal."),
        ("Unreal Niagara docs", "https://dev.epicgames.com/documentation/unreal-engine/creating-visual-effects-in-niagara-for-unreal-engine", "renderer, particle data, SubUV và Niagara system workflow."),
        ("Real-Time Rendering resources", "https://www.realtimerendering.com/", "graphics pipeline, rasterization, texture, shading và GPU thinking."),
        ("RealTimeVFX forum", "https://realtimevfx.com/", "case production, breakdown, optimization, SubUV, VAT và shader discussions."),
        ("VFXSamplePack / ForgePilot Knowledge", "https://github.com/", "nguồn thực hành nội bộ trong project hiện tại; thay link public nếu repository được publish."),
    ]
    for label, url, suffix in refs:
        _hyper_bullet(doc, label, url, suffix)

    base.h2(doc, "Suggested Reading Path")
    for text in [
        "Nếu đang học visual math từ đầu, bắt đầu với The Book of Shaders: tọa độ, shaping function, pattern và noise. Đừng đọc hết một lượt; hãy dựng lại từng khái niệm bằng grayscale trước.",
        "Nếu đang làm trong Unreal, đọc Material Editor docs song song với một material thật trong project. Mục tiêu là nối thuật ngữ trong docs với node, parameter, blend mode và debug view trong graph.",
        "Nếu đang làm particle/system workflow, đọc Niagara docs sau khi đã hiểu material contract. Niagara nên gửi dữ liệu rõ; Material nên tiêu thụ dữ liệu đó theo lane rõ.",
        "Nếu đang tối ưu hoặc chuẩn bị production review, dùng Real-Time Rendering resources và RealTimeVFX forum để đối chiếu cost, overdraw, breakdown và các failure pattern đã gặp trong project thật.",
    ]:
        base.body(doc, text)

    base.h2(doc, "How to Use These Sources")
    base.body(doc, "Một nguồn tốt không thay thế được test trong scene thật. Khi đọc bất kỳ tài liệu nào, hãy ghi lại ba thứ: khái niệm giải quyết vấn đề gì, debug view nào chứng minh nó đúng, và tradeoff production nằm ở đâu. Cách đọc này giữ tài liệu thành công cụ làm việc thay vì danh sách link.")


def rosetta_stone_v27(doc):
    head = base.h1(doc, "Rosetta Stone: Node Graph, HLSL và Visual")
    _add_toc_entry_field(head, "Rosetta Stone: Node Graph, HLSL và Visual")
    base.body(doc, "Phần này là cầu nối giữa cách đọc trực quan và cách đọc kỹ thuật. Người đọc có thể bắt đầu từ node graph như một câu văn hình ảnh, rồi đối chiếu dần với HLSL để hiểu logic tương đương và biết visual output nên debug như thế nào.")
    base.note(doc, "Foundation Tip", "Nếu code làm bạn sợ, hãy đọc HLSL như pseudocode của node graph. Mỗi dòng nên trả lời: dữ liệu nào đi vào, bị biến đổi ra sao, và output nào cần xem bằng grayscale.", False)
    base.note(doc, "Expert Deep-Dive", "Bảng này không thay thế profiling. Nó chỉ giúp map logic. Trên GPU thật, branch, texture fetch, derivative, precision và platform path vẫn quyết định cost cuối.", True)

    base.h2(doc, "Mask and Shaping Translation")
    _comparison_table(doc, ["Node Graph", "HLSL Sketch", "Visual Result"], [
        ("TextureSample Alpha -> Power -> Saturate", "mask = saturate(pow(alpha, power));", "Core chặt hơn, midtone ít hơn; dùng để làm flare/dissolve rõ shape."),
        ("Distance(UV, Center) -> SmoothStep", "ring = 1 - smoothstep(r, r + soft, abs(d - radius));", "Ring mềm, edge có width và softness rõ."),
        ("Noise -> SmoothStep(A,B) -> Edge Band", "edge = smoothstep(a,b,n) - smoothstep(b,c,n);", "Dissolve có cạnh đọc được thay vì chỉ thủng ngẫu nhiên."),
    ])

    doc.add_page_break()
    base.h2(doc, "Space and Motion Translation")
    _comparison_table(doc, ["Node Graph", "HLSL Sketch", "Visual Result"], [
        ("TexCoord + Panner", "uv = uv + direction * speed * time;", "Texture trượt có hướng; hợp flow, smoke, fire sheet."),
        ("PolarUV -> Add Time -> Frac", "phase = frac(angle + time * speed);", "Swirl/portal chạy quanh tâm; cần giấu seam."),
        ("WorldPos - ImpactPos -> Length", "d = length(worldPos - impactPos);", "Ripple hoặc proximity mask bám world/impact point."),
    ])

    base.h2(doc, "Debug Translation")
    _comparison_table(doc, ["Symptom", "Likely Debug View", "First Fix"], [
        ("Effect biến mất trên nền sáng", "Emissive without bloom + value mask", "Tăng hierarchy shape, không chỉ tăng intensity."),
        ("SubUV mask chạy xuyên nhiều frame", "Local UV R/G trong từng frame", "Tạo local frame UV thay vì dùng atlas UV trực tiếp."),
        ("WPO trượt khi mesh rotate", "Vector direction color + space label", "Kiểm tra local/world/tangent space trước khi tune amplitude."),
    ])


def _vocab_table(doc, rows):
    for term, meaning, debug in rows:
        para = base.p(doc, "", size=9.1, font=base.BODY, before=4, after=5, color=base.INK)
        para.paragraph_format.left_indent = Cm(0.35)
        para.paragraph_format.right_indent = Cm(0.35)
        para.paragraph_format.line_spacing = 1.12
        base.shade(para, "F7F9FB")
        r1 = para.add_run(f"{term}. ")
        base.set_font(r1, size=9.2, name=base.SANS, bold=True, color=base.BLUE)
        r2 = para.add_run(f"{meaning} ")
        base.set_font(r2, size=9.1, name=base.BODY, color=base.INK)
        r3 = para.add_run("Debug: ")
        base.set_font(r3, size=9.1, name=base.SANS, bold=True, color=base.ORANGE)
        r4 = para.add_run(debug)
        base.set_font(r4, size=9.1, name=base.BODY, color=base.INK)


def learning_structure_v27(doc):
    head = base.h1(doc, "How This Book Is Structured")
    _add_toc_entry_field(head, "How This Book Is Structured")
    base.body(
        doc,
        "Cuốn sách này được chia theo đường học tăng dần độ khó. Đầu sách giải thích cách nhìn hình ảnh như dữ liệu: value, mask, range, alpha, grayscale và threshold. Giữa sách chuyển sang vật lý thị giác, shader graph, renderer, cost và debug production. Cuối sách là case study, failure atlas và math node atlas để tra cứu nhanh khi làm việc thật.",
    )
    base.note(
        doc,
        "Reading Rule",
        "Nếu một thuật ngữ còn mơ hồ, đừng nhảy thẳng vào node. Hãy hỏi ba câu: dữ liệu này nằm trong range nào, nó điều khiển lane nào, và debug view nào chứng minh nó đúng?",
        False,
    )
    _comparison_table(
        doc,
        ["Chapter Block", "Mục tiêu", "Độ khó"],
        [
            ("Chapter 01 - Foundations", "Đọc hình như dữ liệu: grayscale, alpha, mask, range, curve, vector và wave.", "Foundation-friendly"),
            ("Chapter 02 - Visual Physics", "Hiểu ánh sáng, màu, fire, water, impact, noise và sampling ở mức đủ dùng cho art direction.", "Foundation -> Intermediate"),
            ("Chapter 03 - Defensive Shader", "Viết graph khó hỏng: clamp, safe divide, SafeNormalize, power/log guard và profiling cơ bản.", "Intermediate"),
            ("Chapter 04 - Renderer Contract", "Hiểu material gặp renderer như thế nào: blend, depth, sorting, Niagara data và production constraints.", "Intermediate -> Advanced"),
            ("Chapter 05 - Case and Failure Atlas", "Đọc graph thật, tách lane, nhìn debug view và truy nguyên artifact theo triệu chứng.", "Production-heavy"),
            ("Appendix - Math Node Atlas", "Tra cứu node theo ý nghĩa, VFX use, lỗi thường gặp và mini exercise.", "Reference / Advanced"),
        ],
    )
    base.body(
        doc,
        "Cách đọc khuyến nghị: đọc tuần tự đến hết Chapter 03 nếu bạn mới học Material VFX. Từ Chapter 04 trở đi, hãy vừa đọc vừa mở graph thật để đối chiếu. Appendix không cần đọc một mạch; dùng nó như từ điển khi gặp node hoặc lỗi cụ thể.",
    )


def technical_vocabulary_v27(doc):
    head = base.h1(doc, "Core Technical Vocabulary")
    _add_toc_entry_field(head, "Core Technical Vocabulary")
    base.body(
        doc,
        "Phần này đặt nền cho các chương sau. Những từ dưới đây xuất hiện liên tục trong material, shader, Niagara và renderer debug. Hiểu đúng chúng sẽ giúp bạn đọc graph chậm hơn một chút ở đầu sách, nhưng nhanh hơn rất nhiều khi sửa lỗi production.",
    )
    base.h2(doc, "Visual Value and Mask Terms")
    _vocab_table(
        doc,
        [
            ("Grayscale", "Ảnh hoặc debug view chỉ dùng độ sáng từ đen đến trắng. Đen thường là 0, trắng là 1, xám là giá trị ở giữa.", "Dùng để kiểm mask trước khi thêm màu, bloom hoặc particle color. Nếu grayscale sai, final màu đẹp cũng chỉ đang che lỗi."),
            ("Value", "Độ sáng hoặc cường độ đọc được của một pixel. Value thường quan trọng hơn hue vì mắt nhận silhouette và contrast trước.", "Khi effect mất shape, kiểm value hierarchy trước: core có sáng hơn glow không, edge có tách khỏi background không."),
            ("Alpha", "Kênh điều khiển độ trong suốt hoặc coverage. Trong VFX, alpha thường là dữ liệu shape chứ không chỉ là opacity cuối.", "Quên nhân alpha là lỗi nền tảng phổ biến: color/emissive vẫn hiện nhưng silhouette, fade hoặc soft edge bị sai."),
            ("Mask", "Một trường giá trị 0..1 dùng để chọn vùng nào được tác động. Mask có thể đến từ texture, distance, noise, vertex color hoặc math.", "Luôn debug raw mask, remapped mask và shaped mask riêng. Đừng debug thẳng final beauty."),
            ("Threshold", "Ngưỡng cắt. Giá trị dưới ngưỡng bị loại, trên ngưỡng được giữ hoặc chuyển trạng thái.", "Threshold quá cứng tạo edge răng cưa hoặc flicker. Dùng SmoothStep/fwidth khi edge cần mềm và ổn định theo pixel."),
            ("Artifact", "Lỗi nhìn thấy trên render: flicker, seam, banding, overbright flash, sorting sai, shimmer, frame SubUV sai hoặc edge bị răng cưa.", "Đọc artifact như triệu chứng, không đọc như phán đoán. Triệu chứng dẫn tới debug view, debug view mới dẫn tới nguyên nhân."),
        ],
    )
    base.h2(doc, "Space, Sampling and Renderer Terms")
    _vocab_table(
        doc,
        [
            ("UV", "Tọa độ 2D dùng để đọc texture. U là ngang, V là dọc; thường nằm trong 0..1 nhưng có thể tile hoặc offset.", "UV sai thường tạo texture trượt, seam, scale lệch hoặc SubUV đọc nhầm frame."),
            ("Local UV", "UV đã được quy về một frame hoặc một vùng nhỏ trong atlas, thay vì toàn bộ texture lớn.", "Trong SubUV, atlas UV đúng chưa đủ. Mask nội bộ mỗi frame cần local UV đúng."),
            ("SDF", "Signed Distance Field: giá trị cho biết khoảng cách tới biên. Gần biên là vùng quan trọng để tạo ring, glow, outline hoặc soft edge.", "SDF tốt cho edge có width rõ. Nếu ring dày/mỏng bất thường, kiểm scale, radius, width và softness."),
            ("Derivative", "Độ thay đổi của giá trị giữa các pixel gần nhau. Trong shader, derivative giúp biết edge đang sắc hay mềm theo screen space.", "Dùng cho anti-alias edge procedural. Thiếu derivative/fwidth thường gây shimmer khi camera xa."),
            ("Mip", "Phiên bản texture nhỏ hơn được GPU chọn khi texture ở xa hoặc bị thu nhỏ trên màn hình.", "Mip sai làm texture blur quá mức hoặc shimmer. Noise tần số cao cần kiểm mip và sampling."),
            ("Overdraw", "Một pixel bị vẽ nhiều lần bởi nhiều lớp alpha card hoặc particle chồng lên nhau.", "Shader ít instruction vẫn có thể đắt nếu card quá lớn hoặc layer count cao, nhất là smoke/fire translucent."),
        ],
    )
    base.h2(doc, "Production Debug Terms")
    _vocab_table(
        doc,
        [
            ("Debug View", "Một cách nhìn tạm để cô lập dữ liệu: raw mask, alpha, normal, depth fade, frame index, overdraw hoặc shader complexity.", "Debug view tốt trả lời một câu hỏi nhỏ. Nếu view cố trả lời mọi thứ, nó lại thành final view thứ hai."),
            ("Lane", "Một luồng ý nghĩa trong material: UV lane, mask lane, color lane, opacity lane, motion lane, interaction lane hoặc cost lane.", "Khi graph rối, tách lane giúp biết node nào tạo hình, node nào chỉ đổi range, node nào gây cost."),
            ("Blend Mode", "Cách material trộn vào framebuffer: opaque, masked, translucent, additive... Mỗi mode có luật depth/sorting/cost khác nhau.", "Nhiều lỗi không nằm ở texture mà nằm ở blend/depth. Additive đẹp trong preview nhưng có thể cháy trắng trong scene thật."),
            ("Depth Fade", "Kỹ thuật làm mềm giao tuyến giữa particle/material và geometry dựa trên khoảng cách depth.", "Nếu particle cắt vào mesh quá rõ, kiểm depth fade. Nếu bị mất thân effect, fade distance có thể quá lớn."),
            ("Precision", "Độ chính xác số học của kiểu dữ liệu. Half/float, mobile/desktop và giá trị rất lớn/nhỏ có thể cho kết quả khác nhau.", "Precision error thường xuất hiện như jitter, banding, NaN flash hoặc mask đổi bất thường theo camera/platform."),
            ("NaN / Inf", "Giá trị số không hợp lệ hoặc vô hạn, thường sinh ra từ divide by zero, normalize vector zero, log âm hoặc power với input xấu.", "Nếu effect flash trắng/đen một frame, kiểm safe divide, SafeNormalize, clamp trước Power/Log."),
        ],
    )
    base.note(
        doc,
        "Foundation Tip",
        "Khi mới học, hãy tự ép mình export mọi mask ra grayscale. Bạn không cần hiểu toàn bộ graph ngay; chỉ cần biết giá trị đang đi từ đâu, nằm trong range nào, và có đúng shape chưa.",
        False,
    )
    base.note(
        doc,
        "Expert Deep-Dive",
        "Ở production level, thuật ngữ không chỉ là định nghĩa. Mỗi term cần gắn với một debug action: xem buffer nào, isolate lane nào, đo cost nào, và giả thuyết nào bị loại sau mỗi lần kiểm.",
        True,
    )


CHAPTER_BLOCKS = {
    1: ("Chapter 01", "Foundations: Reading Images as Data", "Từ đây sách đi chậm và giải thích kỹ. Mục tiêu là biến hình ảnh thành dữ liệu có thể đọc: value, mask, grayscale, threshold, range, vector và motion.", "Foundation-friendly", ["Grayscale", "Alpha", "Mask", "Range 0..1", "Curve", "Vector", "Phase"]),
    9: ("Chapter 02", "Visual Physics: Light, Color and Natural Cues", "Các bài này nối visual math với cảm giác vật lý: ánh sáng, màu, heat, smoke, water, impact, noise và sampling.", "Foundation -> Intermediate", ["Value hierarchy", "Temperature", "Refraction", "Impulse", "Noise", "Sampling", "Moiré"]),
    16: ("Chapter 03", "Defensive Shader and Numerical Stability", "Từ đây bắt đầu viết graph như production code: dữ liệu xấu phải được chặn trước khi nó lan ra final render.", "Intermediate", ["Safe divide", "SafeNormalize", "Power guard", "Log guard", "Precision", "Profiling"]),
    21: ("Chapter 04", "Renderer Contract: Graph, Lanes, Blend and Depth", "Material không sống một mình. Nó phải ký hợp đồng với renderer, Niagara, blend mode, depth, sorting và platform budget.", "Intermediate -> Advanced", ["Shader language", "Lane", "Blend mode", "Depth fade", "Sorting", "Platform budget"]),
    26: ("Chapter 05", "Production Cases, Validation and Failure Atlas", "Phần này đọc graph như một ca debug thật: tách lane, xem output preview, kiểm mask grayscale, tìm artifact và xác nhận bằng checklist.", "Production-heavy", ["Material graph", "Output preview", "Failure atlas", "Artifact", "Debug view", "Checklist"]),
    31: ("Chapter 06", "Applied Material Concepts", "Các khái niệm cốt lõi được đào lại ở mức thực hành hơn: texture, UV, SDF, polar, vector math, transforms, mip, Niagara và WPO.", "Advanced practice", ["Texture", "Local UV", "SDF", "Polar UV", "Transform", "Mip", "WPO"]),
    48: ("Appendix A", "Math Node Atlas", "Phần tra cứu nhanh. Đọc theo ý nghĩa, VFX use, lỗi thường gặp và mini exercise; không cần học alphabet.", "Reference / Advanced", ["Clamp", "SmoothStep", "Power", "Frac", "Dot", "Lerp", "Transform"]),
}


def chapter_block_divider_v27(doc, marker, title, summary, difficulty, terms):
    doc.add_page_break()
    p = base.p(doc, marker.upper(), size=9.8, font=base.SANS, bold=True, color=base.ORANGE, before=4, after=4)
    _add_toc_entry_field(p, f"{marker}: {title}", level=1)
    head = base.p(doc, title, size=19.4, font=base.HEAD, bold=True, color=base.BLUE, after=8, style="Heading 1")
    head.paragraph_format.line_spacing = 1.04
    base.body(doc, summary)
    base.note(doc, "Difficulty", difficulty, False)
    base.note(doc, "Key Terms", ", ".join(terms), False)


def interactive_companion_v27(doc):
    head = base.h1(doc, "Interactive Companion Links")
    _add_toc_entry_field(head, "Interactive Companion Links")
    base.body(doc, "Các link dưới đây được thiết kế như điểm móc cho bản public GitHub và ShaderLex. Khi URL chính thức sẵn sàng, hãy cập nhật lại để người đọc mở shader mẫu, xem debug view và dựng lại bài tập ngay.")
    _link_paragraph(doc, "Realtime VFX Brain repository", "https://github.com/search?q=realtime-vfx-brain&type=repositories", "repo public chứa PDF, note, diagram và companion material.")
    _link_paragraph(doc, "ShaderLex companion", "https://github.com/search?q=ShaderLex&type=repositories", "điểm vào cho shader mẫu tương tác sau khi publish.")
    _link_paragraph(doc, "RealTimeVFX forum", "https://realtimevfx.com/", "đọc breakdown thật và đối chiếu failure pattern.")
    base.note(doc, "Foundation Tip", "Khi có companion shader, hãy mở bản grayscale trước. Đừng bắt đầu bằng final color; final color thường giấu lỗi range, alpha và mask.", False)
    base.note(doc, "Expert Deep-Dive", "Một companion tốt nên cho phép bật/tắt từng lane: raw input, remap, shaping, mask, emissive, alpha, cost proxy. Nếu chỉ có final preview, nó chưa đủ để dạy debug.", True)


def chapter_v27(doc, idx, data):
    chapter_title = _title_case_heading(data["title"])
    head = base.h1(doc, chapter_title, idx)
    _add_toc_entry_field(head, chapter_title, level=2)
    if data.get("image"):
        base.figure(doc, data["image"][0], data["image"][1], width=10.2)
    is_atlas_entry = data.get("title", "").startswith("Math Atlas:")
    for title, paras in data["sections"]:
        display_title = _title_case_heading(title)
        if is_atlas_entry and display_title in TITLE_EXACT:
            base.p(doc, display_title, size=11.2, font=base.SANS, bold=True, color=base.INK, before=8, after=3)
        else:
            base.h2(doc, display_title)
        for text in paras:
            base.body(doc, text)
    for title, text, warning in data.get("notes", []):
        base.note(doc, _title_case_heading(title), text, warning)
    if data.get("case"):
        base.p(doc, "Production Diagnostic", size=11.2, font=base.SANS, bold=True, color=base.INK, before=8, after=3)
        for text in data["case"]:
            base.body(doc, text)
    if data.get("code"):
        base.code(doc, data["code"], "Implementation Logic")


def closing_v27(doc):
    doc.add_page_break()
    base.h1(doc, "Kết Luận: Xây Phản Xạ Bằng Chính Mắt Mình")
    base.body(
        doc,
        "Người học Material thường muốn đi nhanh tới kết quả cuối cùng. Điều đó tự nhiên, vì VFX là nghề của hình ảnh. Nhưng tốc độ thật không đến từ việc nối node nhanh hơn; nó đến từ khả năng nhìn đúng vấn đề ngay từ đầu. Trước khi thêm node, hãy nhìn effect như dữ liệu. Trước khi tăng emissive, hãy kiểm tra hierarchy. Trước khi gọi một effect là xấu, hãy hỏi nó đang hỏng ở shape, mask, motion, color, interaction hay cost.",
    )
    base.body(
        doc,
        "Một field guide tốt không nên chỉ là danh sách công thức. Công thức cần thiết, nhưng công thức phải nằm trong câu chuyện của visual và production. Người học cần biết vì sao một kỹ thuật tồn tại, nó giải quyết vấn đề gì, nó hỏng trong trường hợp nào, và khi nào nên bỏ nó để dùng một giải pháp rẻ hơn.",
    )
    base.body(
        doc,
        "Con đường học Shader VFX nên được xây bằng nhiều bài nhỏ nhưng thật: dựng một Mask Grayscale, làm một Edge Band, đọc một Atlas bằng NumberGrid, tạo một Additive Flare rồi kiểm tra trên nền sáng, làm một Fire Card có DepthFade, viết một hàm HLSL nhỏ rồi so với Node Graph. Mỗi bài tập để lại một phần phản xạ. Khi phản xạ đủ nhiều, bạn không chỉ biết làm một hiệu ứng cụ thể; bạn biết cách nghĩ khi gặp một hiệu ứng mới.",
    )


def _chapter_dict(title, image=None, caption="", sections=None, note=None, case=None, code=None):
    return {
        "title": title,
        "image": (image, caption) if image else None,
        "sections": sections or [],
        "notes": [note] if note else [],
        "case": case or [],
        "code": code or [],
    }


TITLE_REPLACEMENTS = {
    "Ghi nhớ": "Key Takeaway",
    "Ghi Nhớ": "Key Takeaway",
    "Lỗi thường gặp": "Common Failure Mode",
    "Lỗi Thường Gặp": "Common Failure Mode",
    "Cảnh báo production": "Production Risk",
    "Cảnh Báo Production": "Production Risk",
    "Direction của Impact": "Impact Directionality",
    "Foam là nơi nước tương tác": "Foam Interaction Zone",
    "Noise là random có tương quan không gian": "Spatially Correlated Noise",
    "The Matrix of Art - toán học như cách nặn cảm giác": "Visual Math as Art Direction",
    "The Matrix of Art - Toán Học như Cách Nặn Cảm Giác": "Visual Math as Art Direction",
    "The Matrix Of Art - Toán Học như Cách Nặn Cảm Giác": "Visual Math as Art Direction",
    "Từ Một Giá Trị 0-1 Tới Một Visual Đọc Được": "Từ Giá Trị 0-1 đến Visual Đọc Được",
    "Range và Normalization: Vật Liệu Bắt Đầu từ 0 đến 1": "Range và Normalization: Material Bắt Đầu từ 0 đến 1",
    "Remap là Kỹ Năng Sống Còn": "Remap là Kỹ Năng Nền Tảng",
    "Physics là Constraint, không Phải Nhà Tù": "Physics là Constraint, không Phải Luật Tuyệt Đối",
    "Gameplay Thắng Simulation": "Gameplay Ưu Tiên Hơn Simulation",
    "Ngọn ngành để ứng biến feedback": "Dịch Feedback Thành Quyết Định Kỹ Thuật",
    "Ngọn Ngành để Ứng Biến Feedback": "Dịch Feedback Thành Quyết Định Kỹ Thuật",
    "Impact là Dữ Liệu không Gian và Thời Gian": "Impact Data: Space và Time",
    "Blend Mode là Render Decision": "Blend Mode là Quyết Định Render Path",
    "Đọc Hiểu Chưa Đủ: Cần Rebuild": "Đọc Hiểu Chưa Đủ: Cần Dựng Lại",
    "Forum Threads như Failure Pattern Log": "Forum Threads như Nhật Ký Lỗi Production",
    "Breakdown Thể Hiện Thinking, không chỉ Final Result": "Breakdown Thể Hiện Tư Duy, không Chỉ Kết Quả Cuối",
    "SmoothStep là Dao Mổ của Threshold": "SmoothStep là Công Cụ Kiểm Soát Threshold Mềm",
    "Ring là Distance Band, không chỉ là Hình Tròn": "Ring là Dải Khoảng Cách, không Chỉ là Hình Tròn",
    "CPU Chuẩn Bị Công Việc, GPU Nghiền Số Lượng Lớn": "CPU Chuẩn Bị Công Việc, GPU Xử Lý Song Song Khối Lượng Lớn",
    "Noise: Organic không Đồng Nghĩa Random Vô Tội Vạ": "Noise: Organic không Đồng Nghĩa Random Vô Kiểm Soát",
    "Kết luận - viết lại bằng chính mắt của bạn": "Kết Luận: Xây Phản Xạ Bằng Chính Mắt Mình",
    "Ánh sáng vật lý tối giản cho VFX Artist": "Physical Lighting Fundamentals cho VFX Artist",
    "Intensity và energy không giống cảm giác sáng": "Intensity và Energy Không Đồng Nhất Với Perceived Brightness",
    "Scattering và absorption ở mức trực quan": "Scattering và Absorption: Mô Hình Trực Quan",
    "Exposure làm effect đổi đời": "Exposure và Tonemapping Trong Scene Thật",
    "Hàm số như bút vẽ - curve điều khiển cảm giác": "Shaping Functions: Curve Điều Khiển Visual Response",
    "Linear hiếm khi đủ đẹp": "Linear Response và Giới Hạn Art Direction",
    "Độ dốc của curve quyết định vùng nhạy": "Curve Slope Quyết Định Vùng Nhạy",
    "Vector là hướng, lực và quan hệ không gian": "Vector: Hướng, Lực và Quan Hệ Không Gian",
    "Đạo hàm, gradient và cảm giác thay đổi": "Derivative và Gradient: Đọc Tốc Độ Thay Đổi",
    "Đạo hàm không chỉ dành cho toán học trên giấy": "Derivative Trong Shader Practice",
    "Sóng - nhịp, phase, frequency và interference": "Wave: Phase, Frequency và Interference",
    "Time không phải animation; phase mới là animation": "Time, Phase và Animation Loop",
    "Trường lực và vector field - motion có cấu trúc": "Vector Field: Motion Có Cấu Trúc",
    "Màu sắc - từ RGB tới cảm giác nhiệt, hiếm và nguy hiểm": "Color: RGB, Temperature và Visual Rarity",
    "RGB là cách lưu, không phải cách nghĩ duy nhất": "RGB Là Storage Model, Không Phải Art Model Duy Nhất",
    "Value contrast đọc trước hue": "Value Contrast Đọc Trước Hue",
    "Nhiệt, lửa và khói - vật lý đủ dùng cho shader": "Heat, Fire và Smoke: Physics Core Cho Shader",
    "Lửa là vùng phản ứng, không chỉ là texture cam": "Fire: Reaction Zone, Không Chỉ Là Orange Texture",
    "Khói là density và ánh sáng": "Smoke: Density và Light Interaction",
    "Nước, bọt và khúc xạ - surface physics cho VFX": "Water, Foam và Refraction: Surface Physics Cho VFX",
    "Nước đọc bằng normal và reflection nhiều hơn màu": "Water Readability: Normal và Reflection Trước Color",
    "Va chạm, impulse và hit response": "Impact, Impulse và Hit Response",
    "Noise, randomness và xác suất thị giác": "Noise, Randomness và Visual Probability",
    "Random cần distribution": "Random Distribution",
    "Sampling theory - vì sao hình bị răng cưa, rung và moiré": "Sampling Theory: Aliasing, Shimmer và Moiré",
    "Pixel là mẫu, không phải thế giới liên tục": "Pixel Sampling và Continuous Signal",
    "Ổn định số và defensive shader": "Numerical Stability và Defensive Shader",
    "Shader cũng cần phòng thủ": "Defensive Shader Boundaries",
    "Từ vật lý sang art direction - biết đủ để bẻ luật": "Từ Physics Sang Art Direction",
    "Vật lý là điểm tựa, không phải nhà tù": "Physics Là Constraint, Không Phải Nhà Tù",
    "Unreal Math Node Atlas - từ node tới tư duy shader": "Unreal Math Node Atlas: Từ Node Đến Shader Thinking",
    "Cách đọc nhóm Math trong Material Editor": "Reading Math Nodes Trong Material Editor",
    "Cơ bắp của realtime - CPU, GPU và cái giá của pixel": "Realtime Cost Model: CPU, GPU và Pixel Cost",
    "Một Effect không chạy trong sách vở": "Effect Trong Runtime Context",
    "Tối ưu không có nghĩa là làm xấu đi": "Optimization Không Đồng Nghĩa Giảm Chất Lượng",
    "Ngôn ngữ Shader - khi Node Graph không còn đủ": "Shader Language: Khi Node Graph Không Còn Đủ Rõ",
    "Code shader không phải nghi thức thần bí": "Shader Code Là Data Transform",
    "Types, parameters và texture sample": "Types, Parameters và Texture Sampling",
    "Custom node trong Unreal và giới hạn của nó": "Custom Node Trong Unreal và Giới Hạn Production",
    "Mask, UV và motion - ba lớp dễ bị sửa nhầm": "Mask, UV và Motion: Ba Lane Dễ Sửa Nhầm",
    "Blend, emissive và depth - nơi material gặp renderer": "Blend, Emissive và Depth: Material Gặp Renderer",
    "Blend mode là quyết định render, không chỉ là mỹ thuật": "Blend Mode Là Render Decision",
    "Production stories - mobile, stylized, iGaming và AAA": "Production Strategies: Mobile, Stylized, iGaming và AAA",
    "Mobile không tha thứ cho sự hào nhoáng không cần thiết": "Mobile Budget và Runtime Discipline",
    "Stylized không có nghĩa là đơn giản": "Stylized VFX Không Đồng Nghĩa Đơn Giản",
    "AAA readability và bài học khi scene quá phức tạp": "AAA Readability Trong Scene Phức Tạp",
    "Production wisdom từ cộng đồng RealTimeVFX": "Production Wisdom Từ RealTimeVFX",
    "Đọc forum như đọc nhật ký lỗi của ngành": "Forum Threads Như Failure Pattern Log",
    "Shader Complexity là bản đồ nhiệt, không phải bản án": "Shader Complexity Là Heatmap, Không Phải Bản Án",
    "Breakdown tốt cho thấy tư duy, không chỉ khoe kết quả": "Breakdown Thể Hiện Thinking, Không Chỉ Final Result",
    "Case study - đọc một Material mẫu trong Unreal": "Case Study: Đọc Material Mẫu Trong Unreal",
    "Case 1: M_Flare như một material additive có kiểm soát": "Case 1: M_Flare Additive Material Có Kiểm Soát",
    "Case 2: SubUV và NumberGrid": "Case 2: SubUV và NumberGrid",
    "Lộ trình luyện tập - từ hiểu tới phản xạ": "Practice Roadmap: Từ Hiểu Đến Phản Xạ",
    "Đọc không đủ, phải dựng lại": "Đọc Hiểu Chưa Đủ: Cần Rebuild",
    "Texture không chỉ là hình ảnh": "Texture Không Chỉ Là Hình Ảnh",
    "Texture mỹ thuật và Texture dữ liệu khác nhau ở mục đích": "Art Texture và Data Texture Khác Nhau Ở Mục Đích",
    "Tọa độ UV - mặt phẳng nhỏ điều khiển rất nhiều thứ": "UV Coordinates: Mặt Phẳng Nhỏ Điều Khiển Nhiều Thứ",
    "UV là bản đồ, không phải chỉ là cặp số U và V": "UV Là Coordinate Map, Không Chỉ Là U và V",
    "Distance, SDF và tư duy hình học 2D": "Distance, SDF và 2D Geometry Thinking",
    "Distance là cách hỏi một điểm cách một hình bao xa": "Distance Field: Đo Khoảng Cách Tới Shape",
    "Ring không phải chỉ là hình tròn": "Ring Là Distance Band, Không Chỉ Là Hình Tròn",
    "Polar UV, góc và chuyển động xoáy": "Polar UV: Angle, Radius và Swirl Motion",
    "Seam là cái giá của angle": "Angle Seam Là Chi Phí Của Polar Mapping",
    "Twirl và swirl là biến dạng tọa độ": "Twirl và Swirl Là Coordinate Distortion",
    "Dot Product là câu hỏi hai hướng giống nhau đến mức nào": "Dot Product Đo Alignment Giữa Hai Hướng",
    "Matrix và Space Transform - đổi hệ quy chiếu để giải đúng bài toán": "Matrix và Space Transform: Đổi Hệ Quy Chiếu",
    "Space là bối cảnh của một con số": "Space Là Context Của Một Giá Trị",
    "GPU và CPU - cơ bắp thật phía sau một pixel": "GPU và CPU: Cost Thật Phía Sau Pixel",
    "HLSL/GLSL - khi Node Graph không còn đủ rõ": "HLSL/GLSL: Khi Node Graph Không Còn Đủ Rõ",
    "Uniform, varying và câu chuyện dữ liệu đi vào shader": "Uniform, Varying và Shader Inputs",
    "Blend Mode và Sorting - nơi Material gặp Renderer": "Blend Mode và Sorting: Material Gặp Renderer",
    "Opaque, Masked, Translucent và Additive không cùng một thế giới": "Opaque, Masked, Translucent và Additive Là Các Render Path Khác Nhau",
    "Niagara và Material - hợp đồng dữ liệu phải rõ": "Niagara và Material: Data Contract Rõ Ràng",
    "Particle Color không chỉ là màu": "Particle Color Không Chỉ Là Color",
    "Dynamic Parameter nên có chuẩn đặt tên": "Dynamic Parameter Naming Standard",
    "Ai sở hữu timing?": "Timing Ownership",
    "WPO - khi Material bắt đầu đẩy hình học": "WPO: Khi Material Đẩy Geometry",
    "POM, Height và ảo giác chiều sâu": "POM, Height và Depth Illusion",
    "Height Map là bản đồ độ cao giả": "Height Map Là Encoded Height Data",
    "Noise - organic không có nghĩa là ngẫu nhiên vô tội vạ": "Noise: Organic Không Đồng Nghĩa Random Vô Tội Vạ",
    "Color, HDR và Bloom - ánh sáng cần hierarchy": "Color, HDR và Bloom: Light Hierarchy",
    "HDR cho phép sáng hơn 1, nhưng không miễn trách nhiệm art direction": "HDR Cho Phép Giá Trị > 1, Nhưng Vẫn Cần Art Direction",
    "Tonemapper và Exposure có thể đổi cảm giác effect": "Tonemapper và Exposure Thay Đổi Effect Perception",
    "Optimization Cookbook - tối ưu theo nguyên nhân": "Optimization Cookbook: Tối Ưu Theo Nguyên Nhân",
    "Nếu vấn đề là Overdraw": "Khi Bottleneck Là Overdraw",
    "Nếu vấn đề là Texture Fetch": "Khi Bottleneck Là Texture Fetch",
    "Nếu vấn đề là Shader Instruction": "Khi Bottleneck Là Shader Instruction",
    "Nếu vấn đề là CPU/Niagara": "Khi Bottleneck Là CPU/Niagara",
    "Cách đọc chung của Math Node Atlas": "Math Node Atlas Reading Framework",
    "Đọc theo lane thay vì đọc theo alphabet": "Đọc Theo Lane, Không Theo Alphabet",
    "Debug bằng giá trị trung gian": "Debug Bằng Intermediate Values",
    "Số, range và normalization - vật liệu bắt đầu từ 0 đến 1": "Range và Normalization: Vật Liệu Bắt Đầu Từ 0 Đến 1",
    "Vì sao 0..1 xuất hiện khắp shader": "Vì Sao 0..1 Xuất Hiện Khắp Shader",
    "Remap là kỹ năng sống còn": "Remap Là Kỹ Năng Sống Còn",
    "Precision và lỗi nhỏ thành visual lớn": "Precision và Lỗi Nhỏ Thành Visual Lớn",
}


TECH_TERMS = {
    "shader": "Shader",
    "material": "Material",
    "node graph": "Node Graph",
    "custom node": "Custom Node",
    "texture": "Texture",
    "texture sample": "Texture Sample",
    "sampling": "Sampling",
    "sample": "Sample",
    "renderer": "Renderer",
    "render": "Render",
    "realtime": "Realtime",
    "effect": "Effect",
    "emissive": "Emissive",
    "intensity": "Intensity",
    "energy": "Energy",
    "exposure": "Exposure",
    "tonemapper": "Tonemapper",
    "absorption": "Absorption",
    "scattering": "Scattering",
    "gradient": "Gradient",
    "curve": "Curve",
    "motion": "Motion",
    "depth": "Depth",
    "height": "Height",
    "noise": "Noise",
    "random": "Random",
    "mask": "Mask",
    "opacity": "Opacity",
    "value": "Value",
    "hue": "Hue",
    "saturation": "Saturation",
    "frequency": "Frequency",
    "phase": "Phase",
    "interference": "Interference",
    "vector": "Vector",
    "matrix": "Matrix",
    "space": "Space",
    "local": "Local",
    "atlas": "Atlas",
    "frame": "Frame",
    "flipbook": "Flipbook",
    "runtime": "Runtime",
    "production": "Production",
    "mobile": "Mobile",
    "stylized": "Stylized",
}


def _apply_tech_caps(text):
    text = TITLE_REPLACEMENTS.get(text, text)
    # Longer keys first so "Node Graph" is handled before "Node".
    for raw, fixed in sorted(TECH_TERMS.items(), key=lambda kv: len(kv[0]), reverse=True):
        text = re.sub(rf"(?<![A-Za-z]){re.escape(raw)}(?![A-Za-z])", fixed, text, flags=re.IGNORECASE)
    return text


TITLE_STOPWORDS = {
    "và", "cho", "của", "trong", "từ", "đến", "như", "khi",
    "để", "ở", "bằng", "theo", "vì", "vào", "ra",
    "trên", "dưới", "giữa", "hay", "hoặc", "hơn", "trước", "sau", "nơi", "qua",
    "là", "có", "với", "không", "chỉ",
    "of", "and", "as", "the",
}

TITLE_PRESERVE = {
    "UV", "CPU", "GPU", "VFX", "AAA", "SDF", "HDR", "VAT", "POM", "WPO", "LOD",
    "HLSL", "GLSL", "HLSL/GLSL", "RGB", "HSV", "sRGB", "iGaming", "fwidth",
    "NaN", "Inf", "2D", "0..1", "0-1", "M_Flare", "SubUV", "DepthFade",
    "SmoothStep", "OneMinus", "DotProduct", "CrossProduct", "VectorToRadialValue",
    "InverseLerp", "RemapValueRange", "ComponentMask", "AppendVector",
    "AddComponents", "Transform3x3Matrix", "TransformToClipSpace",
    "ProjectVectorOntoPlane", "MakeVectorsOrthogonal", "HsvToRgb", "RgbToHsv",
    "FindSaturation", "WrapFloat", "WrapFloat2", "WrapFloat3", "WrapFloat4",
}

TITLE_EXACT = {
    "Visual Math as Art Direction",
    "Ý nghĩa trong graph",
    "VFX use",
    "Failure mode",
    "Mini exercise",
}


def _cap_word(word):
    if not word:
        return word
    if word in TITLE_PRESERVE:
        return word
    if "/" in word:
        return "/".join(_cap_word(part) for part in word.split("/"))
    if "-" in word and not re.match(r"^\d", word):
        return "-".join(_cap_word(part) for part in word.split("-"))
    if any(ch.isdigit() for ch in word) and not word.isalpha():
        return word
    return word[:1].upper() + word[1:]


def _title_case_heading(text):
    text = _apply_tech_caps(text)
    if text in TITLE_EXACT:
        return text
    out = []
    after_break = True
    for token in re.split(r"(\s+)", text):
        if not token or token.isspace():
            out.append(token)
            continue
        leading = re.match(r"^[^\wÀ-ỹ]*", token).group(0)
        trailing = re.search(r"[^\wÀ-ỹ]*$", token).group(0)
        core = token[len(leading):len(token) - len(trailing) if trailing else len(token)]
        if not core:
            out.append(token)
            after_break = "/" in token or token.rstrip().endswith(("-", ":", ";", "("))
            continue
        low = core.lower()
        if core in TITLE_PRESERVE:
            fixed = core
        elif low in TITLE_STOPWORDS and not after_break:
            fixed = low
        else:
            fixed = _cap_word(core)
        out.append(leading + fixed + trailing)
        after_break = token.rstrip().endswith(("-", ":", ";", "/", "("))
    return "".join(out)


def _polish_titles(chapters):
    for chapter in chapters:
        chapter["title"] = _title_case_heading(chapter.get("title", ""))
        sections = []
        for title, paras in chapter.get("sections", []):
            sections.append((_title_case_heading(title), paras))
        chapter["sections"] = sections
        if chapter.get("notes"):
            chapter["notes"] = [(_title_case_heading(title), text, warning) for title, text, warning in chapter["notes"]]


ATLAS_TEXT_OVERRIDES = {
    "Add, Subtract, Multiply, Divide": {
        "role": "Đây là nhóm toán nền để dịch, cân, khuếch đại hoặc chuẩn hóa dữ liệu trong graph. Khi đọc một material phức tạp, hãy xem bốn node này như các thao tác biên tập tín hiệu: Add tạo bias, Subtract tạo sai khác, Multiply đặt trọng số, Divide đưa giá trị về một tỷ lệ có thể kiểm soát.",
        "examples": "Trong VFX, nhóm này thường nằm ở các điểm quyết định readability. Add offset UV hoặc đẩy edge glow; Subtract so threshold hoặc tạo khoảng cách quanh radius; Multiply nhân mask với color, alpha, Particle Color hoặc distortion; Divide normalize khoảng cách theo radius hay kích thước asset. Câu hỏi production là giá trị này đang được đẩy mạnh hơn, bù lệch đi, hay đưa về scale chung.",
        "danger": "Lỗi thường đến từ thứ tự thao tác. Multiply emissive qua nhiều lớp dễ làm flare cháy trắng, Divide với mẫu số nhỏ tạo output bất ổn, còn Add/Subtract đặt trước hay sau Clamp sẽ cho kết quả khác hẳn. Với node nhóm này, hãy debug bằng grayscale sau từng bước thay vì nhìn final color.",
        "exercise": "Dựng một radial mask có radius parameter. Dùng Subtract để tạo sai khác quanh radius, Multiply để tăng contrast, Add để bias vùng nhìn thấy, rồi Divide để normalize distance theo radius. Xuất từng bước ra grayscale để thấy thao tác nào thật sự thay đổi silhouette.",
    },
    "Abs, Sign, OneMinus, Min, Max": {
        "role": "Nhóm này dùng để đảo, gấp, chọn và giới hạn quan hệ giữa các giá trị. Nó giúp biến một tín hiệu đơn giản thành logic có hướng: bên trong hay bên ngoài, trái hay phải, vùng nào thắng, vùng nào bị loại.",
        "examples": "OneMinus thường dùng để đảo Fresnel, alpha hoặc mask edge. Abs biến khoảng cách có dấu thành khoảng cách hai phía nên rất hợp cho stripe đối xứng, slash mark và ring band. Sign giữ thông tin phía của trục, còn Min/Max giúp gộp nhiều mask theo logic lấy vùng tối hơn hoặc sáng hơn.",
        "danger": "OneMinus chỉ sạch khi input nằm trong 0..1; nếu input vượt range, output có thể âm hoặc lớn hơn 1. Abs làm mất thông tin hướng, vì vậy không nên dùng khi còn cần biết bên nào của trục. Min/Max cũng dễ che mất mask yếu hơn nếu không kiểm tra từng input riêng.",
        "exercise": "Tạo một slash mask từ UV.x. Dùng Abs để mirror hai cạnh, OneMinus để đảo vùng chọn, rồi Max để gộp core và edge. Sau đó tắt từng layer để kiểm tra mask nào đang quyết định hình cuối.",
    },
    "Clamp, Saturate, Step, SmoothStep": {
        "role": "Đây là nhóm kiểm soát range và threshold. Nó quyết định dữ liệu được giữ an toàn trong biên, bị cắt cứng, hay chuyển mềm qua một vùng đệm có thể art direct.",
        "examples": "Saturate là lựa chọn nhanh khi cần giữ mask trong 0..1. Step tạo cutout cứng cho masked dissolve hoặc logic bật/tắt. SmoothStep tạo vùng chuyển mềm cho dissolve, foam, soft intersection và shockwave. Clamp dùng khi range cần kiểm soát không phải 0..1, ví dụ distance hoặc depth.",
        "danger": "Clamp quá sớm làm mất dữ liệu mà các bước sau có thể cần. SmoothStep với hai ngưỡng quá gần sẽ thành cạnh sắc và dễ aliasing. Step đặt trên smoke hoặc translucent layer thường làm texture bị cứng giả, trừ khi đó là chủ ý style.",
        "exercise": "Dùng cùng một noise source để tạo ba dissolve: Step cho cạnh cứng, SmoothStep cho cạnh mềm, và Edge Band bằng hiệu của hai SmoothStep. Đặt ba kết quả cạnh nhau để thấy threshold ảnh hưởng thế nào đến cảm giác material.",
    },
    "Power, SquareRoot, Exponential, Logarithm": {
        "role": "Nhóm này không đổi dữ liệu gốc mà đổi cách dữ liệu được phân bố. Trong production, nó là công cụ tạo cảm giác core, glow, falloff và contrast mà không phải vẽ lại texture.",
        "examples": "Power cao làm core flare chặt hơn và giảm midtone. Power thấp hoặc SquareRoot mở rộng vùng chuyển, hữu ích cho glow mềm. Exponential tạo response tăng nhanh cho pulse hoặc energy spike. Logarithm nén dynamic range, ít dùng trực tiếp nhưng hữu ích khi cần hiểu cách giá trị sáng bị nén.",
        "danger": "Power trên giá trị âm có thể tạo kết quả không mong muốn. Power quá cao làm mask mất midtone và khiến effect chỉ còn một đốm nhỏ. Exponential dễ sinh giá trị quá lớn, nên luôn kiểm tra emissive trước và sau scale.",
        "exercise": "Tạo một flare từ cùng radial mask: core dùng Power 6, glow dùng Power 0.5, edge dùng SmoothStep. Tắt bloom và so sánh xem shape có còn đọc được hay chỉ đẹp nhờ post process.",
    },
    "Frac, Floor, Ceil, Round, Fmod": {
        "role": "Nhóm này biến giá trị liên tục thành vòng lặp hoặc bậc rời rạc. Nó là nền cho phase animation, scanline, frame index, glitch và các pattern có nhịp rõ.",
        "examples": "Frac(Time * Speed) tạo phase 0..1 lặp lại. Floor chọn ô atlas hoặc frame index. Ceil và Round tạo cảm giác bật theo nấc. Fmod giữ giá trị trong một chu kỳ tùy ý, hữu ích khi chu kỳ không phải 1 hoặc cần đồng bộ với gameplay time.",
        "danger": "Các node rời rạc dễ tạo giật nếu áp lên motion cần mượt. Frac luôn có seam tại điểm reset từ 1 về 0; nếu seam nằm trên polar UV hoặc texture không tile, người xem sẽ thấy đường cắt. Hãy đặt seam ở vùng ít nhìn hoặc che bằng mask.",
        "exercise": "Tạo một shockwave lặp bằng Frac(Time * Speed). Dùng SmoothStep quanh phase để tạo ring chạy ra ngoài, sau đó cố tình xoay seam để thấy vị trí reset ảnh hưởng tới visual ra sao.",
    },
    "Sine, Cosine, Tangent và LinearSine": {
        "role": "Nhóm lượng giác tạo chuyển động tuần hoàn: pulse, wobble, shimmer, wave và rotation. Giá trị chính của nó không nằm ở công thức, mà ở khả năng tạo nhịp có kiểm soát theo time hoặc phase.",
        "examples": "Sine tạo flicker hoặc wave theo thời gian. Cosine kết hợp với Sine để tạo vector xoay hoặc offset tròn. LinearSine trong Unreal tiện cho oscillation 0..1. Tangent ít dùng hơn trong VFX phổ thông vì dễ tăng rất mạnh gần điểm cực.",
        "danger": "Sine dùng trực tiếp thường quá sạch và dễ lộ nhịp máy. Flicker đẹp cần phase offset, noise nhẹ hoặc curve shaping để tránh cảm giác đều. Tangent cần giới hạn range rõ ràng, nếu không có thể tạo spike khó debug.",
        "exercise": "Tạo emissive pulse bằng Sine, sau đó thêm noise nhỏ vào phase hoặc amplitude. So sánh bản nhịp đều với bản có variation để thấy khác biệt giữa motion máy móc và motion có thiết kế.",
    },
    "Arctangent2, Radians, Degrees và Pi": {
        "role": "Nhóm này đưa graph vào không gian góc. Khi cần portal, vortex, radial wipe hoặc pattern xoay quanh tâm, bạn đang cần angle chứ không chỉ UV.x và UV.y. Đây là nhóm nên đọc cùng Polar UV: radius cho biết điểm cách tâm bao xa, angle cho biết điểm nằm ở hướng nào.",
        "examples": "Arctangent2 lấy angle từ vector 2D và là nền cho Polar UV. Pi và TwoPi giúp normalize góc về 0..1 hoặc đổi giữa vòng tròn và phase. Degrees/Radians hữu ích khi parameter cho artist nhập theo độ nhưng công thức chạy bằng radian. Trong production, nhóm này thường nằm trước scroll, rotate, spiral, radial dissolve hoặc mask theo cung tròn.",
        "danger": "Atan2 luôn có seam angle. Nếu texture không tile hoặc seam nằm ngay vùng focal, portal sẽ có đường cắt rõ. Nhầm Degrees với Radians làm rotation sai tốc độ hoặc sai góc, thường nhìn như graph đúng nhưng art control không ăn.",
        "exercise": "Output angle từ Atan2 ra grayscale, tìm seam, rồi xoay seam bằng offset. Sau đó đặt seam ra sau vùng bị mask để thấy vì sao Polar UV cần được debug riêng.",
    },
    "DotProduct, CrossProduct và Normalize": {
        "role": "Đây là nhóm đọc hướng. Dot biến quan hệ giữa hai vector thành một số dễ dùng cho mask; Cross tạo hướng vuông góc; Normalize làm phép so hướng không bị lẫn với độ dài vector.",
        "examples": "Dot giữa Normal và ViewDir tạo Fresnel hoặc rim response. Dot giữa Velocity và Direction đọc alignment của motion. Cross giữa BeamDir và Up tạo side vector để dựng slash, beam hoặc ribbon. Normalize nên xuất hiện trước các phép so hướng quan trọng.",
        "danger": "Dot không normalize sẽ trộn góc với độ dài, khiến mask thay đổi theo scale ngoài ý muốn. Cross gần zero khi hai vector gần song song. Normalize vector zero có thể gây bất ổn, nên cần fallback hoặc clamp độ dài trong các graph nhạy.",
        "exercise": "Dựng Fresnel thủ công bằng Dot và OneMinus. Sau đó cố tình đổi NormalWS sang một space khác để thấy lỗi space làm rim response sai rõ thế nào.",
    },
    "Distance, Length và VectorToRadialValue": {
        "role": "Nhóm này biến vị trí thành khoảng cách. Nó là nền cho circle, ring, shockwave, radial fade, proximity mask và mọi hiệu ứng cần biết pixel đang cách một điểm bao xa.",
        "examples": "Length(UV - 0.5) tạo radial mask trong UV space. Distance giữa WorldPos và ImpactPos tạo hit ripple hoặc vùng ảnh hưởng quanh điểm va chạm. VectorToRadialValue hỗ trợ workflow radial khi cần chuyển vector 2D thành thông tin radius/angle.",
        "danger": "Distance luôn phụ thuộc space. World distance, object distance và UV distance không cùng scale; nếu mesh scale hoặc unit không rõ, effect có thể lệch giữa asset. Trước khi tune radius, hãy xác nhận space và đơn vị.",
        "exercise": "Tạo một hit ripple từ World Position và Impact Position, rồi tạo cùng logic trong UV. So sánh hai bản trên mesh scale khác nhau để thấy khi nào nên dùng object/world/UV space.",
    },
    "Lerp, InverseLerp và RemapValueRange": {
        "role": "Nhóm này chuyển đổi range và pha trộn trạng thái. Lerp không chỉ là blend màu; nó là cách graph chọn giữa hai đáp án theo một alpha có thể debug được. Khi một value đến từ distance, speed, age hoặc noise, thường cần remap trước rồi mới dùng làm control.",
        "examples": "Lerp pha color, distortion, opacity hoặc hai normal/detail layer. InverseLerp biến một range bất kỳ thành 0..1 để dùng như mask. RemapValueRange đưa dữ liệu từ space đo lường sang space art direction, ví dụ distance thành alpha hoặc speed thành intensity. Trong VFX, đây là cầu nối giữa dữ liệu vật lý và cảm giác visual: cùng một speed có thể map thành glow nhẹ, shock mạnh hoặc fade nhanh tùy art direction.",
        "danger": "Alpha của Lerp cần được kiểm tra như một mask riêng. Nếu alpha không clamp, kết quả có thể vượt ngoài A/B. Remap sai input range làm control cảm giác không tuyến tính hoặc khiến parameter gần như không có tác dụng.",
        "exercise": "Lấy distance 0..500, dùng InverseLerp đưa về 0..1, sau đó Lerp giữa glow yếu và glow mạnh. Thay đổi input range để thấy art control nhạy hay lì như thế nào.",
    },
    "ComponentMask, AppendVector và AddComponents": {
        "role": "Nhóm này quản lý cách dữ liệu được đóng gói và tháo gỡ. Trong production, rất nhiều texture không còn là hình ảnh trực tiếp mà là container dữ liệu: R có thể là dissolve, G là edge, B là noise phụ, A là opacity hoặc height.",
        "examples": "ComponentMask lấy đúng channel cần đọc. AppendVector ghép scalar thành UV, float3 hoặc float4. AddComponents cộng nhiều channel thành một scalar debug, hữu ích khi cần xem tổng năng lượng hoặc kiểm tra texture packed có channel nào sai.",
        "danger": "Lỗi thường gặp là đọc nhầm channel hoặc quên sRGB/linear khi texture được dùng như data. Một mask packed sai compression có thể nhìn đúng trong preview nhưng ra giá trị sai khi sample. Nếu một effect bỗng mất alpha, kiểm tra channel trước khi sửa logic.",
        "exercise": "Tạo một texture packed giả bằng bốn mask grayscale. Dùng ComponentMask để đọc từng kênh, AppendVector để dựng UV debug, rồi xuất từng channel ra preview trước khi dùng trong final material.",
    },
    "If, Switch và logic chọn nhánh": {
        "role": "Nhóm này biểu diễn lựa chọn. Nó hữu ích khi cần variant hoặc điều kiện rõ ràng, nhưng trong shader realtime, lựa chọn không phải lúc nào cũng rẻ như trong code CPU. Hãy đọc nó như quyết định thiết kế pipeline: có thật sự cần nhánh, hay chỉ cần một mask liên tục có thể tune được.",
        "examples": "If có thể chọn output theo threshold hoặc mode, nhưng nhiều trường hợp Step, SmoothStep hoặc Lerp cho graph ổn định hơn. Switch phù hợp cho static option, quality tier hoặc material variant khi quyết định được compile-time. Ví dụ: low quality tắt distortion bằng static switch; dissolve edge thì thường dùng SmoothStep/Lerp để giữ edge mềm và dễ debug.",
        "danger": "Lạm dụng If/Switch làm graph khó đọc và có thể tạo nhiều permutation. Branch động trong pixel shader đôi khi không tiết kiệm như kỳ vọng, đặc biệt khi các pixel trong cùng vùng đi theo nhiều nhánh khác nhau.",
        "exercise": "Dựng cùng một dissolve bằng If và bằng SmoothStep/Lerp. So sánh debug view, độ mượt edge và khả năng tune parameter để thấy khi nào nên dùng logic branch, khi nào nên dùng mask liên tục.",
    },
    "Transform3x3Matrix và TransformToClipSpace": {
        "role": "Nhóm này xử lý đổi hệ tọa độ. Nó xuất hiện khi graph cần hiểu dữ liệu đang ở local, world, tangent, view hay clip space thay vì chỉ thao tác trên UV. Đây là phần dễ gây lỗi âm thầm vì con số vẫn hợp lệ, nhưng ý nghĩa không còn đúng context.",
        "examples": "Transform3x3Matrix dùng cho vector basis, tangent/world conversion hoặc custom orientation. TransformToClipSpace đưa vị trí tới không gian projection, hữu ích cho screen-space effect, decal/projection logic hoặc debug vị trí theo camera. Trong VFX, space transform thường quyết định effect bám mesh, bám world, bám camera hay bám surface.",
        "danger": "Position và direction không transform giống nhau. Nếu dùng matrix sai loại, direction có thể bị translation hoặc normal bị scale làm sai. Sai space thường không crash; nó chỉ làm effect lệch, trôi theo camera hoặc khác giữa mesh.",
        "exercise": "Tạo một mask phụ thuộc World Up, sau đó transform direction về local space để so sánh trên mesh xoay. Ghi lại trường hợp nào visual nên bám world, trường hợp nào nên bám object.",
    },
    "ProjectVectorOntoPlane và MakeVectorsOrthogonal": {
        "role": "Nhóm này giúp dựng basis và projection sạch. Nó quan trọng khi effect cần hướng phụ: slash có side axis, beam có mặt phẳng riêng, decal cần vector nằm trên surface. Khi một effect cần phản ứng theo hướng impact hoặc hướng di chuyển, nhóm này giúp biến vector thô thành coordinate system dùng được.",
        "examples": "ProjectVectorOntoPlane loại bỏ thành phần theo normal để giữ motion nằm trên mặt phẳng. MakeVectorsOrthogonal giúp dựng hai hướng vuông góc khi cần coordinate system nhỏ cho slash, ribbon hoặc procedural decal. Ví dụ, impact direction từ gameplay có thể được project lên surface normal để tạo scorch mark đúng hướng thay vì chọc xuyên qua bề mặt.",
        "danger": "Nếu vector gần song song hoặc gần zero, basis dựng ra có thể không ổn định. Projection sai normal sẽ làm decal trượt hoặc slash xoay sai. Với các node hình học kiểu này, debug bằng vector màu hoặc gizmo là cần thiết.",
        "exercise": "Dựng một slash direction trên surface: project impact direction lên plane của normal, tạo side axis bằng cross product, rồi dùng hai trục đó để sample mask slash trong local 2D space.",
    },
    "HsvToRgb, RgbToHsv và FindSaturation": {
        "role": "Nhóm này tách màu khỏi cách RGB lưu trữ. HSV giúp bạn nói về hue, saturation và value như các trục art direction, thay vì chỉ kéo ba kênh RGB độc lập. Nó đặc biệt hữu ích khi feedback nói về cảm giác màu: quá gắt, quá bẩn, thiếu nhiệt, thiếu contrast hoặc lệch palette.",
        "examples": "HsvToRgb hữu ích khi muốn hue shift theo lifetime hoặc particle data. RgbToHsv giúp lấy hue/saturation từ texture hoặc color ramp. FindSaturation tạo mask dựa trên độ rực, dùng cho debug palette hoặc giảm màu ở vùng ít quan trọng. Với stylized VFX, hue có thể đổi nhẹ theo lifetime trong khi value vẫn giữ hierarchy đọc hình.",
        "danger": "Hue shift quá mạnh dễ phá brand/style hoặc làm effect mất hierarchy. Saturation cao không đồng nghĩa đẹp hơn; nhiều effect cần core gần trắng, mid tone mang hue chính và accent giữ hue phụ. Hãy test trên background sáng, tối và cùng màu.",
        "exercise": "Tạo một color ramp bằng HSV: hue đổi nhẹ theo lifetime, saturation giảm ở outer glow, value giữ core sáng. Sau đó xuất saturation mask để kiểm tra vùng nào thật sự đang rực.",
    },
    "WrapFloat, WrapFloat2, WrapFloat3, WrapFloat4": {
        "role": "Nhóm này đưa giá trị quay vòng trong một khoảng. Nó hữu ích cho phase, UV, hue hoặc bất kỳ dữ liệu nào cần loop ổn định thay vì tăng mãi theo thời gian. Wrap không chỉ là tiện ích toán học; nó là cách đặt điểm reset có chủ đích.",
        "examples": "WrapFloat có thể giữ phase trong 0..1 hoặc 0..2π. WrapFloat2/3/4 dùng khi nhiều channel cần loop cùng quy ước, ví dụ UV offset, color phase hoặc vector data đóng gói. Trong flipbook, pulse, scanline hoặc hue cycle, wrap giúp loop đoán được và tránh giá trị time lớn làm graph khó ổn định.",
        "danger": "Wrap tạo seam tại điểm quay vòng. Nếu seam nằm ở vùng thấy rõ, visual sẽ giật hoặc cắt. Với time dài, wrap cũng giúp tránh precision drift, nhưng cần chọn chu kỳ phù hợp để motion không lặp quá lộ.",
        "exercise": "Tạo hue cycle bằng WrapFloat trên phase, sau đó đặt seam vào thời điểm alpha gần 0. So sánh với bản không wrap khi time chạy lâu để thấy lợi ích về ổn định và kiểm soát loop.",
    },
}


def _atlas_text(profile, name, key):
    override = ATLAS_TEXT_OVERRIDES.get(name, {}).get(key)
    return override or profile[key]


def _rewrite_feedback_section(chapters):
    for chapter in chapters:
        if chapter.get("title") != "Từ Physics Sang Art Direction":
            continue

        sections = []
        for title, paras in chapter.get("sections", []):
            if "Feedback" in title or "Ngọn" in title:
                sections.append((
                    "Dịch feedback thành quyết định kỹ thuật",
                    [
                        "Feedback trong review thường bắt đầu bằng cảm giác: mạnh hơn, nặng hơn, organic hơn, sạch hơn hoặc có lực hơn. Nếu phản ứng đầu tiên là tăng intensity, thêm noise hoặc kéo contrast, bạn đang sửa theo triệu chứng. Cách làm chắc hơn là tách feedback thành layer có thể kiểm tra: shape, value, color, timing, motion, depth interaction hoặc cost.",
                        "Khi lead nói effect cần mạnh hơn, hãy xác định mạnh ở đâu: core sáng hơn, silhouette rõ hơn, timing nhanh hơn, impulse nặng hơn hay contrast với background tốt hơn. Khi feedback là organic hơn, hãy kiểm tra noise frequency, field motion, phase variation, edge breakup, particle variation và decay timing. Mỗi hướng sửa phải đi kèm một debug view hoặc một parameter cụ thể.",
                        "Đó là lý do nên học nguyên lý thay vì học thuộc công thức mẫu. Công thức mẫu chỉ cho một kết quả trong một context. Nguyên lý cho khả năng phân tích context mới, chọn lane cần sửa, và giải thích vì sao thay đổi đó phục vụ readability, art direction hoặc performance.",
                    ],
                ))
            else:
                sections.append((title, paras))

        chapter["sections"] = sections
        chapter["notes"] = [(
            "Key Takeaway",
            "Nguyên lý tốt giúp biến feedback cảm tính thành quyết định kỹ thuật có thể kiểm tra: shape, timing, motion, color, depth hoặc cost.",
            False,
        )]
        chapter["case"] = [
            "Một stylized slash không cần motion blur vật lý chính xác, nhưng cần hướng lực rõ, taper sạch, timing có nhịp và value contrast đủ đọc trong gameplay.",
            "Một black hole fantasy không cần mô phỏng hấp dẫn thật, nhưng vẫn cần inward flow, lensing/distortion, core hierarchy và particle scale cue để người xem tin rằng effect có luật nội bộ.",
        ]
        chapter["code"] = [
            "Feedback signal -> technical lane:",
            "stronger  -> core value / silhouette / timing / contrast",
            "organic   -> noise frequency / phase variation / field motion",
            "heavier   -> slower decay / stronger impulse / larger scale cue",
        ]


def _fill_layout_gap_sections(chapters):
    for chapter in chapters:
        title = chapter.get("title", "")
        if title == "Blend, Emissive và Depth: Material Gặp Renderer":
            chapter.setdefault("case", []).extend([
                "Checklist nhanh cho review: tắt bloom, đổi background sáng/tối, xem alpha riêng, xem emissive trước tonemapping, kiểm DepthFade ở góc camera thấp, rồi mới quyết định tăng Intensity hay đổi Blend path.",
                "Một material blend tốt nên có fallback rõ. Nếu bản full dùng translucent, distortion và soft particles, bản rẻ hơn có thể dùng masked hoặc dithered edge, giảm card area hoặc bake motion vào flipbook. Quyết định này nên được ghi trong material note thay vì để người sau đoán.",
            ])
        if title == "WPO: Khi Material Đẩy Geometry":
            chapter.setdefault("case", []).extend([
                "Checklist nhanh cho WPO: kiểm bounds trong viewport, bật wireframe để xem mesh density, test camera gần/xa, test LOD, tắt color để chỉ đọc silhouette, và so sánh với phương án normal/distortion nếu silhouette không cần đổi.",
                "WPO đáng dùng khi nó tạo hình mà pixel shader không thay được: outline phồng, surface ripple thật, slash mesh có hướng, hoặc shield hit đẩy geometry. Nếu nó chỉ tạo nhiễu nhỏ trong vùng không ai nhìn, hãy chuyển chi phí sang texture/normal hoặc bỏ hẳn.",
            ])


TEXT_REPLACEMENTS = {
    "Đây là lý do học recipe không đủ; bạn phải học cách đọc art direction.": "Đây là lý do học công thức mẫu là chưa đủ; bạn phải học cách đọc art direction.",
    "Cũng đừng chỉ copy recipe mà không hiểu cost.": "Cũng đừng chỉ sao chép công thức mẫu mà không hiểu cost.",
    "copy recipe": "sao chép công thức mẫu",
    "học recipe": "học công thức mẫu",
}


def _clean_text(text):
    for old, new in TEXT_REPLACEMENTS.items():
        text = text.replace(old, new)
    text = re.sub(r"Defensive shader lÃ  má»™t pháº§n cá»§a seniority\.", "Defensive shader lÃ  má»™t pháº§n cá»§a tÆ° duy production.", text, flags=re.IGNORECASE)
    text = re.sub(r"khÃ´ng dÃ¹ng HLSL Ä‘á»ƒ tá» ra senior", "khÃ´ng dÃ¹ng HLSL Ä‘á»ƒ chá»©ng tá» trÃ¬nh Ä‘á»™", text, flags=re.IGNORECASE)
    text = re.sub(r"KhÃ´ng dÃ¹ng code chá»‰ Ä‘á»ƒ trÃ´ng senior hÆ¡n", "KhÃ´ng dÃ¹ng code chá»‰ Ä‘á»ƒ trÃ´ng phá»©c táº¡p hÆ¡n", text)
    text = re.sub(r"\bseniority\b", "production maturity", text, flags=re.IGNORECASE)
    text = re.sub(r"\bsenior\b", "advanced", text, flags=re.IGNORECASE)
    text = re.sub(r"\bjunior\b", "foundation", text, flags=re.IGNORECASE)
    return text


def _clean_problem_phrases(chapters):
    for chapter in chapters:
        chapter["title"] = _clean_text(chapter.get("title", ""))
        chapter["sections"] = [
            (_clean_text(title), [_clean_text(para) for para in paras])
            for title, paras in chapter.get("sections", [])
        ]
        chapter["notes"] = [
            (_clean_text(title), _clean_text(text), warning)
            for title, text, warning in chapter.get("notes", [])
        ]
        chapter["case"] = [_clean_text(text) for text in chapter.get("case", [])]
        chapter["code"] = [_clean_text(text) for text in chapter.get("code", [])]


def _compact_atlas_chapters():
    intro = _chapter_dict(
        "Cách đọc chung của Math Node Atlas",
        None,
        "",
        [
            ("Đọc theo lane thay vì đọc theo alphabet", [
                "Math node không nên được học như một danh sách tên rời rạc. Hãy hỏi node đang biến đổi lane nào: UV, Mask, Color, Opacity, Interaction hay Cost. Cùng một phép toán đặt ở lane khác nhau sẽ tạo ý nghĩa production khác nhau.",
                "Với mỗi nhóm node, ghi bốn dòng: ý nghĩa dữ liệu, VFX use, lỗi thường gặp và mini exercise. Cách ghi này biến Material Editor thành ngôn ngữ thay vì menu dài.",
                "Khi đọc một graph lạ, hãy đi từ output ngược về nguồn: output đang cần opacity, emissive, normal, world offset hay custom data; lane nào tạo ra nó; node nào chỉ chuyển range; node nào thật sự quyết định hình. Cách đọc này giúp tránh bị kẹt ở tên node mà quên mục đích visual.",
            ]),
            ("Debug bằng giá trị trung gian", [
                "Một node chỉ đáng tin khi bạn nhìn được input và output của nó. Hãy xuất raw mask, remapped value, shaped value và final result ra grayscale trước khi thêm màu hoặc bloom.",
                "Nếu visual sai, đừng sửa node cuối cùng theo phản xạ. Kiểm tra source data, range, space và channel trước.",
                "Mỗi lần thêm một node toán, nên biết nó làm ba việc gì: thay đổi range, thay đổi shape, hay đổi coordinate/space. Nếu không phân loại được, graph sẽ khó review và khó handoff. Nếu phân loại được, bạn có thể giải thích lỗi bằng ngôn ngữ production thay vì nói chung chung rằng material nhìn chưa đúng.",
            ]),
        ],
        ("Ghi nhớ", "Atlas này là phụ lục thực hành. Mục tiêu là tạo phản xạ đọc dữ liệu, không phải học thuộc tên node.", False),
        [],
        ["lane -> input range -> transform -> debug view -> final use"],
    )
    chapters = [intro]
    for name in v24.FAMILY_NAMES:
        profile = v24.FAMILY_PROFILES[name]
        chapters.append(_chapter_dict(
            f"Math Atlas: {name}",
            None,
            "",
            [
                ("Ý nghĩa trong graph", [_atlas_text(profile, name, "role")]),
                ("VFX use", [_atlas_text(profile, name, "examples")]),
                ("Failure mode", [_atlas_text(profile, name, "danger")]),
                ("Mini exercise", [_atlas_text(profile, name, "exercise")]),
            ],
            None,
            [],
            profile.get("code", [])[:3],
        ))
    return chapters


def _case_study_v27():
    return _chapter_dict(
        "Case study - đọc một Material mẫu trong Unreal",
        "v27_case_m_flare.png",
        "Diagram tái dựng từ workflow VFXSamplePack: graph đơn giản, output preview, debug mask và breakdown lane. Đây không phải screenshot Unreal thật.",
        [
            ("Case Study Diagram: M_Flare", [
                "Diagram trên tái dựng một flare additive theo các lane chính: shape, falloff, particle color, intensity, fade và output. Nó dùng để đọc contract của material, không phải để mô phỏng chính xác layout node trong Unreal.",
                "Điểm cần quan sát đầu tiên không phải node nào đẹp, mà là dữ liệu nào đang quyết định visual cuối. Nếu debug mask đã yếu, tăng emissive chỉ làm lỗi sáng hơn. Nếu Particle Color không đi qua alpha, Niagara fade sẽ không có tác dụng. Nếu DepthFade nằm quá muộn, intersection vẫn có thể gắt dù output alpha đã mềm.",
            ]),
            ("Case 1: M_Flare như một material additive có kiểm soát", [
                "M_Flare nên được đọc như một hệ thống nhỏ: Shape tạo silhouette, Falloff tạo core/glow, Particle Color điều khiển tint/fade, Intensity điều khiển năng lượng, Camera hoặc Depth Fade xử lý interaction với scene.",
                "Bài test production là tắt Bloom, đặt flare trên nền tối và nền sáng, rồi xem core có còn đọc được không. Nếu flare chỉ đẹp nhờ Bloom, vấn đề nằm ở hierarchy của mask và emissive.",
                "Một flare tốt thường có ít nhất ba lớp quyết định: core đủ chặt để giữ điểm nhấn, glow đủ rộng để tạo cảm giác năng lượng, và alpha/fade đủ sạch để không lộ card. Ba lớp này nên tune độc lập. Khi tất cả cùng nằm trong một scalar intensity, artist chỉ có một cần gạt và rất khó sửa feedback cụ thể.",
            ]),
            ("Breakdown 6 lane", [
                "Shape: radial/texture alpha quyết định silhouette. UV: TexCoord hoặc panner chỉ nên phục vụ motion rõ mục đích. Mask: alpha/falloff/edge phải xem được ở grayscale. Color/Emissive: tint, particle color và intensity tách vai trò. Interaction: camera fade/depth fade tránh giao tuyến gắt. Cost: additive card rộng có thể đắt vì overdraw.",
                "Khi review, hãy đọc từng lane theo thứ tự. Nếu silhouette chưa đúng, đừng sửa color. Nếu mask chưa có hierarchy, đừng sửa bloom. Nếu interaction với scene xấu, kiểm tra depth/camera fade trước khi thêm noise. Nếu cost cao, nhìn card area và layer count trước khi tối ưu vài phép toán nhỏ.",
            ]),
            ("Case 2: SubUV và NumberGrid", [
                "Với flipbook, renderer chọn frame, material đọc frame, Niagara điều khiển thời điểm đổi frame. Nếu ba phần này không cùng quy ước, graph đúng vẫn ra visual sai.",
                "NumberGrid giúp tách lỗi. Nếu số frame sai, kiểm SubUV grid, frame index hoặc Niagara curve. Nếu số frame đúng nhưng radial fade chạy xuyên qua nhiều ô, kiểm local UV trong frame.",
                "Một debug tốt nên đi từ renderer contract tới material contract. Đầu tiên xác nhận atlas có đúng rows/columns. Sau đó xác nhận frame order và interpolation. Cuối cùng mới kiểm local UV, mask, distortion và color. Nếu đảo thứ tự, bạn có thể sửa material rất lâu trong khi lỗi thật nằm ở frame selection.",
            ]),
        ],
        ("Ghi nhớ", "Case study này dùng diagram tái dựng để khóa tư duy đọc material. Khi có screenshot Unreal thật, có thể thay diagram nhưng giữ cùng cấu trúc breakdown.", False),
        ["Đừng sửa parameter trước khi viết được 6 lane breakdown.", "Đừng dùng texture đẹp để debug SubUV; dùng NumberGrid trước."],
        ["Material reading lanes:", "Shape / UV / Mask / Color-Emissive / Interaction / Cost"],
    )


def _case_subuv_v27():
    return _chapter_dict(
        "Case Study Diagram: SubUV / NumberGrid",
        "v27_case_subuv_numbergrid.png",
        "Diagram tái dựng từ workflow VFXSamplePack: atlas frame selection, local UV, và cách tách lỗi frame sai với lỗi local UV sai. Đây không phải screenshot Unreal thật.",
        [
            ("Frame selection và local UV là hai hợp đồng khác nhau", [
                "NumberGrid giúp nhìn ngay renderer đang chọn frame nào. Nếu số không đúng, hãy kiểm SubUV grid, frame index, frame interpolation hoặc Niagara curve trước khi chỉnh material.",
                "Khi số đúng nhưng radial fade, border hoặc distortion chạy xuyên qua nhiều ô, vấn đề thường nằm ở local UV. Material cần tạo lại không gian 0..1 bên trong frame hiện tại thay vì dùng toàn bộ atlas UV.",
                "Điểm cần nhớ là atlas UV trả lời câu hỏi sample texture ở đâu, còn local UV trả lời câu hỏi pixel này nằm ở đâu bên trong frame hiện tại. Một graph dùng nhầm hai khái niệm này vẫn có thể chạy, nhưng mask theo tâm, edge fade, rotation hoặc distortion sẽ lệch theo atlas thay vì theo frame.",
            ]),
            ("Debug order", [
                "Bước 1: thay texture đẹp bằng NumberGrid. Bước 2: xác nhận frame order. Bước 3: xuất local UV hoặc radial mask trong từng frame. Bước 4: chỉ quay lại color/distortion khi frame và local UV đã đúng.",
                "Nếu frame nhảy sai nhịp, đừng sửa material trước khi kiểm animation curve và interpolation. Nếu frame đúng nhưng shape méo, hãy xuất local UV R/G thành màu để xem nó có chạy từ 0 tới 1 trong từng ô không. Nếu local UV đúng mà final vẫn sai, lúc đó mới kiểm texture compression, packed channel, blend mode và particle alpha.",
            ]),
        ],
        ("Ghi nhớ", "SubUV không chỉ là texture animation. Nó là hợp đồng giữa Niagara renderer, atlas layout và material sampling.", False),
        ["Nếu NumberGrid sai số, sửa renderer/frame contract. Nếu NumberGrid đúng nhưng mask sai vị trí, sửa local UV."],
        ["frameUV = (atlasUV - frameOffset) / frameSize", "localMask = radial(frameUV)"],
    )


def _validation_chapter_v27():
    return _chapter_dict(
        "Production Validation Checklist",
        None,
        "",
        [
            ("Visual debug pass", [
                "Trước khi gọi một material là xong, hãy kiểm tra nó qua các debug view tối thiểu: source mask, remapped mask, shaped mask, alpha cuối, emissive không bloom, và final trong scene thật. Nếu một effect chỉ đọc được ở final beauty view, team sẽ khó sửa khi feedback tới.",
                "Debug pass cần chạy trong nhiều context: nền sáng, nền tối, nền cùng hue, camera gần, camera xa, motion nhanh và motion chậm. Một material đẹp trong preview sphere nhưng mất silhouette trong gameplay vẫn chưa đạt production quality.",
            ]),
            ("Performance and profiling pass", [
                "Tối ưu không nên bắt đầu bằng cảm giác. Hãy xác định bottleneck trước: overdraw, texture fetch, instruction count, particle count, draw call, sorting, hay CPU/Niagara update. Mỗi loại chi phí có cách sửa khác nhau; giảm node không giúp nhiều nếu vấn đề thật là alpha card quá rộng.",
                "Trong Unreal, Shader Complexity, Quad Overdraw, GPU timing, particle count và platform preview đều nên được dùng như tín hiệu chẩn đoán. Không có một heatmap nào là bản án cuối cùng, nhưng nếu nhiều tool cùng chỉ về một vấn đề, đó là hướng sửa ưu tiên.",
            ]),
            ("Handoff and maintenance pass", [
                "Một material production cần người khác mở lên vẫn hiểu được. Parameter phải có tên rõ, range hợp lý, default an toàn và grouping theo lane: Shape, UV, Mask, Color, Interaction, Cost. Nếu artist phải hỏi scalar nào điều khiển alpha hoặc edge width, graph chưa được handoff tốt.",
                "Handoff tốt cũng cần ghi chú tradeoff: layer nào có thể tắt trên low spec, texture nào là data texture, channel nào packed mask, parameter nào được Niagara sở hữu, và test nào bắt buộc trước khi ship. Đây là phần biến một graph đẹp thành asset sống được trong project.",
            ]),
        ],
        ("Key Takeaway", "Production validation là cầu nối giữa material đẹp và material dùng được: debug được, đo được, handoff được, và giảm chất lượng có kiểm soát.", False),
        [
            "Nếu một effect chỉ đẹp trong một camera, nó chưa đủ. Nếu nó không có debug path, nó khó bảo trì. Nếu nó không có quality tier, nó khó sống trên nhiều platform.",
            "Một review tốt nên kết thúc bằng quyết định rõ: sửa visual lane nào, đo lại cost nào, và test lại trong context nào.",
        ],
        [
            "Ship checklist:",
            "debug views -> background tests -> bloom off -> motion test",
            "shader complexity -> overdraw -> texture fetch -> particle count",
            "parameter names -> safe ranges -> quality tier -> owner notes",
        ],
    )


def _defensive_shader_chapter_v27():
    return _chapter_dict(
        "Defensive Shader: Viết Graph Khó Hỏng",
        None,
        "",
        [
            ("Division by Zero", [
                "Divide là node rất hữu ích nhưng cũng là nguồn lỗi âm thầm. Khi mẫu số có thể về 0 hoặc rất nhỏ, output sẽ bùng lên thành giá trị cực lớn, Inf hoặc NaN. Trong VFX, lỗi này thường nhìn như flash trắng, mask cháy, distortion văng hoặc material biến mất.",
                "Quy tắc production: không divide trực tiếp bằng dữ liệu chưa bảo vệ. Dùng max(abs(denom), epsilon) hoặc clamp range trước khi divide. Epsilon phải hợp scale của bài toán; quá lớn thì làm mất độ nhạy, quá nhỏ thì không bảo vệ đủ.",
            ]),
            ("SafeNormalize", [
                "Normalize vector zero là lỗi phổ biến trong direction math. Khi velocity, impact direction hoặc cross product gần 0, Normalize có thể tạo output không ổn định. SafeNormalize kiểm tra length trước, rồi fallback về một hướng mặc định.",
                "Nếu direction điều khiển slash, beam, WPO hoặc projection, hãy debug vector bằng màu trước khi dùng nó để sample mask. Một direction sai space hoặc zero-length thường nhìn như effect xoay lạ, nhảy frame hoặc biến mất ở vài góc camera.",
            ]),
            ("Safe Power and Log", [
                "Power và Log cần input hợp lệ. Power trên giá trị âm với exponent không nguyên có thể tạo NaN. Log của 0 hoặc số âm không có nghĩa trong shader thông thường. Nếu mask đến từ noise/remap có thể vượt range, hãy saturate hoặc max trước khi đưa vào Power/Log.",
                "Không phải lúc nào clamp cũng đặt ở cuối. Với defensive shader, có những clamp đặt trước phép toán nguy hiểm để bảo vệ graph. Sau đó vẫn nên debug output trung gian để chắc rằng clamp không giết mất art control.",
            ]),
        ],
        ("Foundation Tip", "Nếu một material đột nhiên flash trắng hoặc mất toàn bộ opacity, kiểm Divide, Normalize, Power và Log trước khi sửa color.", False),
        [
            "Mini exercise: cố tình đặt denominator về 0 trong một radial fade, xem output bị cháy ra sao, rồi sửa bằng safe divide.",
            "Mini exercise: cho velocity bằng zero vào Normalize, sau đó thay bằng SafeNormalize với fallback WorldUp.",
            "Mini exercise: đưa noise remap thành giá trị âm rồi dùng Power. Quan sát lỗi, sau đó thêm max(x, 0) hoặc saturate trước Power.",
        ],
        [
            "safeDenom = max(abs(denom), 1e-4)",
            "safeDiv   = numerator / safeDenom",
            "safeNorm  = len(v) > 1e-4 ? v / len(v) : float3(0,0,1)",
            "safePow   = pow(max(x, 0), exponent)",
            "safeLog   = log(max(x, 1e-4))",
        ],
    )


def _profiling_chapter_v27():
    return _chapter_dict(
        "Production Profiling: Từ Heatmap đến GPU Capture",
        None,
        "",
        [
            ("Shader Complexity không Đủ Một Mình", [
                "Shader Complexity và Quad Overdraw là điểm bắt đầu tốt, nhưng không phải toàn bộ sự thật. Một material instruction thấp vẫn có thể đắt nếu alpha card phủ diện tích lớn. Một shader instruction cao có thể ổn nếu nó chạy trên vùng pixel nhỏ và ít layer.",
                "Profiling tốt cần hỏi đúng bottleneck: pixel cost, vertex cost, texture bandwidth, particle update, draw call, sorting hay post-process. Mỗi bottleneck có cách sửa khác nhau.",
            ]),
            ("RenderDoc và Frame Capture", [
                "RenderDoc hoặc công cụ frame capture tương đương giúp bạn nhìn draw call, render target, texture sample, pass order và state thật. Đây là nơi kiểm chứng material có thật sự là thủ phạm hay chỉ là thứ dễ bị đổ lỗi.",
                "Khi capture một VFX, hãy cô lập frame có effect nặng nhất, xem draw call liên quan, kiểm render target size, blend state, texture count và vùng screen coverage. Nếu effect chỉ tốn nhiều khi phủ nửa màn hình, vấn đề có thể là area/overdraw chứ không phải công thức.",
            ]),
            ("Mobile Tile-Based GPU", [
                "Nhiều GPU mobile dùng tile-based rendering. Chúng chia màn hình thành tile nhỏ để giảm bandwidth, nhưng translucent overdraw, depth complexity và render target switching vẫn có thể rất nhạy. Một effect ổn trên PC không tự động ổn trên mobile.",
                "Trên mobile, hãy ưu tiên card area nhỏ, layer count thấp, texture fetch ít, atlas hợp lý và fallback rõ. Bake motion vào flipbook khi runtime procedural không đem lại khác biệt đủ lớn.",
            ]),
        ],
        ("Expert Deep-Dive", "Đừng tối ưu node trước khi biết bottleneck. RenderDoc/frame capture trả lời câu hỏi pass nào, draw call nào, texture nào và vùng pixel nào đang trả giá.", True),
        [
            "Profiling checklist: isolate effect -> disable bloom -> view shader complexity -> view quad overdraw -> capture frame -> compare platform preview.",
            "Nếu PC ổn nhưng mobile tụt frame, kiểm alpha area, texture bandwidth, render target format và số translucent layer trước.",
        ],
        [
            "Cost question:",
            "pixel area? layer count? texture fetch? instruction? particle count?",
            "capture frame -> find pass -> inspect draw -> decide fix",
        ],
    )


def _failure_atlas_chapter_v27():
    return _chapter_dict(
        "Failure Atlas: Triệu Chứng, Nguyên Nhân và Debug View",
        None,
        "",
        [
            ("Artifact Map", [
                "Failure Atlas là bảng chẩn đoán nhanh. Mục tiêu không phải đoán đúng ngay, mà là chọn debug view đầu tiên hợp lý để thu hẹp nguyên nhân.",
                "Hãy đọc lỗi theo triệu chứng visual: nhấp nháy, cháy trắng, mất alpha, seam, banding, sorting, frame sai, WPO culling hoặc edge răng cưa. Mỗi triệu chứng nên dẫn tới một debug view cụ thể.",
            ]),
            ("Common Failures", [
                "Flash trắng thường liên quan Divide, Power, HDR emissive hoặc exposure. Edge răng cưa thường liên quan Step quá cứng, thiếu fwidth hoặc texture mip. SubUV frame sai thường nằm ở renderer/frame index, còn mask sai trong frame thường nằm ở local UV.",
                "Sorting sai thường không sửa bằng texture đẹp hơn. Nó cần kiểm blend mode, depth write/test, particle sort key, pivot, camera angle và layer order. WPO biến mất thường cần kiểm bounds trước khi kiểm opacity.",
            ]),
            ("Break It on Purpose", [
                "Một cách học nhanh là cố tình phá shader. Đưa dữ liệu sai space, tắt saturate, đặt denominator về 0, tăng power quá cao, bỏ local UV hoặc đưa vector zero vào Normalize. Sau đó ghi lại triệu chứng và debug view tương ứng.",
                "Người học xây được phản xạ nhận lỗi. Người review có checklist để kiểm tra. Team có ngôn ngữ chung để nói 'đây là lỗi space' thay vì 'nhìn nó kỳ kỳ'.",
            ]),
        ],
        ("Foundation Tip", "Mỗi lần gặp lỗi visual, chụp một ảnh final và một ảnh debug grayscale. Không có debug view thì rất dễ sửa nhầm layer.", False),
        [
            "Mini exercise: đổi World Position thành Object Position trong hit ripple và mô tả triệu chứng.",
            "Mini exercise: bỏ fwidth ở một ring nhỏ rồi zoom camera xa để thấy aliasing.",
            "Mini exercise: đảo rows/columns của SubUV atlas rồi dùng NumberGrid để bắt lỗi.",
        ],
        [
            "Failure Atlas rows:",
            "symptom -> likely cause -> first debug view -> first fix",
            "white flash -> unsafe math/HDR -> clamp emissive/debug mask -> safe math",
            "wrong SubUV -> frame contract/local UV -> NumberGrid/local UV -> fix renderer/material",
        ],
    )


def _patch_chapters():
    v26.polish_chapters()
    image_map = {
        "Số, range": ("v27_curve_01.png", "Curve 0..1, remap và shaping cho thấy vì sao range sạch giúp material dễ tune."),
        "Shaping Function": ("v27_curve_01.png", "Cùng mask 0..1, curve khác nhau tạo cảm giác core, glow và fade khác nhau."),
        "Distance, SDF": ("v27_sdf_ring.png", "SDF ring tách rõ radius, width, softness và edge band."),
        "Polar UV": ("v27_polar_uv_unwrap.png", "Polar UV đổi UV vuông thành angle/radius; seam cần được quản lý có chủ đích."),
        "Production wisdom": ("v27_overdraw_layers.png", "Overdraw là câu chuyện diện tích pixel và số lớp chồng nhau, không chỉ instruction count."),
        "Optimization Cookbook": ("v27_overdraw_layers.png", "Tối ưu bắt đầu từ đo loại chi phí: area, layer count, texture fetch hay instruction."),
        "Tọa độ UV": ("v27_subuv_local_uv.png", "Atlas UV và local frame UV là hai không gian khác nhau trong SubUV workflow."),
        "Blend, emissive": ("v27_blend_depth_sorting.png", "Blend mode, depth và sorting quyết định cách material gặp renderer."),
        "Ánh sáng vật lý": ("v27_light_energy_curves.png", "Đồ thị falloff và hierarchy giúp đọc emissive mà không phụ thuộc ảnh flare cinematic."),
        "Ngôn ngữ Shader": ("v27_shader_contract_blocks.png", "Một shader function tốt có input, transform, output và debug path rõ ràng."),
        "Niagara và Material": ("v27_niagara_material_contract_2d.png", "Niagara gửi attribute; Material nhận từng channel như một hợp đồng dữ liệu rõ nghĩa."),
        "POM, Height": ("v27_pom_height_cross_section.png", "POM nhìn theo mặt cắt: ray step qua height field nhưng silhouette thật không đổi."),
        "Production stories": ("v27_production_strategy_matrix.png", "Ma trận production cho thấy style/platform quyết định layer nào giữ, bake hoặc bỏ."),
    }

    patched = []
    for chapter in base.CHAPTERS:
        title = chapter.get("title", "")
        if title.startswith("Atlas node Unreal -"):
            continue
        if title.startswith("Unreal Math Node Atlas"):
            continue
        if title.startswith("Case study - đọc một Material mẫu trong Unreal"):
            patched.append(_case_study_v27())
            patched.append(_case_subuv_v27())
            continue
        new_chapter = dict(chapter)
        for prefix, image in image_map.items():
            if title.startswith(prefix):
                new_chapter["image"] = image
                break
        patched.append(new_chapter)
        title_key = title.casefold()
        if title_key.startswith("ổn định số"):
            patched.append(_defensive_shader_chapter_v27())
        if title_key.startswith("cơ bắp của realtime"):
            patched.append(_profiling_chapter_v27())
        if (
            title_key.startswith("practice roadmap")
            or title_key.startswith("lộ trình luyện tập")
        ):
            patched.append(_validation_chapter_v27())
            patched.append(_failure_atlas_chapter_v27())
    base.CHAPTERS = patched + _compact_atlas_chapters()
    _polish_titles(base.CHAPTERS)
    _rewrite_feedback_section(base.CHAPTERS)
    _clean_problem_phrases(base.CHAPTERS)
    _fill_layout_gap_sections(base.CHAPTERS)


def export_pdf_with_word(docx_path, pdf_path):
    docx = str(Path(docx_path).resolve())
    pdf = str(Path(pdf_path).resolve())
    ps = f"""
$ErrorActionPreference = 'Stop'
$word = New-Object -ComObject Word.Application
$word.Visible = $false
$doc = $null
try {{
    $doc = $word.Documents.Open('{docx}')
    $doc.Fields.Update() | Out-Null
    foreach ($toc in $doc.TablesOfContents) {{ $toc.Update() | Out-Null }}
    try {{ $doc.BuiltInDocumentProperties('Title').Value = '{TITLE}' }} catch {{}}
    try {{ $doc.BuiltInDocumentProperties('Subject').Value = '{SUBTITLE}' }} catch {{}}
    try {{ $doc.BuiltInDocumentProperties('Author').Value = 'ForgePilot / VFXSamplePack' }} catch {{}}
    try {{ $doc.BuiltInDocumentProperties('Keywords').Value = 'Material VFX, Realtime VFX, Unreal Material, Niagara, Shader, HLSL, SubUV, SDF, Overdraw' }} catch {{}}
    $doc.Save()
    $doc.ExportAsFixedFormat('{pdf}', 17, $false, 0, 0, 1, 1, 0, $true, $true, 1, $true, $true, $false)
}} finally {{
    if ($doc -ne $null) {{ $doc.Close($false) | Out-Null }}
    $word.Quit() | Out-Null
}}
"""
    subprocess.run(
        ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", ps],
        check=True,
    )


def build_v27():
    v26.generate_v26_diagrams()
    generate_v27_diagrams()
    _patch_chapters()

    base.cover = cover_v27
    base.note = v26.note_v26
    base.h1 = v26.h1_v26
    base.figure = v26.figure_v26
    base.code = code_box_v27
    base.chapter = chapter_v27
    base.closing = closing_v27

    doc = base.setup()
    _set_metadata(doc)
    base.cover(doc)
    _apply_book_grid(doc)
    _set_footer_with_page_numbers(doc)
    _add_toc(doc)
    base.preface(doc)
    learning_structure_v27(doc)
    technical_vocabulary_v27(doc)
    rosetta_stone_v27(doc)
    for i, chapter in enumerate(base.CHAPTERS, 1):
        if i in CHAPTER_BLOCKS:
            chapter_block_divider_v27(doc, *CHAPTER_BLOCKS[i])
        base.chapter(doc, i, chapter)
    interactive_companion_v27(doc)
    base.closing(doc)
    references_v27(doc)
    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_DOCX)
    export_pdf_with_word(OUT_DOCX, OUT_PDF)
    print(OUT_DOCX)
    print(OUT_PDF)


if __name__ == "__main__":
    build_v27()
