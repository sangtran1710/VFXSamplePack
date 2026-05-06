from pathlib import Path
import importlib.util
import math

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
BASE_SCRIPT = ROOT / "tools" / "build_material_vfx_v25_math_physics_core.py"

spec = importlib.util.spec_from_file_location("material_vfx_v25", BASE_SCRIPT)
v25 = importlib.util.module_from_spec(spec)
spec.loader.exec_module(v25)

base = v25.base
OUT = Path.home() / "Downloads" / "Giao_Trinh_Material_VFX_Realtime_V26_PolishedBook.docx"
base.OUT = OUT

IMG_DIR = ROOT / "Saved" / "material_vfx_v5_images"


def _font(size=30, bold=False):
    candidates = [
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
        "C:/Windows/Fonts/segoeuib.ttf" if bold else "C:/Windows/Fonts/segoeui.ttf",
    ]
    for path in candidates:
        if Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def _canvas(title, subtitle=None, size=(1700, 820)):
    im = Image.new("RGB", size, "#F8F7F3")
    d = ImageDraw.Draw(im)
    d.rectangle([0, 0, size[0], 86], fill="#173B59")
    d.text((58, 24), title, fill="white", font=_font(30, True))
    if subtitle:
        d.text((58, 102), subtitle, fill="#4C5661", font=_font(22, False))
    return im, d


def _save(im, name):
    IMG_DIR.mkdir(parents=True, exist_ok=True)
    path = IMG_DIR / name
    im.save(path, quality=94)


def _line(d, pts, fill="#1F5FBF", width=5):
    d.line(pts, fill=fill, width=width, joint="curve")


def generate_v26_diagrams():
    # Range and remap.
    im, d = _canvas("Range / Remap", "Đưa dữ liệu thô về 0..1 trước khi shaping")
    x0, y0, x1, y1 = 160, 400, 1540, 400
    _line(d, [(x0, y0), (x1, y1)], "#1E293B", 4)
    for t, label in [(0, "inMin"), (0.35, "x"), (1, "inMax")]:
        x = x0 + (x1 - x0) * t
        d.ellipse([x-10, y0-10, x+10, y0+10], fill="#E66A2C")
        d.text((x-38, y0+28), label, fill="#1E293B", font=_font(22, True))
    d.rounded_rectangle([430, 545, 1270, 660], radius=18, outline="#CBD5E1", width=3, fill="#FFFFFF")
    d.text((470, 585), "t = saturate((x - inMin) / (inMax - inMin))", fill="#111827", font=_font(30, True))
    _save(im, "v26_range_remap.png")

    # Derivative / edge.
    im, d = _canvas("Derivative / Edge", "Edge là vùng giá trị đổi nhanh quanh threshold")
    for i in range(1200):
        x = 230 + i
        v = 1 / (1 + math.exp(-(i - 600) / 55))
        y = 640 - int(v * 390)
        if i:
            _line(d, [(px, py), (x, y)], "#2563EB", 4)
        px, py = x, y
    d.line([230, 445, 1430, 445], fill="#EF4444", width=3)
    d.text((1450, 430), "threshold", fill="#EF4444", font=_font(24, True))
    d.double_arrow = None
    d.rectangle([790, 235, 930, 650], outline="#F59E0B", width=4)
    d.text((745, 675), "vùng đạo hàm cao = edge", fill="#92400E", font=_font(24, True))
    _save(im, "v26_derivative_edge.png")

    # Wave.
    im, d = _canvas("Wave / Phase", "Motion procedural thường là quản lý phase")
    origin = (150, 430)
    d.line([120, 430, 1580, 430], fill="#334155", width=3)
    pts = []
    for i in range(1400):
        x = 150 + i
        y = 430 - int(math.sin(i / 70) * 160)
        pts.append((x, y))
    _line(d, pts, "#0EA5E9", 5)
    for x, label in [(320, "phase"), (760, "frequency"), (1190, "wavelength")]:
        d.line([x, 230, x, 640], fill="#CBD5E1", width=2)
        d.text((x-65, 665), label, fill="#0F172A", font=_font(23, True))
    _save(im, "v26_wave_phase.png")

    # Vector field.
    im, d = _canvas("Vector Field", "Radial + Tangential = cảm giác vortex")
    cx, cy = 850, 460
    for yy in range(220, 700, 80):
        for xx in range(260, 1460, 100):
            dx, dy = cx - xx, cy - yy
            length = max((dx*dx + dy*dy) ** 0.5, 1)
            rx, ry = dx / length, dy / length
            tx, ty = -ry, rx
            vx, vy = rx * 32 + tx * 22, ry * 32 + ty * 22
            d.line([xx, yy, xx + vx, yy + vy], fill="#2563EB", width=3)
            d.polygon([(xx+vx, yy+vy), (xx+vx-tx*7-rx*10, yy+vy-ty*7-ry*10), (xx+vx+tx*7-rx*10, yy+vy+ty*7-ry*10)], fill="#2563EB")
    d.ellipse([cx-55, cy-55, cx+55, cy+55], outline="#E66A2C", width=6)
    d.text((735, 540), "inward + swirl", fill="#111827", font=_font(28, True))
    _save(im, "v26_vector_field.png")

    # Sampling aliasing.
    im, d = _canvas("Sampling / Aliasing", "Tín hiệu nhỏ hơn pixel sẽ rung hoặc moiré")
    for i in range(36):
        x = 160 + i * 18
        color = "#111827" if i % 2 == 0 else "#FFFFFF"
        d.rectangle([x, 250, x+18, 620], fill=color)
    d.text((190, 650), "tần số quá cao", fill="#111827", font=_font(25, True))
    for i in range(12):
        x = 880 + i * 48
        color = "#111827" if i % 2 == 0 else "#FFFFFF"
        d.rectangle([x, 250, x+48, 620], fill=color)
    d.text((935, 650), "tần số hợp pixel", fill="#111827", font=_font(25, True))
    _save(im, "v26_sampling_aliasing.png")

    # Color hierarchy.
    im, d = _canvas("Value / Color Hierarchy", "Value đọc trước hue; màu chỉ làm ý đồ rõ hơn")
    for i, r in enumerate([190, 135, 85, 40]):
        shade = int(30 + i * 55)
        d.ellipse([350-r, 420-r, 350+r, 420+r], fill=(shade, shade, shade))
    d.text((230, 650), "grayscale đọc shape", fill="#111827", font=_font(25, True))
    for i, r in enumerate([190, 135, 85, 40]):
        col = [(30,70,140), (40,120,220), (220,110,32), (255,240,210)][i]
        d.ellipse([1120-r, 420-r, 1120+r, 420+r], fill=col)
    d.text((1000, 650), "color làm hierarchy", fill="#111827", font=_font(25, True))
    _save(im, "v26_color_hierarchy.png")

    # Defensive shader.
    im, d = _canvas("Defensive Shader", "Bảo vệ boundary trước khi dữ liệu xấu lan ra toàn graph")
    boxes = [("input", 190), ("remap", 480), ("safe clamp", 770), ("shape", 1060), ("output", 1350)]
    for label, x in boxes:
        d.rounded_rectangle([x, 350, x+190, 455], radius=18, fill="#FFFFFF", outline="#CBD5E1", width=3)
        d.text((x+30, 386), label, fill="#111827", font=_font(25, True))
    for (_, x0), (_, x1) in zip(boxes, boxes[1:]):
        d.line([x0+190, 402, x1, 402], fill="#E66A2C", width=4)
        d.polygon([(x1,402), (x1-16,392), (x1-16,412)], fill="#E66A2C")
    d.text((430, 560), "protect: divide, normalize, power, remap", fill="#334155", font=_font(30, True))
    _save(im, "v26_defensive_shader.png")


def _set_paragraph_left_border(paragraph, color="C98E54", size="12"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    left = p_bdr.find(qn("w:left"))
    if left is None:
        left = OxmlElement("w:left")
        p_bdr.append(left)
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), size)
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), color)


def note_v26(doc, title, text, warning=False):
    fill = "FFF4E6" if warning else "EFF6FF"
    color = base.ORANGE if warning else base.BLUE
    para = base.p(doc, "", size=9.8, font=base.BODY, before=8, after=8, color=base.INK)
    para.paragraph_format.left_indent = Cm(0.45)
    para.paragraph_format.right_indent = Cm(0.45)
    para.paragraph_format.line_spacing = 1.18
    base.shade(para, fill)
    _set_paragraph_left_border(para, "C98E54" if warning else "4F83CC", "18")
    r1 = para.add_run(f"{title}. ")
    base.set_font(r1, size=9.8, name=base.SANS, bold=True, color=color)
    r2 = para.add_run(text)
    base.set_font(r2, size=9.8, name=base.BODY, color=base.INK)


MAJOR_OPENERS = (
    "Lời nói đầu",
    "Digital Eyes",
    "Số, range",
    "Unreal Math Node Atlas",
    "Cơ bắp của realtime",
    "Case study",
    "Kết luận",
    "Nguồn tham khảo",
)


def h1_v26(doc, title, number=None):
    is_major = title.startswith(MAJOR_OPENERS)
    if not is_major:
        # Most entries are lessons, not full chapter openers. Let them flow
        # to avoid half-empty tail pages and make the document feel like a
        # technical book instead of one slide per topic.
        head = base.p(doc, title, size=13.5, font=base.HEAD, bold=True, color=base.INK, before=18, after=6, style="Heading 2")
        base.keep_next(head)
        base.border_bottom(head, color="E5E7EB", size="3")
        return head
    doc.add_page_break()
    if number is not None:
        base.p(doc, f"{number}", size=9.2, font=base.BODY, italic=True, color=base.MUTED, before=4, after=2, align=WD_ALIGN_PARAGRAPH.RIGHT)
    head = base.p(doc, title, size=18.2, font=base.HEAD, bold=True, color=base.INK, after=10, align=None, style="Heading 1")
    head.paragraph_format.line_spacing = 1.04
    return head


def figure_v26(doc, filename, caption, width=13.0):
    path = IMG_DIR / filename
    if not path.exists():
        return
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run()
    run.add_picture(str(path), width=Cm(min(width, 9.8)))
    if caption:
        cap = base.p(doc, caption, size=8.6, font=base.BODY, italic=True, color=base.MUTED, before=1, after=7, align=WD_ALIGN_PARAGRAPH.CENTER)
        cap.paragraph_format.left_indent = Cm(1.0)
        cap.paragraph_format.right_indent = Cm(1.0)


def cover_v26(doc):
    base.p(doc, "Giáo trình", size=12, font=base.SANS, bold=True, color=base.ORANGE, before=56, after=5, align=WD_ALIGN_PARAGRAPH.CENTER)
    title = base.p(doc, "Material VFX Realtime", size=32, font=base.HEAD, bold=True, color=base.BLUE, after=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    title.paragraph_format.line_spacing = 1.0
    base.p(doc, "Toán học, vật lý trực quan và tư duy shader cho Technical VFX Artist", size=12.2, font=base.BODY, italic=True, color=base.MUTED, after=22, align=WD_ALIGN_PARAGRAPH.CENTER)
    rule = base.p(doc, "", before=0, after=18)
    base.border_bottom(rule, color="C98E54", size="10")
    base.p(doc, "Space · Vector · Wave · Field · Sampling · Light · Energy", size=14.2, font=base.SANS, bold=True, color=base.BLUE, before=8, after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    summary = base.p(
        doc,
        "Bản polish layout: giảm hình lặp, thêm diagram tối giản đúng chủ đề, tăng điểm nhấn đọc nhanh, và gom phần atlas thành phụ lục liên tục để bớt khoảng trống.",
        size=10.4,
        font=base.BODY,
        color=base.INK,
        before=16,
        after=22,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    summary.paragraph_format.left_indent = Cm(1.0)
    summary.paragraph_format.right_indent = Cm(1.0)
    base.p(doc, "Phiên bản V26 · Polished Book Layout · 2026", size=9.3, font=base.SANS, color=base.MUTED, before=8, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)


IMAGE_BY_TITLE = {
    "Số, range": ("v26_range_remap.png", "Remap đưa dữ liệu thô về 0..1 để các bước shaping phía sau dễ kiểm soát."),
    "Đạo hàm": ("v26_derivative_edge.png", "Edge là vùng dữ liệu đổi nhanh; derivative giúp kiểm soát cạnh và aliasing."),
    "Sóng": ("v26_wave_phase.png", "Wave trong shader thường là phase chạy theo thời gian hoặc khoảng cách."),
    "Trường lực": ("v26_vector_field.png", "Vector field cho motion có luật: hút vào tâm, xoáy quanh tâm hoặc kết hợp cả hai."),
    "Sampling theory": ("v26_sampling_aliasing.png", "Aliasing xuất hiện khi tín hiệu có tần số cao hơn khả năng lấy mẫu của pixel/frame."),
    "Màu sắc": ("v26_color_hierarchy.png", "Value giúp effect đọc được; hue và saturation làm nó có phong cách."),
    "Ổn định số": ("v26_defensive_shader.png", "Defensive shader đặt boundary rõ cho divide, normalize, power và remap."),
}


def polish_chapters():
    seen = set()
    for chapter in base.CHAPTERS:
        title = chapter.get("title", "")
        for prefix, image in IMAGE_BY_TITLE.items():
            if title.startswith(prefix):
                chapter["image"] = image
                break
        image = chapter.get("image")
        if image:
            filename = image[0]
            if filename in seen and not filename.startswith("v26_"):
                chapter["image"] = None
            else:
                seen.add(filename)


def build_v26():
    generate_v26_diagrams()
    polish_chapters()
    base.cover = cover_v26
    base.note = note_v26
    base.h1 = h1_v26
    base.figure = figure_v26
    base.build()


if __name__ == "__main__":
    build_v26()
