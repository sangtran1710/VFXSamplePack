from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path.home() / "Downloads" / "Giao_Trinh_Material_VFX_Realtime_V18_BookSample.docx"
IMAGE_DIR = Path("Saved") / "material_vfx_v5_images"

INK = RGBColor(24, 24, 24)
MUTED = RGBColor(92, 92, 92)
BLUE = RGBColor(22, 58, 88)
ORANGE = RGBColor(152, 78, 30)
BODY = "Georgia"
HEAD = "Aptos Display"
SANS = "Aptos"
CODE = "Consolas"


def set_font(run, size=10.6, name=BODY, bold=False, italic=False, color=INK):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def shade(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def border_bottom(paragraph, color="D6D6D6", size="4"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), color)


def keep_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    if p_pr.find(qn("w:keepNext")) is None:
        p_pr.append(OxmlElement("w:keepNext"))


def p(doc, text="", size=10.7, before=0, after=6.2, font=BODY, bold=False, italic=False, color=INK, align=None, style=None):
    para = doc.add_paragraph()
    if style:
        para.style = style
    para.paragraph_format.space_before = Pt(before)
    para.paragraph_format.space_after = Pt(after)
    para.paragraph_format.line_spacing = 1.33
    if align is not None:
        para.alignment = align
    if text:
        run = para.add_run(text)
        set_font(run, size=size, name=font, bold=bold, italic=italic, color=color)
    return para


def body(doc, text):
    return p(doc, text)


def h1(doc, title, number=None):
    doc.add_page_break()
    if number is not None:
        p(doc, f"Chương {number}", size=9.5, font=SANS, bold=True, color=ORANGE, before=18, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
    head = p(doc, title, size=25, font=HEAD, bold=True, color=BLUE, after=8, align=WD_ALIGN_PARAGRAPH.CENTER, style="Heading 1")
    head.paragraph_format.line_spacing = 1.05
    border_bottom(head, color="C98E54", size="8")
    return head


def h2(doc, title):
    para = p(doc, title, size=14.0, font=HEAD, bold=True, color=BLUE, before=12, after=4, style="Heading 2")
    keep_next(para)
    return para


def h3(doc, title):
    para = p(doc, title, size=10.4, font=SANS, bold=True, color=INK, before=8, after=3)
    keep_next(para)
    return para


def note(doc, title, text, warning=False):
    fill = "F5F2EC" if not warning else "F8EDE6"
    color = BLUE if not warning else ORANGE
    head = p(doc, title, size=9.3, font=SANS, bold=True, color=color, before=7, after=1.5)
    head.paragraph_format.left_indent = Cm(0.45)
    head.paragraph_format.right_indent = Cm(0.45)
    shade(head, fill)
    para = p(doc, text, size=9.6, font=BODY, before=0, after=6)
    para.paragraph_format.left_indent = Cm(0.45)
    para.paragraph_format.right_indent = Cm(0.45)
    shade(para, fill)


def code(doc, lines, caption=None):
    if caption:
        p(doc, caption, size=8.8, font=SANS, italic=True, color=MUTED, after=1.5)
    for line in lines:
        para = p(doc, line, size=8.7, font=CODE, before=0, after=0.5)
        para.paragraph_format.left_indent = Cm(0.5)
        para.paragraph_format.right_indent = Cm(0.5)
        para.paragraph_format.line_spacing = 1.05
        shade(para, "ECEFF2")


def figure(doc, filename, caption, width=13.0):
    path = IMAGE_DIR / filename
    if not path.exists():
        return
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(8)
    para.paragraph_format.space_after = Pt(2)
    para.add_run().add_picture(str(path), width=Cm(width))
    cap = p(doc, caption, size=8.5, font=SANS, italic=True, color=MUTED, after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    cap.paragraph_format.line_spacing = 1.12


def bullet(doc, text):
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Cm(0.6)
    para.paragraph_format.first_line_indent = Cm(-0.24)
    para.paragraph_format.space_after = Pt(3.2)
    para.paragraph_format.line_spacing = 1.18
    run = para.add_run("- " + text)
    set_font(run, size=10.0, name=BODY)


def setup():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.75)
    sec.bottom_margin = Cm(1.55)
    sec.left_margin = Cm(2.3)
    sec.right_margin = Cm(2.3)
    sec.header_distance = Cm(0.8)
    sec.footer_distance = Cm(0.75)
    styles = doc.styles
    styles["Normal"].font.name = BODY
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), BODY)
    styles["Normal"].font.size = Pt(10.7)
    styles["Normal"].paragraph_format.line_spacing = 1.33
    styles["Normal"].paragraph_format.space_after = Pt(6.2)
    footer = sec.footer.paragraphs[0]
    footer.text = "Giáo Trình Material VFX Realtime"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        set_font(run, size=8.0, name=SANS, color=MUTED)
    return doc


def cover(doc):
    p(doc, "Giáo Trình", size=14, font=SANS, bold=True, color=ORANGE, before=60, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
    p(doc, "Material VFX Realtime", size=31, font=HEAD, bold=True, color=BLUE, after=7, align=WD_ALIGN_PARAGRAPH.CENTER)
    p(doc, "Tư duy material, shader và production cho Technical VFX Artist", size=12.2, font=BODY, italic=True, color=MUTED, after=22, align=WD_ALIGN_PARAGRAPH.CENTER)
    figure(doc, "v14_github_pdf_cover.png", "Material VFX được đọc như một hệ thống dữ liệu: shape, mask, motion, color, interaction và cost.", width=9.6)
    p(doc, "2026", size=9.5, font=SANS, color=MUTED, before=4, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)


def preface(doc):
    h1(doc, "Lời nói đầu")
    body(doc, "Một hiệu ứng realtime tồn tại trong một khoảng thời gian rất ngắn, nhưng nó được tạo ra bởi rất nhiều quyết định. Một vệt lửa nhỏ có thể chứa texture, UV, mask, particle color, blend mode, depth, exposure, bloom và cả giới hạn của GPU. Người xem chỉ thấy nó cháy lên rồi biến mất. Người làm Technical VFX phải thấy được hệ thống phía sau khoảnh khắc đó.")
    body(doc, "Material VFX vì vậy không nên được học như một danh sách node. Một node chỉ có ý nghĩa khi nó biến đổi đúng loại dữ liệu để phục vụ đúng mục tiêu thị giác. Power có thể làm flare sắc hơn, nhưng cũng có thể làm smoke mất mềm. SmoothStep có thể tạo dissolve edge sạch, nhưng cũng có thể làm mask bị giả nếu source không tốt. Fresnel có thể tạo rim đẹp, nhưng cũng có thể biến shield thành một quả cầu nhựa phát sáng nếu thiếu surface breakup và hit response.")
    body(doc, "Cuốn sách này đi từ cách nhìn tới cách làm. Trước hết, ta học cách nhìn hình ảnh như dữ liệu. Sau đó, ta dùng toán shader để nắn dữ liệu thành shape, mask, motion và color. Tiếp theo, ta đặt material vào cơ thể thật của realtime rendering: CPU, GPU, overdraw, texture bandwidth, blend mode và post-process. Cuối cùng, ta quay lại production, nơi mỗi quyết định phải trả lời ba câu hỏi: effect có đọc được không, có tune được không, và có sống được trong budget không.")
    body(doc, "Mục tiêu không phải là biến người học thành người thuộc nhiều công thức nhất. Mục tiêu là giúp người học mở một material bất kỳ, chia nó thành các lớp dễ hiểu, tìm được nguyên nhân khi visual hỏng, và giải thích được quyết định của mình bằng ngôn ngữ mà artist, technical artist và programmer đều có thể trao đổi.")
    note(doc, "Ghi nhớ", "Một material tốt không chỉ tạo ra hình ảnh đẹp trong preview. Nó phải giữ được ý đồ thị giác trong scene thật, dưới camera thật, với post-process thật và giới hạn phần cứng thật.")


CHAPTERS = [
    {
        "title": "Digital Eyes - tập nhìn hình ảnh như dữ liệu",
        "image": ("v9_shader_lane_diagram.png", "Một effect có thể được đọc theo các làn dữ liệu: UV, mask, color, opacity, interaction và cost."),
        "sections": [
            ("Tại sao mắt của Technical VFX Artist phải khác mắt của người xem",
             [
                 "Người xem bình thường nhìn một vệt lửa và nói rằng nó đẹp, yếu, bẩn, mạnh, giả hoặc hấp dẫn. Technical VFX Artist phải nhìn cùng vệt lửa đó và hỏi thêm một lớp câu hỏi khác. Vệt sáng này đến từ texture hay procedural mask? Nó đang cộng sáng bằng additive hay đang alpha blend? Edge của nó sắc vì texture, vì Power, hay vì alpha bị compression? Nó bị mất trên nền sáng vì thiếu contrast hay vì bloom đang nuốt mất hierarchy?",
                 "Đó là lý do tôi gọi chương này là Digital Eyes. Bạn không chỉ nhìn bằng mắt mỹ thuật. Bạn nhìn bằng mắt dữ liệu. Một pixel trên màn hình không phải là màu cuối cùng của riêng material. Nó là kết quả của một chuỗi quyết định: texture import, UV, particle color, material graph, blend mode, depth, exposure, bloom, tonemapper và background. Khi bạn hiểu điều này, bạn sẽ bớt sửa effect bằng cảm giác mơ hồ và bắt đầu sửa bằng nguyên nhân.",
                 "Kỹ năng này rất giống việc nghe một bản nhạc và tách được bass, trống, melody, vocal. Người mới nghe một khối âm thanh. Người làm nghề nghe từng lớp. Với material cũng vậy. Người mới thấy một hiệu ứng phát sáng. Người có Digital Eyes thấy mask nào đang sống, motion nào đang chạy, color nào đang quá mạnh, opacity nào đang làm bẩn silhouette, và cost nào có thể khiến GPU trả giá."
             ]),
            ("Từ visual sang dữ liệu",
             [
                 "Một cách luyện đơn giản là lấy bất kỳ effect nào và tạm thời cấm bản thân dùng từ đẹp hoặc xấu. Thay vào đó, hãy mô tả nó bằng dữ liệu. Shape của effect có đọc được khi thu nhỏ không? Mask có vùng đen trắng rõ không? Motion có cùng nhịp hay bị trượt nhau? Color có phân cấp core, edge và outer glow không? Interaction với scene có mềm hay bị cắt? Cost có nằm ở pixel area, texture fetch hay số lượng particle?",
                 "Khi bạn đặt câu hỏi như vậy, material không còn là một mạng node bí ẩn. Nó trở thành một hệ thống nhỏ. Hệ thống đó nhận dữ liệu, biến đổi dữ liệu, và xuất ra pixel. Điều này nghe đơn giản, nhưng nó thay đổi cách bạn học. Bạn sẽ không hỏi 'node này dùng làm gì' như một câu hỏi độc lập nữa. Bạn sẽ hỏi 'node này đang biến đổi lane nào, và visual nhận được thay đổi gì'.",
                 "Ví dụ, Power không phải là một node để làm mọi thứ sắc hơn. Power là một cách uốn curve 0-1. Khi áp vào radial mask, nó bóp ánh sáng vào core. Khi áp vào alpha smoke, nó có thể làm smoke mất mềm. Khi áp vào Fresnel, nó làm rim hẹp lại. Cùng một node, ba hậu quả khác nhau, vì dữ liệu đầu vào và mục tiêu visual khác nhau."
             ]),
            ("Debug bằng grayscale",
             [
                 "Nếu chỉ có một thói quen cần lấy từ chương này, đó là solo grayscale. Hãy tắt màu, tắt bloom, tắt những thứ làm bạn thấy phấn khích, rồi nhìn mask đen trắng. Mask đó có kể được câu chuyện chưa? Vùng nào sống, vùng nào chết, vùng nào nằm ở edge, vùng nào bị noise phá, vùng nào quá đều? Nếu grayscale đã sai, bản màu chỉ đang trang điểm cho lỗi.",
                 "Trong production, grayscale còn giúp bạn giao tiếp tốt hơn. Khi lead nói dissolve chưa organic, bạn có thể đưa ra source mask và edge band để chỉ rằng vấn đề nằm ở silhouette breakup chứ không phải ở màu edge. Khi programmer hỏi vì sao effect đắt, bạn có thể chỉ alpha card quá rộng thay vì nói chung chung rằng shader nặng.",
                 "Thói quen này cũng giúp bạn chuyển engine. Unreal, Unity, custom engine hay WebGL có tên node khác nhau, nhưng mask grayscale thì ở đâu cũng là mask grayscale. Một vùng 0-1 tốt sẽ sống được qua nhiều pipeline."
             ]),
        ],
        "notes": [
            ("Ghi nhớ", "Đừng để màu và bloom quyết định quá sớm. Một material VFX tốt thường có grayscale tốt trước khi nó có màu đẹp.", False),
        ],
        "case": [
            "Trong một buổi review thực tế, feedback về flare hoặc dissolve thường không đến dưới dạng công thức. Lead có thể nói rằng effect nhìn yếu, nhìn bẩn hoặc chưa có lực. Nếu bạn chỉ nghe phần cảm giác, bạn sẽ dễ tăng intensity hoặc thêm noise. Nếu bạn nghe bằng Digital Eyes, bạn sẽ kiểm tra grayscale trước: shape có đọc không, mask có bị phẳng không, alpha có quá rộng không, và bloom có đang che lỗi không.",
            "Cách đọc này làm cuộc nói chuyện chuyên nghiệp hơn. Bạn không cần tranh luận cảm giác ngay lập tức. Bạn có thể chỉ vào source mask, edge band hoặc alpha card và nói: vấn đề đang nằm ở đây. Khi một feedback mơ hồ được kéo về một layer cụ thể, việc sửa effect trở nên nhanh hơn rất nhiều."
        ],
        "code": [
            "rawMask    = TextureMask(UV)",
            "visible    = smoothstep(Threshold, Threshold + Softness, rawMask)",
            "edge       = smoothstep(Threshold - Width, Threshold, rawMask) - visible",
            "finalColor = BaseColor * visible + EdgeColor * edge"
        ],
    },
    {
        "title": "The Matrix of Art - toán học như cách nặn cảm giác",
        "image": ("v9_power_shaping_flare.png", "Power không chỉ là phép toán. Nó thay đổi cảm giác falloff của một vùng sáng."),
        "sections": [
            ("Toán trong shader không bắt đầu từ nỗi sợ",
             [
                 "Nhiều artist nghe tới toán shader là thấy xa lạ. Điều này dễ hiểu, vì toán thường được dạy như bài kiểm tra, không phải như công cụ tạo hình. Trong material VFX, toán nên được học theo cảm giác trước. SmoothStep là một lưỡi dao mềm. Power là bàn tay bóp ánh sáng vào core. Dot product là câu hỏi hai hướng có đang nhìn cùng phía không. Distance là cách biến không gian thành shape.",
                 "Khi bạn học toán theo cách đó, công thức không còn đứng riêng. Nó gắn với một hành động thị giác. Remap đổi thang đo để một mask nhạt trở thành mask rõ. Clamp giữ giá trị không vượt khỏi vùng an toàn. Lerp pha hai trạng thái. Sine tạo nhịp thở. Atan2 mở cánh cửa vào polar UV, nơi portal, vortex và radial slash bắt đầu có ngôn ngữ riêng.",
                 "Một Technical VFX Artist không cần biến mình thành nhà toán học thuần túy để làm shader tốt. Nhưng bạn cần đủ thân với các phép toán cơ bản để không sợ graph. Khi feedback tới, bạn cần biết mình đang thiếu contrast, thiếu edge width, thiếu falloff, sai space hay sai curve."
             ]),
            ("Từ một giá trị 0-1 tới một visual đọc được",
             [
                 "Rất nhiều material VFX bắt đầu từ một giá trị 0-1. Một noise sample trả về vùng xám. Một radial gradient trả về khoảng cách từ tâm. Một Fresnel trả về độ nghiêng giữa normal và hướng nhìn. Một depth comparison trả về khoảng cách tới geometry. Bản thân các giá trị này chưa phải effect. Effect xuất hiện khi bạn remap, clamp, threshold, làm mềm, nhân màu, đưa vào opacity hoặc emissive.",
                 "Hãy tưởng tượng một dissolve. Source mask là một bản đồ địa hình. Threshold là mực nước đang dâng. Vùng thấp biến mất trước, vùng cao còn lại sau. Edge band là bờ nước, nơi bạn có thể cho ánh sáng cháy lên. Nếu bạn hiểu hình ảnh đó, công thức không còn khô nữa. Nó chỉ là cách viết chính xác cho một câu chuyện rất trực quan.",
                 "Điểm quan trọng là edge phải đến từ cùng source mask với visible mask. Nếu bạn vẽ edge bằng một texture khác hoặc một outline tùy tiện, dissolve sẽ không bám logic. Người xem có thể không biết công thức, nhưng mắt họ sẽ thấy có gì đó giả."
             ]),
            ("SDF, polar và những shape không cần texture",
             [
                 "Signed Distance Field là một ví dụ đẹp về việc toán có thể vẽ hình. Thay vì lưu sẵn một vòng tròn trong texture, bạn tính khoảng cách từ pixel tới tâm. Khoảng cách đó trở thành dữ liệu. Từ dữ liệu đó, bạn tạo circle, ring, soft edge, outline hoặc shockwave. Khi cần portal, bạn thêm angle bằng atan2. Khi cần motion xoáy, bạn offset angle theo time hoặc radius.",
                 "Điều này không có nghĩa là texture không quan trọng. Texture vẫn rất mạnh, đặc biệt khi bạn cần chi tiết organic, hand-painted shape hoặc baked simulation. Nhưng khi bạn biết dựng shape bằng toán, bạn có thêm một lớp tự do. Bạn có thể làm prototype nhanh, tạo parameter sạch, và chỉ đưa texture vào khi texture thật sự mang thông tin mà toán không nên gánh.",
                 "Trong production, sự cân bằng giữa procedural và texture là một quyết định. Procedural dễ tune và nhẹ memory, nhưng có thể tốn instruction hoặc nhìn quá sạch. Texture giàu chi tiết, nhưng tốn bandwidth và dễ bị compression/mip phá. Người giỏi không chọn một phe; họ chọn theo effect."
             ]),
        ],
        "notes": [
            ("Cảnh báo production", "Đừng biến mọi thứ thành procedural chỉ vì nó trông thông minh. Nếu một flipbook hoặc texture nhỏ giải quyết visual tốt hơn và rẻ hơn, đó là lựa chọn production hợp lý.", True),
        ],
        "case": [
            "Một portal không texture là bài tập rất tốt, nhưng nó cũng dễ tạo ảo giác rằng procedural luôn tốt hơn. Trong prototype, bạn có thể dựng ring bằng distance, thêm polar UV, rồi làm edge bằng SmoothStep. Kết quả rất sạch và dễ tune. Nhưng khi art direction cần vết nứt, ký tự cổ, vân năng lượng hoặc cảm giác hand-painted, texture có thể là dữ liệu đúng hơn.",
            "Quyết định production không phải là procedural hay texture. Quyết định đúng là phần nào nên được tính, phần nào nên được vẽ, và phần nào nên được bake. Một Technical VFX Artist giỏi biết chuyển qua lại giữa ba cách đó mà không xem cách nào là tôn giáo."
        ],
        "code": [
            "r     = length(UV - Center)",
            "ring  = 1 - smoothstep(Width, Width + Softness, abs(r - Radius))",
            "edge  = ring * EdgeIntensity",
            "color = ring * Tint"
        ],
    },
    {
        "title": "Cơ bắp của realtime - CPU, GPU và cái giá của pixel",
        "image": ("v9_shader_complexity_overdraw.png", "Overdraw là một trong những chi phí lớn nhất của translucent VFX."),
        "sections": [
            ("Một effect không chạy trong sách vở",
             [
                 "Một material đẹp trong editor vẫn có thể là một vấn đề trong game. Lý do rất đơn giản: game không render một material trong chân không. Game render hàng nghìn object, ánh sáng, shadow, post-process, UI, particle, animation và gameplay logic. CPU phải chuẩn bị scene. GPU phải vẽ scene. Mỗi frame chỉ có một ngân sách thời gian rất nhỏ.",
                 "CPU thường lo gameplay, transform, culling, submit draw call, sort translucent object, update particle hoặc gửi parameter. GPU nhận công việc và xử lý theo pipeline: vertex, rasterization, fragment/pixel, blending và các pass sau đó. Khi bạn làm VFX, bạn thường chạm mạnh vào pixel shader, blending và overdraw. Đây là vùng rất dễ đắt, nhất là với fire, smoke, glow, aura và UI effect có alpha rộng.",
                 "Điều đáng chú ý là một shader có instruction count không quá cao vẫn có thể đắt nếu nó phủ quá nhiều pixel. Một card smoke lớn, alpha rộng và chồng nhiều lớp có thể làm GPU vẽ đi vẽ lại cùng một vùng màn hình. Đây là lý do overdraw thường quan trọng hơn vài phép Multiply nhỏ."
             ]),
            ("Translucent VFX và bài toán diện tích",
             [
                 "Opaque mesh thường ghi depth và giúp GPU loại bỏ nhiều pixel phía sau. Translucent VFX lại phức tạp hơn, vì nó cần trộn với background và thường không ghi depth theo cách opaque. Nếu bạn có mười card smoke chồng lên nhau, một pixel có thể bị shader xử lý nhiều lần. Người xem chỉ thấy khói mềm; GPU thấy một hóa đơn nhiều lớp.",
                 "Khi tối ưu VFX, nhiều người mới nhìn instruction count trước. Điều đó không sai, nhưng chưa đủ. Bạn cần hỏi effect chiếm bao nhiêu phần trăm màn hình, có bao nhiêu layer chồng nhau, alpha card có bị rỗng quá nhiều không, texture có quá lớn so với kích thước hiển thị không, và effect xuất hiện bao nhiêu lần trong một frame.",
                 "Với mobile hoặc iGaming, những câu hỏi này trở nên thực tế hơn nữa. Một sparkle nhỏ có thể rẻ khi xuất hiện một lần, nhưng đắt khi có hàng trăm instance. Một glow đẹp có thể ổn trên PC, nhưng trên web/mobile nó trở thành lý do khiến animation giật."
             ]),
            ("Tối ưu không có nghĩa là làm xấu đi",
             [
                 "Tối ưu tốt không phải là xóa bừa node. Tối ưu tốt là giữ phần người xem thật sự đọc được và bỏ phần họ không cần. Nếu một projectile cần hướng bay rõ, hãy giữ core, head, tail và timing. Nếu noise nhỏ trên shell không đọc được ở camera gameplay, bạn có thể giảm nó hoặc bake nó. Nếu glow phủ quá rộng, crop alpha và giảm card area có thể hiệu quả hơn việc tranh luận thêm bớt vài instruction.",
                 "Một nguyên tắc thực dụng là tối ưu theo thứ tự: silhouette, area, layer, texture fetch, instruction. Silhouette quyết định effect có đọc được không. Area và layer quyết định overdraw. Texture fetch quyết định bandwidth. Instruction quyết định độ nặng tính toán. Thứ tự này không tuyệt đối, nhưng nó giúp bạn tránh tối ưu nhầm chỗ.",
                 "Technical VFX Artist giỏi thường biết nói câu này: bản đẹp nhất chưa chắc là bản production tốt nhất. Bản production tốt là bản đạt visual, giữ gameplay readable, dễ tune, và sống được trong budget."
             ]),
        ],
        "notes": [
            ("Ghi nhớ", "Nếu effect là translucent và chiếm nhiều màn hình, hãy nghĩ tới overdraw trước khi nghĩ tới vài phép toán nhỏ trong graph.", False),
        ],
        "case": [
            "Một win effect trong iGaming có thể trông rất nhỏ trong layout UI, nhưng nó xuất hiện nhiều lần, trên nhiều thiết bị, và thường đi cùng animation khác. Nếu mỗi sparkle là một translucent quad rộng, nếu mỗi glow dùng texture lớn, và nếu mọi layer đều additive, chi phí cộng dồn sẽ đến nhanh hơn cảm giác của artist.",
            "Trong trường hợp này, tối ưu tốt thường không bắt đầu bằng HLSL. Nó bắt đầu bằng asset authoring: crop alpha sát hơn, pack mask, bake motion phụ, giảm layer runtime và giới hạn số particle cùng lúc. Shader thông minh không cứu được một thiết kế asset lãng phí."
        ],
        "code": [
            "ApproxCost = PixelArea * LayerCount * ShaderCost",
            "If PixelArea is large:",
            "    crop alpha, reduce card size, reduce overlapping layers",
            "If TextureFetch is high:",
            "    pack masks, lower resolution, check mip/compression"
        ],
    },
    {
        "title": "Ngôn ngữ shader - khi node graph không còn đủ",
        "image": ("v11_realtime_contract_simple.png", "Material tốt cần hợp đồng dữ liệu rõ với Niagara, Blueprint hoặc hệ runtime tương đương."),
        "sections": [
            ("Code shader không phải nghi thức thần bí",
             [
                 "Có một giai đoạn trong hành trình học material khi node graph bắt đầu giống một mê cung. Lúc đó, nhiều artist nghĩ HLSL hoặc GLSL là một thế giới khác, chỉ dành cho programmer. Thực tế bình tĩnh hơn. Shader code cũng chỉ nhận dữ liệu, biến đổi dữ liệu và trả kết quả. Điều khác biệt là code buộc bạn viết rõ ý hơn.",
                 "Một fragment shader tối giản có thể chỉ trả về một màu. Từ đó, bạn thêm UV, sample texture, nhân tint, tính alpha, thêm time, thêm noise, thêm normal, thêm depth. Mỗi bước vẫn là một câu rất nhỏ. Nếu bạn đọc được câu nhỏ đó trong code, bạn sẽ đọc node graph tốt hơn, vì bạn hiểu ý tưởng nằm dưới hình dạng node.",
                 "Điểm quan trọng là không dùng HLSL để tỏ ra senior. Một node graph rõ ràng vẫn rất tốt, đặc biệt trong team art. HLSL nên xuất hiện khi graph lặp quá nhiều, khi công thức diễn đạt bằng code sạch hơn, hoặc khi bạn cần một function nhỏ có input/output rõ."
             ]),
            ("Types, parameters và texture sample",
             [
                 "Shader language có vài từ vựng cơ bản mà bạn cần thân quen. float là một số. float2 thường là UV. float3 thường là position, normal hoặc color RGB. float4 thường là color có alpha. Texture không tự sinh ra màu; bạn phải sample texture bằng UV. Parameter là dữ liệu từ ngoài shader đưa vào, như DissolveAmount, EdgeWidth, Tint hoặc DistortionStrength.",
                 "Unreal gọi tên theo hệ material của Unreal. Unity Shader Graph có cách đóng gói riêng. GLSL có vec3, mix, fract. HLSL có float3, lerp, frac. Tên khác nhau nhưng tư duy giống nhau. Khi bạn hiểu dữ liệu, bạn không bị kẹt vào tên node.",
                 "Một shader production tốt cần input rõ. Nếu function EdgeBand nhận mask, threshold, width và softness, nó nên trả về visible mask và edge mask. Nó không nên tự quyết định màu, emissive, Niagara binding và opacity cùng lúc. Function nhỏ sống lâu hơn function tham lam."
             ]),
            ("Custom node trong Unreal và giới hạn của nó",
             [
                 "Trong Unreal, Custom Material Expression cho phép viết HLSL trong material. Công cụ này hữu ích, nhưng nó không phải lối thoát cho mọi graph rối. Nếu bạn chưa hiểu graph, chuyển nó sang code chỉ làm lỗi khó thấy hơn. Nếu bạn hiểu rõ công thức và graph đang lặp hoặc quá dài, custom node có thể làm material sạch hơn.",
                 "Một ví dụ hợp lý là nhiều hit ripple trên shield. Nếu mỗi hit có position, radius, width, intensity và decay, node graph có thể nhanh chóng thành một tấm lưới dây nối. Một loop HLSL nhỏ với count cố định có thể dễ đọc hơn, miễn là team chấp nhận và bạn có debug output.",
                 "Khi dùng HLSL, hãy viết như đang viết cho người khác đọc. Đặt tên biến theo ý nghĩa visual. Clamp output khi cần. Ghi rõ range input. Và luôn có một cách xem kết quả trung gian bằng grayscale hoặc debug color."
             ]),
        ],
        "notes": [
            ("Cảnh báo production", "Code shader tệ nguy hiểm hơn node graph tệ, vì nó trông gọn trong khi che giấu quyết định. Hãy dùng HLSL khi nó làm ý tưởng rõ hơn, không phải khi nó chỉ làm graph ngắn hơn.", True),
        ],
        "case": [
            "Một custom node nhỏ có thể làm graph sạch hơn rất nhiều nếu nó thay thế một cụm logic lặp lại. Ví dụ EdgeBand, Remap01 hoặc PolarUV là các function có đầu vào và đầu ra rõ. Người khác đọc được tên function và hiểu ngay ý định.",
            "Ngược lại, một custom node chứa nửa material bên trong sẽ tạo ra một hộp đen. Artist không biết parameter nào quan trọng. Lead không biết debug ở đâu. Programmer không muốn maintain. Vì vậy, HLSL tốt trong material VFX thường là HLSL nhỏ, rõ và có đường debug."
        ],
        "code": [
            "float EdgeBand(float mask, float t, float width, float soft)",
            "{",
            "    float inner = smoothstep(t - width, t, mask);",
            "    float outer = smoothstep(t, t + soft, mask);",
            "    return saturate(inner - outer);",
            "}"
        ],
    },
    {
        "title": "Mask, UV và motion - ba lớp dễ bị sửa nhầm",
        "image": ("v9_mask_solo_dissolve.png", "Một dissolve tốt bắt đầu từ source mask, visible mask và edge band rõ ràng."),
        "sections": [
            ("Mask quyết định phần nào của effect được phép tồn tại",
             [
                 "Trong rất nhiều material VFX, mask là phần quan trọng nhất nhưng lại bị xem nhẹ nhất. Một mask tốt cho bạn biết vùng nào hiện, vùng nào mất, vùng nào là edge, vùng nào nên phát sáng, vùng nào nên distort, vùng nào nên mềm đi khi gặp depth. Khi mask sai, màu đẹp chỉ làm lỗi khó nhận ra hơn.",
                 "Một dissolve đơn giản có thể bắt đầu bằng noise. Nhưng production dissolve hiếm khi chỉ là noise. Nó có hướng, có tốc độ, có edge width, có breakup, có sync với animation hoặc gameplay. Nếu nhân vật tan từ tay lên vai, source mask cần chứa thông tin hướng. Nếu portal mở từ tâm ra ngoài, source mask nên là distance field hoặc radial gradient. Nếu shield bị hit tại một điểm, source mask có thể là distance từ world position tới hit position.",
                 "Vì vậy, khi sửa effect, hãy hỏi mask đang đại diện cho điều gì. Nó là hình dạng mỹ thuật, dữ liệu gameplay, khoảng cách, độ sâu, lifetime particle hay frame index? Câu trả lời quyết định cách bạn debug."
             ]),
            ("UV motion không phải texture đang chạy",
             [
                 "Khi dùng Panner, texture không thật sự di chuyển. Bạn đang thay đổi địa chỉ đọc texture. Câu này rất quan trọng. Nếu bạn hiểu UV là địa chỉ, bạn sẽ hiểu vì sao scale, rotate, polar, flow map và distortion đều là các cách làm lệch câu hỏi 'tôi đang đọc texture ở đâu'.",
                 "Một ngọn lửa card thường có nhiều motion. Flipbook giữ motion lớn. Panner nhẹ có thể thêm shimmer. Distortion có thể làm edge rung. Nhưng nếu mọi layer chạy cùng tốc độ, effect sẽ trông máy móc. Motion tốt cần phase, offset và hierarchy. Core có thể nhanh hơn edge. Smoke có thể chậm hơn fire. Sparkle có thể bật theo nhịp riêng thay vì trôi đều.",
                 "Một lỗi thường gặp là dùng UV motion để sửa vấn đề shape. Nếu silhouette của effect yếu, panner không cứu được. Nếu alpha quá rộng, rotator không làm nó rẻ hơn. Nếu mask không có hướng, thêm noise pan có thể chỉ làm effect bẩn hơn."
             ]),
            ("Derived 0-1 UV trong SubUV",
             [
                 "Khi làm flipbook, bạn thường có atlas UV trên toàn sheet và local UV bên trong frame hiện tại. Hai khái niệm này cần tách rõ. Atlas UV cho biết bạn đang ở đâu trên toàn texture. Derived 0-1 UV cho biết bạn đang ở đâu trong ô frame hiện tại. Nếu bạn muốn tạo radial fade trong từng frame, bạn cần local UV, không phải atlas UV.",
                 "Đây là lý do các material như M_SubUV_Utilities đáng học. Nó không chỉ là utility. Nó dạy bạn cách renderer, atlas và material ký hợp đồng với nhau. Niagara hoặc Sprite Renderer chọn frame. Material đọc texture. Nếu material cần logic theo từng frame, nó phải suy ra local coordinate.",
                 "NumberGrid là công cụ debug rất tốt ở đây. Khi grid sai, bạn thấy ngay frame bị lệch. Khi local UV sai, radial fade hoặc border sẽ chạy theo toàn atlas thay vì từng ô. Đây là loại lỗi nhỏ nhưng rất dễ làm artist mất nhiều giờ nếu không biết nhìn."
             ]),
        ],
        "notes": [
            ("Ghi nhớ", "UV là địa chỉ đọc dữ liệu. Khi bạn làm méo UV, bạn không di chuyển texture; bạn thay đổi câu hỏi sample.", False),
        ],
        "case": [
            "Khi một flame card nhìn giả, nhiều người thêm noise pan. Nhưng nếu flipbook đã có motion chính, noise pan quá mạnh có thể làm lửa trôi khỏi nhịp gốc. Người xem không gọi tên được lỗi, nhưng họ cảm thấy lửa không còn trọng lượng.",
            "Một cách debug tốt là tắt lần lượt từng motion layer. Giữ flipbook, tắt panner. Giữ panner, tắt distortion. Sau đó xem layer nào thật sự thêm cảm giác và layer nào chỉ thêm nhiễu. Không phải motion nào cũng làm effect sống hơn."
        ],
        "code": [
            "atlasUV = TexCoord",
            "frameSize = 1 / float2(Columns, Rows)",
            "frameOffset = FrameIndexToOffset(FrameIndex)",
            "localUV = (atlasUV - frameOffset) / frameSize"
        ],
    },
    {
        "title": "Blend, emissive và depth - nơi material gặp renderer",
        "image": ("v9_failure_depthfade_before_after.png", "DepthFade không làm smoke đẹp hơn theo nghĩa màu sắc; nó làm giao tuyến với scene bớt giả."),
        "sections": [
            ("Blend mode là quyết định render, không chỉ là mỹ thuật",
             [
                 "Additive, translucent và masked thường được chọn như lựa chọn mỹ thuật. Điều đó đúng một phần, nhưng chưa đủ. Blend mode quyết định cách pixel của material trộn với frame đã có, quyết định depth behavior, sorting và cost. Additive cộng sáng nên hợp với flare, spark, magic streak, nhưng nó dễ biến mất trên nền sáng. Translucent mềm hơn, hợp với smoke, fire card, glass hoặc energy field, nhưng dễ gặp sorting và overdraw. Masked sắc và ổn định hơn, nhưng edge có thể răng cưa nếu không xử lý tốt.",
                 "Khi một lead nói effect không đủ mạnh, nhiều người tăng emissive. Nhưng nếu blend mode là additive và background đã sáng, tăng emissive có thể không giải quyết readability. Bạn có thể cần tăng contrast của mask, thêm core tối hơn ở phía trong, thay đổi color hierarchy, hoặc thêm một layer non-additive để giữ shape.",
                 "Blend mode cũng liên quan tới platform. Một effect translucent nhiều layer có thể ổn trên PC nhưng quá đắt trên mobile. Một masked/dithered variant có thể không mềm bằng, nhưng ổn định và rẻ hơn. Đây không phải chuyện đúng sai tuyệt đối; đó là tradeoff production."
             ]),
            ("Emissive cần hierarchy",
             [
                 "Emissive rất dễ gây nghiện. Chỉ cần tăng intensity, effect lập tức sáng hơn trong preview. Nhưng emissive không có hierarchy sẽ nhanh chóng thành một mảng cháy trắng. Một fire card tốt thường có core nóng, edge chuyển màu, outer smoke hoặc glow yếu hơn. Một magic orb tốt có core, shell, rim và accent. Nếu mọi lớp đều sáng ngang nhau, người xem mất điểm tựa.",
                 "Bloom làm vấn đề này rõ hơn. Bloom khuếch đại vùng sáng và có thể làm một core nhỏ trở nên đẹp. Nhưng nếu bạn dựa hoàn toàn vào bloom, khi post-process thay đổi hoặc scene exposure khác đi, effect sẽ đổi cảm giác mạnh. Vì vậy, hãy tắt bloom tạm thời khi debug shape. Nếu không có bloom mà effect không đọc được gì, bạn đang mượn bloom quá nhiều.",
                 "Một cách nghĩ thực dụng là chia emissive thành ba lớp: core đọc shape, edge đọc transition, accent đọc motion hoặc energy. Mỗi lớp nên có vai trò. Nếu một layer không có vai trò, nó có thể chỉ là noise trang trí."
             ]),
            ("DepthFade và soft particles",
             [
                 "DepthFade giải quyết một lỗi rất cụ thể: translucent card cắt thẳng vào geometry. Smoke, fire, splash, dust và magic cloud thường cần nó. Khi không có DepthFade, giao tuyến với sàn hoặc tường nhìn như một mặt phẳng bị lộ. Khi có DepthFade, opacity giảm dần gần geometry, tạo cảm giác mềm hơn.",
                 "Nhưng DepthFade không phải thuốc chữa mọi thứ. Nếu card quá lớn, alpha quá rộng, sorting sai hoặc smoke texture không có shape tốt, DepthFade chỉ che được một phần. Nó cũng có cost và phụ thuộc depth buffer. Bạn cần test nó với camera gần, xa, góc thấp và scene có nhiều geometry.",
                 "Một lỗi production phổ biến là để fade distance quá lớn. Smoke bắt đầu biến mất trước khi chạm sàn, làm effect bị mỏng. Fade distance quá nhỏ thì giao tuyến vẫn cứng. Giá trị đúng phụ thuộc scale scene, kích thước card và mục tiêu visual."
             ]),
        ],
        "notes": [
            ("Cảnh báo production", "Đừng tăng emissive để sửa mọi feedback. Nhiều khi vấn đề nằm ở mask, hierarchy, blend mode hoặc background contrast.", True),
        ],
        "case": [
            "Một flare additive đẹp trên nền đen có thể biến mất trên nền trời sáng. Nếu bạn chỉ tăng intensity, nó sáng hơn trong một số cảnh nhưng vẫn thiếu shape ở cảnh khác. Vấn đề không phải chỉ là lượng sáng. Vấn đề là flare không có đủ hierarchy để tự đọc trên nhiều background.",
            "Giải pháp có thể là thay falloff, thêm core rõ hơn, giảm outer glow, đổi blend cho một layer phụ, hoặc kiểm soát bloom. Đây là lý do blend mode phải được xem như quyết định render chứ không phải lựa chọn màu sắc."
        ],
        "code": [
            "opacityRaw  = TextureAlpha * ParticleColor.a",
            "opacitySoft = DepthFade(opacityRaw, FadeDistance)",
            "emissive    = Core * CoreIntensity + Edge * EdgeIntensity",
            "final       = emissive * opacitySoft"
        ],
    },
    {
        "title": "Production stories - mobile, stylized, iGaming và AAA",
        "image": ("v11_imagegen_vfx_atlas.png", "Các nhóm effect khác nhau cần quyết định khác nhau. Một công thức không phục vụ mọi style."),
        "sections": [
            ("Mobile không tha thứ cho sự hào nhoáng không cần thiết",
             [
                 "Mobile VFX có một sự thật hơi lạnh lùng: nếu người chơi không thấy chi tiết, GPU vẫn có thể phải trả tiền cho chi tiết đó. Một noise layer nhỏ trong fire, một distortion nhẹ trong win effect, một alpha card rộng hơn silhouette thật, tất cả đều có thể cộng lại thành frame-time xấu khi effect xuất hiện nhiều lần.",
                 "Điều này không có nghĩa mobile phải xấu. Nhiều game mobile stylized rất đẹp vì chúng chọn đúng thứ để giữ. Chúng giữ silhouette rõ, color sạch, timing tốt, sprite hoặc flipbook được author cẩn thận. Chúng không cố mô phỏng mọi thứ bằng procedural shader runtime nếu một atlas nhỏ làm tốt hơn.",
                 "Với iGaming, bài toán còn rõ hơn. Win effect phải sáng, vui, dễ đọc, lặp đi lặp lại và không làm UI khó nhìn. Rất nhiều sparkle, shine sweep hoặc coin burst nên được bake một phần. Runtime chỉ cần tint, scale, fade và chọn frame. Đó không phải là lười. Đó là production."
             ]),
            ("Stylized không có nghĩa là đơn giản",
             [
                 "Stylized VFX thường bị hiểu nhầm là ít thật hơn nên dễ hơn. Thực tế, stylized đòi hỏi quyết định rõ hơn. Anime slash cần hard shape, taper, smear, impact timing và color accent. Một portal stylized cần ring đọc rõ trước, swirl sau, particle accent cuối cùng. Một shield stylized cần rim mạnh nhưng không được thành quả cầu nhựa phát sáng.",
                 "Điểm khó của stylized là mọi thứ lộ. Noise sai lộ. Timing sai lộ. Shape yếu lộ. Nếu realistic VFX đôi khi có thể dựa vào complexity để che lỗi, stylized buộc bạn thiết kế shape thật chắc. Vì vậy, nhiều khi stylized shader cần ít node hơn nhưng nhiều art direction hơn.",
                 "Một cách làm tốt là dựng effect ở grayscale trước. Nếu slash, portal hoặc shield không đọc được khi chỉ có trắng đen, màu sắc sẽ không cứu nó. Sau đó mới thêm palette, emissive, bloom và particle accent."
             ]),
            ("AAA readability và bài học khi scene quá phức tạp",
             [
                 "AAA không có nghĩa là mọi shader đều phức tạp. AAA thường có nghĩa là effect phải sống trong scene phức tạp: lighting thay đổi, camera di chuyển, motion blur, post-process, gameplay đông, nhiều asset cùng lúc. Trong môi trường đó, readability là kỹ năng sống còn.",
                 "Một projectile đẹp trong turntable có thể biến mất trong combat. Một impact đẹp trong black background có thể cháy trắng trong level tuyết. Một smoke đẹp trong material preview có thể cắt sàn khi đặt vào map thật. Đây là lý do production review phải diễn ra trong context thật càng sớm càng tốt.",
                 "Technical VFX Artist ở môi trường này cần biết khi nào vấn đề nằm ở material và khi nào không. Có lúc shader đúng nhưng lighting sai. Có lúc texture tốt nhưng import setting sai. Có lúc material đẹp nhưng Niagara timing làm nó yếu. Có lúc effect đẹp nhưng gameplay không đọc được. Khả năng phân biệt các trường hợp này mới là giá trị thật."
             ]),
        ],
        "notes": [
            ("Ghi nhớ", "Production không thưởng cho shader phức tạp nếu một giải pháp đơn giản đạt visual tốt hơn, rẻ hơn và dễ tune hơn.", False),
        ],
        "case": [
            "Trong một game stylized, slash đẹp thường đến từ shape và timing hơn là simulation. Một mesh arc có taper tốt, một texture hand-painted sạch và một curve scale đúng nhịp có thể thuyết phục hơn một shader procedural nhiều noise.",
            "Trong một game realistic, cùng slash đó có thể cần sparks, debris, lighting interaction và smoke nhỏ. Cùng một loại effect, hai style dẫn tới hai quyết định khác nhau. Đây là lý do học recipe không đủ; bạn phải học cách đọc art direction."
        ],
        "code": [
            "HighSpec:  flipbook + depth fade + distortion + extra particles",
            "LowSpec:   packed atlas + tint + scale + alpha fade",
            "UI/iGame:  baked sequence + deterministic timing + minimal shader"
        ],
    },
    {
        "title": "Đọc VFXSamplePack như một người làm production",
        "image": ("v11_realtime_contract_simple.png", "Khi đọc material mẫu, hãy tìm contract giữa renderer, particle data và material graph."),
        "sections": [
            ("Đừng mở graph rồi đuổi theo dây nối",
             [
                 "Khi mở một material có sẵn như M_Flare, M_Mesh_Add, M_FlamesCard hoặc M_SubUV_Utilities, phản xạ tự nhiên là nhìn dây nối. Càng nhiều node, bạn càng muốn lần theo từng dây từ trái sang phải. Cách này có thể dùng, nhưng rất dễ làm bạn lạc. Một cách đọc production tốt hơn là chia graph thành lane.",
                 "Lane đầu tiên là UV. Material đang đọc texture bằng TexCoord gốc, panner, rotator, polar hay SubUV? Lane thứ hai là mask. Vùng nào quyết định opacity, vùng nào quyết định edge, vùng nào quyết định distortion? Lane thứ ba là color/emissive. Màu đến từ texture, parameter, Particle Color hay gradient? Lane thứ tư là interaction. Có DepthFade, CameraFade, SceneDepth, Fresnel, WPO hay Dynamic Parameter không? Lane cuối cùng là cost.",
                 "Khi bạn đọc như vậy, graph không còn là spaghetti. Nó trở thành một bản đồ. Bạn có thể nói với người khác rằng material này dùng Particle Color để tint/fade, dùng texture alpha làm opacity, dùng Power để chỉnh falloff, dùng DepthFade để mềm giao tuyến, và expose Intensity để Niagara hoặc material instance điều khiển."
             ]),
            ("M_Flare và bài học về additive",
             [
                 "Một flare additive dạy rất nhiều điều cơ bản. Nó thường có texture hoặc radial mask, một tint, một intensity và một cách fade theo particle/camera. Điều quan trọng là flare không chỉ cần sáng. Nó cần đọc được trên nền khác nhau. Nếu background tối, additive flare rất đẹp. Nếu background sáng, nó có thể mất gần hết shape.",
                 "Khi học M_Flare, hãy thử tắt bloom và giảm intensity. Lúc đó, bạn sẽ thấy shape thật của mask. Nếu mask chỉ là một blob mềm không có hierarchy, bloom đang làm quá nhiều việc. Hãy thử Power khác nhau, contrast khác nhau, và test trên nền sáng/tối.",
                 "Bài học production ở đây là additive không miễn phí về mặt readability. Nó cộng sáng, nhưng không tự tạo contrast. Muốn flare sống trong nhiều scene, bạn cần nghĩ tới core, falloff, background và exposure."
             ]),
            ("M_SubUV_Utilities và bài học về atlas",
             [
                 "M_SubUV_Utilities đáng học vì nó chạm vào một lỗi rất thường gặp: material và renderer không thống nhất về frame. Khi SubUV được định nghĩa ở Sprite Renderer hoặc Niagara, material cần biết nó đang đọc frame nào và đôi khi cần local UV trong frame đó. Nếu không, radial fade, border hoặc per-frame logic sẽ chạy sai.",
                 "NumberGrid là một texture debug tuyệt vời trong trường hợp này. Nó không đẹp, nhưng nó nói thật. Nếu bạn thấy số frame sai, grid sai. Nếu gradient chạy qua toàn sheet thay vì từng ô, local UV sai. Nếu frame bị giật, frame interpolation hoặc frame rate có vấn đề.",
                 "Khi học atlas, hãy nhớ rằng texture sheet là dữ liệu đóng gói. Renderer chọn ô. Material đọc ô. Nếu một trong ba phần này không có cùng quy ước, effect sẽ hỏng theo cách rất khó đoán nếu bạn chỉ nhìn final fire/smoke."
             ]),
        ],
        "notes": [
            ("Ghi nhớ", "Khi đọc material có sẵn, hãy viết breakdown bằng sáu dòng: Shape, UV, Mask, Color, Interaction, Cost. Nếu không viết được, bạn chưa thật sự đọc xong.", False),
        ],
        "case": [
            "Khi đọc M_Flare, đừng sửa parameter ngay. Hãy tạo một trang ghi chú nhỏ. Texture nào tạo shape? Particle Color đi vào đâu? Intensity có nằm trước hay sau mask? Có camera fade không? Có chỗ nào phụ thuộc bloom không?",
            "Khi đọc M_SubUV_Utilities, hãy thay texture bằng NumberGrid trước khi nhìn fire hoặc smoke. NumberGrid không đẹp, nhưng nó nói thật. Nó cho bạn biết frame nào đang được đọc và UV local có đúng hay không."
        ],
        "code": [
            "Material reading lanes:",
            "1. Shape / silhouette",
            "2. UV / motion",
            "3. Mask / opacity / edge",
            "4. Color / emissive",
            "5. Interaction / renderer data",
            "6. Cost / platform risk"
        ],
    },
    {
        "title": "Lộ trình luyện tập - từ hiểu tới phản xạ",
        "image": ("v9_space_vectors.png", "Space transform là một trong những phần nên luyện chậm, vì nó xuất hiện ở nhiều shader nâng cao."),
        "sections": [
            ("Đọc không đủ, phải dựng lại",
             [
                 "Shader giống ngoại ngữ ở một điểm rất rõ: đọc hiểu chưa đủ. Bạn có thể hiểu SmoothStep khi nhìn công thức, nhưng nếu chưa tự dùng nó để làm dissolve edge, ring, foam hoặc shockwave, bạn chưa có phản xạ. Bạn có thể hiểu Particle Color trên giấy, nhưng nếu chưa từng quên nhân ParticleColor.a rồi tự debug vì Niagara fade không tác dụng, bài học đó chưa bám vào tay.",
                 "Vì vậy, lộ trình học material nên có output nhỏ liên tục. Mỗi output phải đủ nhỏ để hoàn thành, nhưng đủ thật để có lỗi. Một ngày dựng mask grayscale. Một ngày làm radial ring. Một ngày làm fire card baseline. Một ngày đọc M_Flare. Một ngày viết HLSL EdgeBand. Những bài nhỏ này tạo vốn từ vựng thực chiến.",
                 "Đừng đợi tới khi hiểu hết GPU mới làm effect. Cũng đừng chỉ copy recipe mà không hiểu cost. Học tốt nhất là xen kẽ: đọc một khái niệm, dựng một bản tối giản, phá nó, sửa nó, rồi viết lại bằng lời của mình."
             ]),
            ("Một lộ trình 8 tuần thực dụng",
             [
                 "Tuần đầu tiên nên dành cho grayscale. Bạn dựng mask từ gradient, noise, radial distance, UV.x, UV.y và texture alpha. Mỗi mask phải có threshold, softness và debug view. Tuần thứ hai chuyển sang shaping: Power, SmoothStep, remap, contrast, edge band. Đừng thêm màu quá sớm.",
                 "Tuần ba học UV motion: panner, rotator, polar, distortion và flow. Tuần bốn học blend mode, emissive, opacity, DepthFade và CameraFade. Tuần năm học Particle Color, Dynamic Parameter, SubUV và NumberGrid. Tuần sáu học Fresnel, space transform, WPO và shield hit. Tuần bảy học optimization: overdraw, texture fetch, channel packing và mobile variant. Tuần tám đọc lại material thật trong VFXSamplePack và viết breakdown.",
                 "Điều quan trọng không phải hoàn thành lộ trình hoàn hảo. Điều quan trọng là mỗi tuần có artifact. Một material nhỏ, một screenshot grayscale, một note debug, một câu giải thích. Sau vài tháng, bạn sẽ có thư viện tư duy của riêng mình."
             ]),
            ("Tự kiểm tra bằng câu hỏi production",
             [
                 "Sau mỗi effect, hãy hỏi vài câu khó. Nếu người khác mở material này, họ có hiểu parameter không? Nếu effect spawn 100 lần, nó có còn ổn không? Nếu background sáng, shape có đọc được không? Nếu tắt bloom, visual có sụp không? Nếu chuyển sang mobile, bạn sẽ bỏ layer nào trước? Nếu lead nói organic hơn, bạn sẽ sửa lane nào?",
                 "Những câu hỏi này giúp bạn rời khỏi tư duy làm cho xong. Bạn bắt đầu nghĩ như người chịu trách nhiệm với effect trong production. Đây là ranh giới giữa artist biết node và Technical VFX Artist biết hệ thống.",
                 "Không có con đường tắt để master shader VFX. Nhưng có con đường rõ: nhìn bằng dữ liệu, học toán theo visual, hiểu cơ bắp render, viết được shader khi cần, và luôn kiểm tra effect trong scene thật."
             ]),
        ],
        "notes": [
            ("Ghi nhớ", "Mỗi material học được nên để lại ba thứ: một screenshot debug, một câu giải thích visual, và một ghi chú cost.", False),
        ],
        "case": [
            "Một bài tập tốt không cần hoành tráng. Một radial mask có Power và SmoothStep đủ để bạn học falloff. Một dissolve grayscale đủ để bạn học threshold. Một SubUV NumberGrid đủ để bạn hiểu atlas. Nếu bài tập quá lớn, bạn sẽ nhớ cảm giác hoàn thành nhưng quên mất khái niệm.",
            "Hãy giữ mỗi bài nhỏ tới mức bạn có thể làm lại từ đầu trong 30 phút. Khi bạn làm lại được, giải thích được và phá được nó, khái niệm đã bắt đầu thành phản xạ."
        ],
        "code": [
            "Weekly loop:",
            "Read one concept",
            "Build one tiny material",
            "Break it intentionally",
            "Fix it",
            "Explain it in six lanes"
        ],
    },
]


def chapter(doc, idx, data):
    h1(doc, data["title"], idx)
    if data.get("image"):
        figure(doc, data["image"][0], data["image"][1])
    for title, paras in data["sections"]:
        h2(doc, title)
        for text in paras:
            body(doc, text)
    for title, text, warning in data.get("notes", []):
        note(doc, title, text, warning)
    if data.get("case"):
        h2(doc, "Case đọc chậm")
        for text in data["case"]:
            body(doc, text)
    if data.get("code"):
        code(doc, data["code"], "Sketch logic")


def closing(doc):
    h1(doc, "Kết luận - viết lại bằng chính mắt của bạn")
    body(doc, "Một người học material thường muốn đi nhanh tới kết quả cuối cùng. Điều đó tự nhiên, vì VFX là nghề của hình ảnh. Nhưng càng làm lâu, bạn càng thấy tốc độ thật không đến từ việc nối node nhanh hơn. Tốc độ thật đến từ việc nhìn đúng vấn đề ngay từ đầu. Trước khi thêm node, hãy nhìn effect như dữ liệu. Trước khi tăng emissive, hãy xem hierarchy. Trước khi gọi một effect là xấu, hãy hỏi nó đang hỏng ở shape, mask, motion, color, interaction hay cost.")
    body(doc, "Một cuốn sách material VFX tốt không thể chỉ là danh sách công thức. Công thức cần thiết, nhưng công thức phải nằm trong câu chuyện của visual và production. Người học cần biết vì sao một trick tồn tại, nó giải quyết vấn đề gì, nó hỏng trong trường hợp nào, và khi nào nên bỏ nó để dùng một giải pháp rẻ hơn.")
    body(doc, "Con đường học shader VFX nên được xây bằng nhiều bài nhỏ nhưng thật. Dựng một mask grayscale. Làm một edge band. Đọc một atlas bằng NumberGrid. Tạo một flare additive rồi kiểm tra nó trên nền sáng. Làm một fire card có DepthFade. Viết một function HLSL nhỏ rồi so với graph node. Mỗi bài tập như vậy để lại một phần phản xạ. Khi phản xạ đủ nhiều, bạn không chỉ biết làm một hiệu ứng cụ thể; bạn biết cách nghĩ khi gặp một hiệu ứng mới.")


def references(doc):
    h1(doc, "Nguồn tham khảo và hướng đọc tiếp")
    body(doc, "Các nguồn dưới đây được dùng như định hướng học, không phải để trích dài. Khi đọc tiếp, bạn nên đọc theo vấn đề đang gặp trong production thay vì đọc dàn trải tất cả cùng lúc.")
    refs = [
        "The Book of Shaders - cách tiếp cận shader qua pixel, tọa độ, shaping function, pattern và noise.",
        "Unity URP Shader/VFX e-book - hướng đọc Shader Graph, URP, VFX Graph, Render Graph, mobile và stylized rendering.",
        "Unreal Engine Material documentation - Material Editor, parameter, blend mode, Custom Material Expressions và HLSL.",
        "Unreal Niagara documentation - renderer, Particle Color, Dynamic Parameter, SubUV và dữ liệu runtime.",
        "Real-Time Rendering - nền tảng graphics pipeline, rasterization, texture, shading và GPU thinking.",
        "VFXSamplePack / ForgePilot Knowledge - nguồn thực hành trong project hiện tại."
    ]
    for ref in refs:
        bullet(doc, ref)


def build():
    doc = setup()
    cover(doc)
    preface(doc)
    for i, ch in enumerate(CHAPTERS, 1):
        chapter(doc, i, ch)
    closing(doc)
    references(doc)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
