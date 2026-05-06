from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path.home() / "Downloads" / "Giao_Trinh_Material_VFX_Realtime_V21_AsymmetricGrid.docx"
IMAGE_DIR = Path("Saved") / "material_vfx_v5_images"

INK = RGBColor(24, 24, 24)
MUTED = RGBColor(92, 92, 92)
BLUE = RGBColor(22, 58, 88)
ORANGE = RGBColor(152, 78, 30)
BODY = "Times New Roman"
HEAD = "Arial"
SANS = "Arial"
CODE = "Consolas"


def set_font(run, size=10.6, name=BODY, bold=False, italic=False, color=INK):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def shade(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def border_bottom(paragraph, color="D6D6D6", size="4"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), color)


def keep_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    if p_pr.find(qn("w:keepNext")) is None:
        p_pr.append(OxmlElement("w:keepNext"))


def clear_cell(cell):
    tc = cell._tc
    for child in list(tc):
        tc.remove(child)


def set_cell_margins(cell, top=120, start=120, bottom=120, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, val="single", color="DDE3EA", size="2"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            borders.append(node)
        node.set(qn("w:val"), val)
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def cm_to_twips(value):
    return str(int(value * 567))


def set_column_widths(table, widths_cm):
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_cm:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), cm_to_twips(width))
        grid.append(col)
    for col_idx, width in enumerate(widths_cm):
        for cell in table.columns[col_idx].cells:
            cell.width = Cm(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), cm_to_twips(width))
            tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table, color="DDE3EA"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "2")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def p(doc, text="", size=10.7, before=0, after=6.2, font=BODY, bold=False, italic=False, color=INK, align=None, style=None):
    para = doc.add_paragraph()
    if style:
        para.style = style
    para.paragraph_format.space_before = Pt(before)
    para.paragraph_format.space_after = Pt(after)
    para.paragraph_format.line_spacing = 1.33
    if align is not None:
        para.alignment = align
    if text:
        run = para.add_run(text)
        set_font(run, size=size, name=font, bold=bold, italic=italic, color=color)
    return para


def body(doc, text):
    return p(doc, text)


def h1(doc, title, number=None):
    doc.add_page_break()
    if number is not None:
        p(doc, f"Chương {number}", size=9.5, font=SANS, bold=True, color=ORANGE, before=18, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
    head = p(doc, title, size=25, font=HEAD, bold=True, color=BLUE, after=8, align=WD_ALIGN_PARAGRAPH.CENTER, style="Heading 1")
    head.paragraph_format.line_spacing = 1.05
    border_bottom(head, color="C98E54", size="8")
    return head


def h2(doc, title):
    para = p(doc, title, size=14.0, font=HEAD, bold=True, color=BLUE, before=12, after=4, style="Heading 2")
    keep_next(para)
    return para


def h3(doc, title):
    para = p(doc, title, size=10.4, font=SANS, bold=True, color=INK, before=8, after=3)
    keep_next(para)
    return para


def note(doc, title, text, warning=False):
    fill = "F5F2EC" if not warning else "F8EDE6"
    color = BLUE if not warning else ORANGE
    head = p(doc, title, size=9.3, font=SANS, bold=True, color=color, before=7, after=1.5)
    head.paragraph_format.left_indent = Cm(0.45)
    head.paragraph_format.right_indent = Cm(0.45)
    shade(head, fill)
    para = p(doc, text, size=9.6, font=BODY, before=0, after=6)
    para.paragraph_format.left_indent = Cm(0.45)
    para.paragraph_format.right_indent = Cm(0.45)
    shade(para, fill)


def code(doc, lines, caption=None):
    if caption:
        p(doc, caption, size=8.8, font=SANS, italic=True, color=MUTED, after=1.5)
    for line in lines:
        para = p(doc, line, size=8.7, font=CODE, before=0, after=0.5)
        para.paragraph_format.left_indent = Cm(0.5)
        para.paragraph_format.right_indent = Cm(0.5)
        para.paragraph_format.line_spacing = 1.05
        shade(para, "ECEFF2")


def figure(doc, filename, caption, width=13.0):
    path = IMAGE_DIR / filename
    if not path.exists():
        return
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(8)
    para.paragraph_format.space_after = Pt(2)
    para.add_run().add_picture(str(path), width=Cm(width))
    cap = p(doc, caption, size=8.5, font=SANS, italic=True, color=MUTED, after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    cap.paragraph_format.line_spacing = 1.12


def figure_cell(cell, filename, caption=None, width=4.6):
    path = IMAGE_DIR / filename
    if not path.exists():
        return
    para = cell.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_after = Pt(3)
    para.add_run().add_picture(str(path), width=Cm(width))
    if caption:
        cap = p(cell, caption, size=7.2, font=SANS, italic=True, color=MUTED, after=6, align=WD_ALIGN_PARAGRAPH.CENTER)
        cap.paragraph_format.line_spacing = 1.05


def bullet(doc, text):
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Cm(0.6)
    para.paragraph_format.first_line_indent = Cm(-0.24)
    para.paragraph_format.space_after = Pt(3.2)
    para.paragraph_format.line_spacing = 1.18
    run = para.add_run("- " + text)
    set_font(run, size=10.0, name=BODY)


def term_hits(text):
    terms = [
        "Overdraw", "Shader Complexity", "SubUV", "NumberGrid", "DepthFade", "Fresnel",
        "WPO", "VAT", "HLSL", "GLSL", "SDF", "Polar UV", "Particle Color",
        "Dynamic Parameter", "Blend Mode", "Texture", "Mask", "UV", "Atlas",
        "Power", "SmoothStep", "Niagara", "Material Instance", "Opacity",
        "Emissive", "Instruction", "Frame Rate", "Frame Interpolation"
    ]
    found = []
    low = text.lower()
    for term in terms:
        if term.lower() in low and term not in found:
            found.append(term)
    return found[:7]


def margin_notes(title, paras):
    joined = " ".join(paras)
    notes = []
    terms = term_hits(title + " " + joined)
    if terms:
        notes.append("Từ khóa: " + ", ".join(terms))
    low = joined.lower()
    if "debug" in low or "kiểm tra" in low:
        notes.append("Debug: cô lập dữ liệu trước khi chỉnh màu hoặc tăng cường độ.")
    if "overdraw" in low or "shader complexity" in low or "mobile" in low:
        notes.append("Performance: nhìn diện tích pixel và số lớp Translucent trước Instruction.")
    if "subuv" in low or "atlas" in low or "flipbook" in low:
        notes.append("Atlas: phân biệt Atlas UV và local UV trong frame hiện tại.")
    if "hlsl" in low or "custom node" in low:
        notes.append("Code: Custom Node chỉ tốt khi input/output và debug view rõ ràng.")
    if "fresnel" in low or "normal" in low or "space" in low:
        notes.append("Space: luôn hỏi dữ liệu đang ở Object, World, View hay Tangent Space.")
    if "mask" in low or "grayscale" in low:
        notes.append("Visual check: tắt Bloom, xem Mask Grayscale trước.")
    if not notes:
        notes.append("Câu hỏi chính: dữ liệu nào đi vào, biến đổi gì, xuất ra pixel nào?")
    return notes[:5]


def grid_block(doc, title, paras, margin_title="Margin", margin_items=None, image=None, code_lines=None):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    set_column_widths(table, [12.35, 4.95])
    main, side = table.rows[0].cells
    main.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    side.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    clear_cell(main)
    clear_cell(side)
    set_cell_margins(main, top=80, start=80, bottom=130, end=220)
    set_cell_margins(side, top=130, start=170, bottom=130, end=130)
    set_cell_borders(main)
    set_cell_borders(side)

    head = p(main, title, size=13.2, font=HEAD, bold=True, color=BLUE, before=4, after=5)
    keep_next(head)
    for text in paras:
        para = p(main, text, size=10.35, font=BODY, before=0, after=5.6)
        para.paragraph_format.line_spacing = 1.27

    if image:
        figure_cell(side, image[0], image[1], width=4.2)
    side_head = p(side, margin_title, size=8.7, font=SANS, bold=True, color=ORANGE, before=0, after=4)
    border_bottom(side_head, color="D8DEE4", size="3")
    for item in (margin_items or []):
        para = p(side, "• " + item, size=7.6, font=SANS, color=INK, before=0, after=3)
        para.paragraph_format.line_spacing = 1.08
    if code_lines:
        p(side, "Sketch", size=7.8, font=SANS, bold=True, color=BLUE, before=5, after=2)
        for line in code_lines[:8]:
            para = p(side, line, size=6.8, font=CODE, color=INK, before=0, after=1)
            para.paragraph_format.line_spacing = 1.0
            shade(para, "E9EEF2")
    p(doc, "", size=1, before=0, after=2)


def setup():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.55)
    sec.bottom_margin = Cm(1.35)
    sec.left_margin = Cm(1.65)
    sec.right_margin = Cm(1.35)
    sec.header_distance = Cm(0.8)
    sec.footer_distance = Cm(0.75)
    styles = doc.styles
    styles["Normal"].font.name = BODY
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), BODY)
    styles["Normal"].font.size = Pt(10.7)
    styles["Normal"].paragraph_format.line_spacing = 1.33
    styles["Normal"].paragraph_format.space_after = Pt(6.2)
    footer = sec.footer.paragraphs[0]
    footer.text = "Giáo trình Material VFX Realtime"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        set_font(run, size=8.0, name=SANS, color=MUTED)
    return doc


def cover(doc):
    p(doc, "Giáo trình", size=12, font=SANS, bold=True, color=ORANGE, before=74, after=5, align=WD_ALIGN_PARAGRAPH.CENTER)
    title = p(doc, "Material VFX Realtime", size=32, font=HEAD, bold=True, color=BLUE, after=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    title.paragraph_format.line_spacing = 1.0
    p(doc, "Tư duy Material, Shader và Production cho Technical VFX Artist", size=12.2, font=BODY, italic=True, color=MUTED, after=26, align=WD_ALIGN_PARAGRAPH.CENTER)

    rule = p(doc, "", before=0, after=22)
    border_bottom(rule, color="C98E54", size="10")

    block = p(doc, "Shape · Mask · Motion · Color · Interaction · Cost", size=14.2, font=SANS, bold=True, color=BLUE, before=8, after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    block.paragraph_format.line_spacing = 1.05
    p(doc, "Unreal Material · Niagara · Shader Graph · HLSL", size=11.0, font=SANS, color=INK, after=26, align=WD_ALIGN_PARAGRAPH.CENTER)

    summary = p(
        doc,
        "Một tài liệu học Material VFX theo cách đọc dữ liệu: bắt đầu từ hình ảnh, tách thành các lớp kỹ thuật, rồi đưa từng quyết định trở lại production.",
        size=10.4,
        font=BODY,
        color=INK,
        before=16,
        after=26,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    summary.paragraph_format.left_indent = Cm(1.15)
    summary.paragraph_format.right_indent = Cm(1.15)
    p(doc, "Phiên bản V21 · Asymmetric Grid · 2026", size=9.3, font=SANS, color=MUTED, before=8, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)


def preface(doc):
    h1(doc, "Lời nói đầu")
    body(doc, "Một hiệu ứng thời gian thực tồn tại trong một khoảng thời gian rất ngắn, nhưng nó được tạo ra bởi rất nhiều quyết định. Một vệt lửa nhỏ có thể chứa Texture, UV, Mask, Particle Color, Blend Mode, Depth, Exposure, Bloom và cả giới hạn của GPU. Người xem chỉ thấy nó cháy lên rồi biến mất. Người làm Technical VFX phải thấy được hệ thống phía sau khoảnh khắc đó.")
    body(doc, "Material VFX vì vậy không nên được học như một danh sách node rời rạc. Một node chỉ có ý nghĩa khi nó biến đổi đúng loại dữ liệu để phục vụ đúng mục tiêu thị giác. Power có thể làm Flare sắc hơn, nhưng cũng có thể làm Smoke mất độ mềm. SmoothStep có thể tạo dải biên Dissolve sạch, nhưng cũng có thể làm Mask bị giả nếu nguồn dữ liệu không tốt. Fresnel có thể tạo Rim đẹp, nhưng cũng có thể biến Shield thành một quả cầu nhựa phát sáng nếu thiếu Surface Breakup và Hit Response.")
    body(doc, "Cuốn sách này đi từ cách nhìn tới cách làm. Trước hết, ta học cách nhìn hình ảnh như dữ liệu. Sau đó, ta dùng toán Shader để nắn dữ liệu thành Shape, Mask, Motion và Color. Tiếp theo, ta đặt Material vào bối cảnh thật của Realtime Rendering: CPU, GPU, Overdraw, băng thông Texture, Blend Mode và Post-process. Cuối cùng, ta quay lại production, nơi mỗi quyết định phải trả lời ba câu hỏi: Effect có đọc được không, có tinh chỉnh được không, và có sống được trong ngân sách hiệu năng không.")
    body(doc, "Mục tiêu không phải là biến người học thành người thuộc nhiều công thức nhất. Mục tiêu là giúp người học mở một Material bất kỳ, chia nó thành các lớp dễ hiểu, tìm được nguyên nhân khi visual bị hỏng, và giải thích được quyết định của mình bằng ngôn ngữ mà Artist, Technical Artist và Programmer đều có thể trao đổi.")
    note(doc, "Ghi nhớ", "Một Material tốt không chỉ tạo ra hình ảnh đẹp trong khung Preview. Nó phải giữ được ý đồ thị giác trong Scene thật, dưới camera thật, với Post-process thật và giới hạn phần cứng thật.")


CHAPTERS = [
    {
        "title": "Digital Eyes - tập nhìn hình ảnh như dữ liệu",
        "image": ("v9_shader_lane_diagram.png", "Một Effect có thể được đọc theo các làn dữ liệu: UV, Mask, Color, Opacity, Interaction và Cost."),
        "sections": [
            ("Tại sao mắt của Technical VFX Artist phải khác mắt của người xem",
             [
                 "Người xem bình thường nhìn một vệt lửa và nói rằng nó đẹp, yếu, bẩn, mạnh, giả hoặc hấp dẫn. Technical VFX Artist phải nhìn cùng vệt lửa đó và hỏi thêm một lớp câu hỏi khác. Vệt sáng này đến từ Texture hay Procedural Mask? Nó đang cộng sáng bằng Additive hay đang Alpha Blend? Biên (Edge) của nó sắc vì Texture, vì Power, hay vì Alpha bị nén Texture? Nó bị mất trên nền sáng vì thiếu Contrast hay vì Bloom đang nuốt mất Hierarchy?",
                 "Đó là lý do tôi gọi chương này là Digital Eyes. Bạn không chỉ nhìn bằng mắt mỹ thuật. Bạn nhìn bằng mắt dữ liệu. Một pixel trên màn hình không phải là màu cuối cùng của riêng Material. Nó là kết quả của một chuỗi quyết định: Texture Import, UV, Particle Color, Material Graph, Blend Mode, Depth, Exposure, Bloom, Tonemapper và background. Khi bạn hiểu điều này, bạn sẽ bớt sửa Effect bằng cảm giác mơ hồ và bắt đầu sửa bằng nguyên nhân.",
                 "Kỹ năng này rất giống việc nghe một bản nhạc và tách được bass, trống, melody, vocal. Người mới nghe một khối âm thanh. Người làm nghề nghe từng lớp. Với Material cũng vậy. Người mới thấy một hiệu ứng phát sáng. Người có Digital Eyes thấy Mask nào đang sống, Motion nào đang chạy, Color nào đang quá mạnh, Opacity nào đang làm bẩn Silhouette, và Cost nào có thể khiến GPU trả giá."
             ]),
            ("Từ Visual sang dữ liệu",
             [
                 "Một cách luyện đơn giản là lấy bất kỳ Effect nào và tạm thời cấm bản thân dùng từ đẹp hoặc xấu. Thay vào đó, hãy mô tả nó bằng dữ liệu. Shape của Effect có đọc được khi thu nhỏ không? Mask có vùng đen trắng rõ không? Motion có cùng nhịp hay bị trượt nhau? Color có phân cấp vùng lõi (Core), dải biên (Edge) và vầng sáng ngoài (Outer Glow) không? Interaction với scene có mềm hay bị cắt? Cost có nằm ở diện tích pixel và lượt đọc Texture hay số lượng particle?",
                 "Khi bạn đặt câu hỏi như vậy, Material không còn là một Node Graph bí ẩn. Nó trở thành một hệ thống nhỏ. Hệ thống đó nhận dữ liệu, biến đổi dữ liệu, và xuất ra pixel. Điều này nghe đơn giản, nhưng nó thay đổi cách bạn học. Bạn sẽ không hỏi 'node này dùng làm gì' như một câu hỏi độc lập nữa. Bạn sẽ hỏi 'node này đang biến đổi làn dữ liệu nào, và visual nhận được thay đổi gì'.",
                 "Ví dụ, Power không phải là một node để làm mọi thứ sắc hơn. Power là một cách uốn curve 0-1. Khi áp vào Radial Mask, nó bóp ánh sáng vào vùng lõi (Core). Khi áp vào Alpha Smoke, nó có thể làm Smoke mất độ mềm. Khi áp vào Fresnel, nó làm rim hẹp lại. Cùng một node, ba hậu quả khác nhau, vì dữ liệu đầu vào và mục tiêu thị giác khác nhau."
             ]),
            ("Debug bằng grayscale",
             [
                 "Nếu chỉ có một thói quen cần lấy từ chương này, đó là Solo Grayscale. Hãy tắt màu, tắt bloom, tắt những thứ làm bạn thấy phấn khích, rồi nhìn Mask đen trắng. Mask đó có kể được câu chuyện chưa? Vùng nào sống, vùng nào chết, vùng nào nằm ở vùng biên (Edge), vùng nào bị noise phá, vùng nào quá đều? Nếu Grayscale đã sai, bản màu chỉ đang trang điểm cho lỗi.",
                 "Trong production, grayscale còn giúp bạn giao tiếp tốt hơn. Khi Lead nói Dissolve chưa organic, bạn có thể đưa ra mask nguồn (Source Mask) và dải biên (Edge Band) để chỉ rằng vấn đề nằm ở Silhouette Breakup chứ không phải ở màu vùng biên (Edge). Khi Programmer hỏi vì sao Effect đắt, bạn có thể chỉ card Alpha quá rộng thay vì nói chung chung rằng shader nặng.",
                 "Thói quen này cũng giúp bạn chuyển engine. Unreal, Unity, Custom Engine hay WebGL có tên node khác nhau, nhưng Mask Grayscale thì ở đâu cũng là Mask Grayscale. Một vùng 0-1 tốt sẽ sống được qua nhiều pipeline."
             ]),
        ],
        "notes": [
            ("Ghi nhớ", "Đừng để màu và bloom quyết định quá sớm. Một Material VFX tốt thường có grayscale tốt trước khi nó có màu đẹp.", False),
        ],
        "case": [
            "Trong một buổi review thực tế, feedback về Flare hoặc Dissolve thường không đến dưới dạng công thức. Lead có thể nói rằng Effect nhìn yếu, nhìn bẩn hoặc chưa có lực. Nếu bạn chỉ nghe phần cảm giác, bạn sẽ dễ tăng Intensity hoặc thêm Noise. Nếu bạn nghe bằng Digital Eyes, bạn sẽ kiểm tra Grayscale trước: Shape có đọc không, Mask có bị phẳng không, Alpha có quá rộng không, và Bloom có đang che lỗi không.",
            "Cách đọc này làm cuộc nói chuyện chuyên nghiệp hơn. Bạn không cần tranh luận cảm giác ngay lập tức. Bạn có thể chỉ vào mask nguồn (Source Mask), dải biên (Edge Band) hoặc card Alpha và nói: vấn đề đang nằm ở đây. Khi một feedback mơ hồ được kéo về một lớp cụ thể, việc sửa effect trở nên nhanh hơn rất nhiều."
        ],
        "code": [
            "rawMask    = TextureMask(UV)",
            "visible    = smoothstep(Threshold, Threshold + Softness, rawMask)",
            "edge       = smoothstep(Threshold - Width, Threshold, rawMask) - visible",
            "finalColor = BaseColor * visible + EdgeColor * edge"
        ],
    },
    {
        "title": "The Matrix of Art - toán học như cách nặn cảm giác",
        "image": ("v9_power_shaping_flare.png", "Power không chỉ là phép toán. Nó thay đổi cảm giác falloff của một vùng sáng."),
        "sections": [
            ("Toán trong shader không bắt đầu từ nỗi sợ",
             [
                 "Nhiều artist nghe tới toán Shader là thấy xa lạ. Điều này dễ hiểu, vì toán thường được dạy như bài kiểm tra, không phải như công cụ tạo hình. Trong Material VFX, toán nên được học theo cảm giác trước. SmoothStep là một lưỡi dao mềm. Power là bàn tay bóp ánh sáng vào vùng lõi (Core). Dot Product là câu hỏi hai hướng có đang nhìn cùng phía không. Distance là cách biến không gian thành shape.",
                 "Khi bạn học toán theo cách đó, công thức không còn đứng riêng. Nó gắn với một hành động thị giác. Remap đổi thang đo để một mask nhạt trở thành mask rõ. Clamp giữ giá trị không vượt khỏi vùng an toàn. Lerp pha hai trạng thái. Sine tạo nhịp thở. Atan2 mở cánh cửa vào Polar UV, nơi Portal, Vortex và Radial Slash bắt đầu có ngôn ngữ riêng.",
                 "Một Technical VFX Artist không cần biến mình thành nhà toán học thuần túy để làm shader tốt. Nhưng bạn cần đủ thân với các phép toán cơ bản để không sợ Graph. Khi feedback tới, bạn cần biết mình đang thiếu contrast, thiếu độ rộng dải biên (Edge Width), thiếu falloff, sai Space hay sai Curve."
             ]),
            ("Từ một giá trị 0-1 tới một visual đọc được",
             [
                 "Rất nhiều Material VFX bắt đầu từ một giá trị 0-1. Một mẫu Noise trả về vùng xám. Một Radial Gradient trả về khoảng cách từ tâm. Một Fresnel trả về độ nghiêng giữa Normal và hướng nhìn. Một Depth Comparison trả về khoảng cách tới Geometry. Bản thân các giá trị này chưa phải effect. Effect xuất hiện khi bạn remap, clamp, Threshold, làm mềm, nhân màu, đưa vào Opacity hoặc Emissive.",
                 "Hãy tưởng tượng một dissolve. Mask nguồn (Source Mask) là một bản đồ địa hình. Threshold là mực nước đang dâng. Vùng thấp biến mất trước, vùng cao còn lại sau. Dải biên (Edge Band) là bờ nước, nơi bạn có thể cho ánh sáng cháy lên. Nếu bạn hiểu hình ảnh đó, công thức không còn khô nữa. Nó chỉ là cách viết chính xác cho một câu chuyện rất trực quan.",
                 "Điểm quan trọng là Edge phải đến từ cùng Mask nguồn (Source Mask) với Visible Mask. Nếu bạn vẽ edge bằng một texture khác hoặc một outline tùy tiện, dissolve sẽ không bám logic. Người xem có thể không biết công thức, nhưng mắt họ sẽ thấy có gì đó giả."
             ]),
            ("SDF, Polar và những Shape không cần Texture",
             [
                 "Signed Distance Field là một ví dụ đẹp về việc toán có thể vẽ hình. Thay vì lưu sẵn một vòng tròn trong Texture, bạn tính khoảng cách từ pixel tới tâm. Khoảng cách đó trở thành dữ liệu. Từ dữ liệu đó, bạn tạo Circle, Ring, Soft Edge, Outline hoặc Shockwave. Khi cần Portal, bạn thêm angle bằng atan2. Khi cần Motion xoáy, bạn offset angle theo Time hoặc Radius.",
                 "Điều này không có nghĩa là texture không quan trọng. Texture vẫn rất mạnh, đặc biệt khi bạn cần chi tiết organic, Shape vẽ tay (hand-painted) hoặc baked simulation. Nhưng khi bạn biết dựng shape bằng toán, bạn có thêm một lớp tự do. Bạn có thể làm prototype nhanh, tạo parameter sạch, và chỉ đưa texture vào khi texture thật sự mang thông tin mà toán không nên gánh.",
                 "Trong production, sự cân bằng giữa procedural và texture là một quyết định. Procedural dễ tinh chỉnh và nhẹ bộ nhớ, nhưng có thể tốn Instruction hoặc nhìn quá sạch. Texture giàu chi tiết, nhưng tốn băng thông Texture và dễ bị Compression/Mip phá. Người giỏi không chọn một phe; họ chọn theo Effect."
             ]),
        ],
        "notes": [
            ("Cảnh báo production", "Đừng biến mọi thứ thành procedural chỉ vì nó trông thông minh. Nếu một flipbook hoặc texture nhỏ giải quyết visual tốt hơn và rẻ hơn, đó là lựa chọn production hợp lý.", True),
        ],
        "case": [
            "Một Portal không dùng Texture là bài tập rất tốt, nhưng nó cũng dễ tạo ảo giác rằng procedural luôn tốt hơn. Trong prototype, bạn có thể dựng Ring bằng Distance, thêm Polar UV, rồi làm dải biên (Edge) bằng SmoothStep. Kết quả rất sạch và dễ tinh chỉnh. Nhưng khi art direction cần vết nứt, ký tự cổ, vân năng lượng hoặc cảm giác hand-painted, texture có thể là dữ liệu đúng hơn.",
            "Quyết định production không phải là Procedural hay Texture. Quyết định đúng là phần nào nên được tính, phần nào nên được vẽ, và phần nào nên được Bake. Một Technical VFX Artist giỏi biết chuyển qua lại giữa ba cách đó mà không xem cách nào là tôn giáo."
        ],
        "code": [
            "r     = length(UV - Center)",
            "ring  = 1 - smoothstep(Width, Width + Softness, abs(r - Radius))",
            "edge  = ring * EdgeIntensity",
            "color = ring * Tint"
        ],
    },
    {
        "title": "Cơ bắp của realtime - CPU, GPU và cái giá của pixel",
        "image": ("v9_shader_complexity_overdraw.png", "Overdraw là một trong những chi phí lớn nhất của translucent VFX."),
        "sections": [
            ("Một Effect không chạy trong sách vở",
             [
                 "Một material đẹp trong Editor vẫn có thể là một vấn đề trong game. Lý do rất đơn giản: game không render một Material trong chân không. Game render hàng nghìn object, ánh sáng, shadow, Post-process, UI, Particle, animation và gameplay logic. CPU phải chuẩn bị scene. GPU phải vẽ scene. Mỗi frame chỉ có một ngân sách thời gian rất nhỏ.",
                 "CPU thường lo gameplay, transform, culling, submit draw call, sort translucent object, update Particle hoặc gửi parameter. GPU nhận công việc và xử lý theo pipeline: vertex, rasterization, fragment/pixel, blending và các pass sau đó. Khi bạn làm VFX, bạn thường chạm mạnh vào Pixel Shader, blending và overdraw. Đây là vùng rất dễ đắt, nhất là với Fire, Smoke, Glow, Aura và UI effect có alpha rộng.",
                 "Điều đáng chú ý là một shader có Instruction Count không quá cao vẫn có thể đắt nếu nó phủ quá nhiều pixel. Một card Smoke lớn, alpha rộng và chồng nhiều lớp có thể làm GPU vẽ đi vẽ lại cùng một vùng màn hình. Đây là lý do overdraw thường quan trọng hơn vài phép Multiply nhỏ."
             ]),
            ("Translucent VFX và bài toán diện tích",
             [
                 "Opaque Mesh thường ghi Depth và giúp GPU loại bỏ nhiều pixel phía sau. Translucent VFX lại phức tạp hơn, vì nó cần trộn với background và thường không ghi Depth theo cách opaque. Nếu bạn có mười card Smoke chồng lên nhau, một pixel có thể bị shader xử lý nhiều lần. Người xem chỉ thấy khói mềm; GPU thấy một hóa đơn nhiều lớp.",
                 "Khi tối ưu VFX, nhiều người mới nhìn Instruction Count trước. Điều đó không sai, nhưng chưa đủ. Bạn cần hỏi effect chiếm bao nhiêu phần trăm màn hình, có bao nhiêu layer chồng nhau, card Alpha có bị rỗng quá nhiều không, Texture có quá lớn so với kích thước hiển thị không, và effect xuất hiện bao nhiêu lần trong một frame.",
                 "Với mobile hoặc iGaming, những câu hỏi này trở nên thực tế hơn nữa. Một sparkle nhỏ có thể rẻ khi xuất hiện một lần, nhưng đắt khi có hàng trăm instance. Một glow đẹp có thể ổn trên PC, nhưng trên web/mobile nó trở thành lý do khiến animation giật."
             ]),
            ("Tối ưu không có nghĩa là làm xấu đi",
             [
                 "Tối ưu tốt không phải là xóa bừa node. Tối ưu tốt là giữ phần người xem thật sự đọc được và bỏ phần họ không cần. Nếu một projectile cần hướng bay rõ, hãy giữ Core, Head, Tail và timing. Nếu noise nhỏ trên shell không đọc được ở camera gameplay, bạn có thể giảm nó hoặc bake nó. Nếu glow phủ quá rộng, crop Alpha và giảm diện tích card có thể hiệu quả hơn việc tranh luận thêm bớt vài instruction.",
                 "Một nguyên tắc thực dụng là tối ưu theo thứ tự: Silhouette, diện tích, lớp, lượt đọc Texture, Instruction. Silhouette quyết định effect có đọc được không. Area và layer quyết định overdraw. Texture fetch quyết định bandwidth. Instruction quyết định độ nặng tính toán. Thứ tự này không tuyệt đối, nhưng nó giúp bạn tránh tối ưu nhầm chỗ.",
                 "Technical VFX Artist giỏi thường biết nói câu này: bản đẹp nhất chưa chắc là bản production tốt nhất. Bản production tốt là bản đạt visual, giữ gameplay readable, dễ tinh chỉnh, và sống được trong budget."
             ]),
        ],
        "notes": [
            ("Ghi nhớ", "Nếu effect là translucent và chiếm nhiều màn hình, hãy nghĩ tới Overdraw trước khi nghĩ tới vài phép toán nhỏ trong Graph.", False),
        ],
        "case": [
            "Một Win Effect trong iGaming có thể trông rất nhỏ trong layout UI, nhưng nó xuất hiện nhiều lần, trên nhiều thiết bị, và thường đi cùng animation khác. Nếu mỗi sparkle là một Translucent Quad rộng, nếu mỗi Glow dùng Texture lớn, và nếu mọi layer đều additive, chi phí cộng dồn sẽ đến nhanh hơn cảm giác của artist.",
            "Trong trường hợp này, tối ưu tốt thường không bắt đầu bằng HLSL. Nó bắt đầu bằng asset authoring: crop Alpha sát hơn, pack mask, bake motion phụ, giảm layer runtime và giới hạn số particle cùng lúc. Shader thông minh không cứu được một thiết kế asset lãng phí."
        ],
        "code": [
            "ApproxCost = PixelArea * LayerCount * ShaderCost",
            "If PixelArea is large:",
            "    crop Alpha, reduce card size, reduce overlapping layers",
            "If TextureFetch is high:",
            "    pack masks, lower resolution, check mip/compression"
        ],
    },
    {
        "title": "Ngôn ngữ Shader - khi Node Graph không còn đủ",
        "image": ("v11_realtime_contract_simple.png", "Material tốt cần hợp đồng dữ liệu rõ với Niagara, Blueprint hoặc hệ runtime tương đương."),
        "sections": [
            ("Code shader không phải nghi thức thần bí",
             [
                 "Có một giai đoạn trong hành trình học Material khi Node Graph bắt đầu giống một mê cung. Lúc đó, nhiều artist nghĩ HLSL hoặc GLSL là một thế giới khác, chỉ dành cho programmer. Thực tế bình tĩnh hơn. Shader code cũng chỉ nhận dữ liệu, biến đổi dữ liệu và trả kết quả. Điều khác biệt là code buộc bạn viết rõ ý hơn.",
                 "Một fragment shader tối giản có thể chỉ trả về một màu. Từ đó, bạn thêm UV, sample texture, nhân tint, tính alpha, thêm time, thêm Noise, thêm normal, thêm depth. Mỗi bước vẫn là một câu rất nhỏ. Nếu bạn đọc được câu nhỏ đó trong code, bạn sẽ đọc Node Graph tốt hơn, vì bạn hiểu ý tưởng nằm dưới hình dạng node.",
                 "Điểm quan trọng là không dùng HLSL để tỏ ra senior. Một Node Graph rõ ràng vẫn rất tốt, đặc biệt trong team art. HLSL nên xuất hiện khi Graph lặp quá nhiều, khi công thức diễn đạt bằng code sạch hơn, hoặc khi bạn cần một function nhỏ có input/output rõ."
             ]),
            ("Types, parameters và texture sample",
             [
                 "Shader language có vài từ vựng cơ bản mà bạn cần thân quen. float là một số. float2 thường là UV. float3 thường là position, normal hoặc color RGB. float4 thường là color có alpha. Texture không tự sinh ra màu; bạn phải sample texture bằng UV. Parameter là dữ liệu từ ngoài shader đưa vào, như DissolveAmount, EdgeWidth, Tint hoặc DistortionStrength.",
                 "Unreal gọi tên theo hệ material của Unreal. Unity Shader Graph có cách đóng gói riêng. GLSL có vec3, mix, fract. HLSL có float3, lerp, frac. Tên khác nhau nhưng tư duy giống nhau. Khi bạn hiểu dữ liệu, bạn không bị kẹt vào tên node.",
                 "Một shader production tốt cần input rõ. Nếu function EdgeBand nhận mask, Threshold, width và softness, nó nên trả về Visible Mask và edge mask. Nó không nên tự quyết định màu, emissive, Niagara binding và opacity cùng lúc. Function nhỏ sống lâu hơn function tham lam."
             ]),
            ("Custom node trong Unreal và giới hạn của nó",
             [
                 "Trong Unreal, Custom Material Expression cho phép viết HLSL trong material. Công cụ này hữu ích, nhưng nó không phải lối thoát cho mọi Graph rối. Nếu bạn chưa hiểu graph, chuyển nó sang code chỉ làm lỗi khó thấy hơn. Nếu bạn hiểu rõ công thức và Graph đang lặp hoặc quá dài, Custom Node có thể làm material sạch hơn.",
                 "Một ví dụ hợp lý là nhiều hit ripple trên shield. Nếu mỗi hit có position, radius, width, intensity và decay, Node Graph có thể nhanh chóng thành một tấm lưới dây nối. Một loop HLSL nhỏ với count cố định có thể dễ đọc hơn, miễn là team chấp nhận và bạn có debug output.",
                 "Khi dùng HLSL, hãy viết như đang viết cho người khác đọc. Đặt tên biến theo ý nghĩa visual. Clamp output khi cần. Ghi rõ range input. Và luôn có một cách xem kết quả trung gian bằng grayscale hoặc debug color."
             ]),
        ],
        "notes": [
            ("Cảnh báo production", "Code shader tệ nguy hiểm hơn Node Graph tệ, vì nó trông gọn trong khi che giấu quyết định. Hãy dùng HLSL khi nó làm ý tưởng rõ hơn, không phải khi nó chỉ làm Graph ngắn hơn.", True),
        ],
        "case": [
            "Một Custom Node nhỏ có thể làm Graph sạch hơn rất nhiều nếu nó thay thế một cụm logic lặp lại. Ví dụ EdgeBand, Remap01 hoặc PolarUV là các function có đầu vào và đầu ra rõ. Người khác đọc được tên function và hiểu ngay ý định.",
            "Ngược lại, một Custom Node chứa nửa material bên trong sẽ tạo ra một hộp đen. Artist không biết parameter nào quan trọng. Lead không biết debug ở đâu. Programmer không muốn maintain. Vì vậy, HLSL tốt trong Material VFX thường là HLSL nhỏ, rõ và có đường debug."
        ],
        "code": [
            "float EdgeBand(float mask, float t, float width, float soft)",
            "{",
            "    float inner = smoothstep(t - width, t, mask);",
            "    float outer = smoothstep(t, t + soft, mask);",
            "    return saturate(inner - outer);",
            "}"
        ],
    },
    {
        "title": "Mask, UV và motion - ba lớp dễ bị sửa nhầm",
        "image": ("v9_mask_solo_dissolve.png", "Một Dissolve tốt bắt đầu từ mask nguồn (Source Mask), Visible Mask và dải biên (Edge Band) rõ ràng."),
        "sections": [
            ("Mask quyết định phần nào của effect được phép tồn tại",
             [
                 "Trong rất nhiều Material VFX, mask là phần quan trọng nhất nhưng lại bị xem nhẹ nhất. Một mask tốt cho bạn biết vùng nào hiện, vùng nào mất, vùng nào là edge, vùng nào nên phát sáng, vùng nào nên distort, vùng nào nên mềm đi khi gặp depth. Khi mask sai, màu đẹp chỉ làm lỗi khó nhận ra hơn.",
                 "Một Dissolve đơn giản có thể bắt đầu bằng Noise. Nhưng Dissolve trong production hiếm khi chỉ là Noise. Nó có hướng, có tốc độ, có độ rộng dải biên (Edge Width), có Breakup, có đồng bộ với Animation hoặc Gameplay. Nếu nhân vật tan từ tay lên vai, mask nguồn (Source Mask) cần chứa thông tin hướng. Nếu Portal mở từ tâm ra ngoài, mask nguồn nên là Distance Field hoặc Radial Gradient. Nếu Shield bị hit tại một điểm, mask nguồn có thể là Distance từ World Position tới Hit Position.",
                 "Vì vậy, khi sửa effect, hãy hỏi mask đang đại diện cho điều gì. Nó là hình dạng mỹ thuật, dữ liệu gameplay, khoảng cách, độ sâu, lifetime particle hay frame index? Câu trả lời quyết định cách bạn debug."
             ]),
            ("UV Motion không phải texture đang chạy",
             [
                 "Khi dùng Panner, texture không thật sự di chuyển. Bạn đang thay đổi địa chỉ đọc texture. Câu này rất quan trọng. Nếu bạn hiểu UV là địa chỉ, bạn sẽ hiểu vì sao scale, rotate, polar, flow map và distortion đều là các cách làm lệch câu hỏi 'tôi đang đọc texture ở đâu'.",
                 "Một ngọn lửa card thường có nhiều motion. Flipbook giữ motion lớn. Panner nhẹ có thể thêm shimmer. Distortion có thể làm edge rung. Nhưng nếu mọi layer chạy cùng tốc độ, effect sẽ trông máy móc. Motion tốt cần phase, offset và Hierarchy. Core có thể nhanh hơn edge. Smoke có thể chậm hơn fire. Sparkle có thể bật theo nhịp riêng thay vì trôi đều.",
                 "Một lỗi thường gặp là dùng UV Motion để sửa vấn đề shape. Nếu silhouette của effect yếu, panner không cứu được. Nếu alpha quá rộng, rotator không làm nó rẻ hơn. Nếu mask không có hướng, thêm Noise pan có thể chỉ làm effect bẩn hơn."
             ]),
            ("Derived 0-1 UV trong SubUV",
             [
                 "Khi làm flipbook, bạn thường có Atlas UV trên toàn sheet và local UV bên trong frame hiện tại. Hai khái niệm này cần tách rõ. Atlas UV cho biết bạn đang ở đâu trên toàn texture. Derived 0-1 UV cho biết bạn đang ở đâu trong ô frame hiện tại. Nếu bạn muốn tạo radial fade trong từng frame, bạn cần local UV, không phải Atlas UV.",
                 "Đây là lý do các material như M_SubUV_Utilities đáng học. Nó không chỉ là utility. Nó dạy bạn cách renderer, Atlas và material ký hợp đồng với nhau. Niagara hoặc Sprite Renderer chọn frame. Material đọc texture. Nếu material cần logic theo từng frame, nó phải suy ra local coordinate.",
                 "NumberGrid là công cụ debug rất tốt ở đây. Khi grid sai, bạn thấy ngay frame bị lệch. Khi local UV sai, radial fade hoặc border sẽ chạy theo toàn Atlas thay vì từng ô. Đây là loại lỗi nhỏ nhưng rất dễ làm artist mất nhiều giờ nếu không biết nhìn."
             ]),
        ],
        "notes": [
            ("Ghi nhớ", "UV là địa chỉ đọc dữ liệu. Khi bạn làm méo UV, bạn không di chuyển texture; bạn thay đổi câu hỏi sample.", False),
        ],
        "case": [
            "Khi một flame card nhìn giả, nhiều người thêm Noise pan. Nhưng nếu flipbook đã có motion chính, noise pan quá mạnh có thể làm lửa trôi khỏi nhịp gốc. Người xem không gọi tên được lỗi, nhưng họ cảm thấy lửa không còn trọng lượng.",
            "Một cách debug tốt là tắt lần lượt từng motion layer. Giữ flipbook, tắt panner. Giữ panner, tắt distortion. Sau đó xem layer nào thật sự thêm cảm giác và layer nào chỉ thêm nhiễu. Không phải motion nào cũng làm effect sống hơn."
        ],
        "code": [
            "AtlasUV = TexCoord",
            "frameSize = 1 / float2(Columns, Rows)",
            "frameOffset = FrameIndexToOffset(FrameIndex)",
            "localUV = (AtlasUV - frameOffset) / frameSize"
        ],
    },
    {
        "title": "Blend, emissive và depth - nơi material gặp renderer",
        "image": ("v9_failure_depthfade_before_after.png", "DepthFade không làm smoke đẹp hơn theo nghĩa màu sắc; nó làm giao tuyến với scene bớt giả."),
        "sections": [
            ("Blend mode là quyết định render, không chỉ là mỹ thuật",
             [
                 "Additive, translucent và masked thường được chọn như lựa chọn mỹ thuật. Điều đó đúng một phần, nhưng chưa đủ. Blend mode quyết định cách pixel của material trộn với frame đã có, quyết định depth behavior, sorting và cost. Additive cộng sáng nên hợp với flare, spark, magic streak, nhưng nó dễ biến mất trên nền sáng. Translucent mềm hơn, hợp với smoke, fire card, glass hoặc energy field, nhưng dễ gặp sorting và overdraw. Masked sắc và ổn định hơn, nhưng edge có thể răng cưa nếu không xử lý tốt.",
                 "Khi một lead nói effect không đủ mạnh, nhiều người tăng emissive. Nhưng nếu Blend Mode là additive và background đã sáng, tăng emissive có thể không giải quyết readability. Bạn có thể cần tăng contrast của mask, thêm core tối hơn ở phía trong, thay đổi color Hierarchy, hoặc thêm một layer non-additive để giữ shape.",
                 "Blend mode cũng liên quan tới platform. Một effect translucent nhiều layer có thể ổn trên PC nhưng quá đắt trên mobile. Một masked/dithered variant có thể không mềm bằng, nhưng ổn định và rẻ hơn. Đây không phải chuyện đúng sai tuyệt đối; đó là tradeoff production."
             ]),
            ("Emissive cần Hierarchy",
             [
                 "Emissive rất dễ gây nghiện. Chỉ cần tăng Intensity, effect lập tức sáng hơn trong khung Preview. Nhưng emissive không có Hierarchy sẽ nhanh chóng thành một mảng cháy trắng. Một fire card tốt thường có core nóng, edge chuyển màu, outer smoke hoặc glow yếu hơn. Một magic orb tốt có core, shell, rim và accent. Nếu mọi lớp đều sáng ngang nhau, người xem mất điểm tựa.",
                 "Bloom làm vấn đề này rõ hơn. Bloom khuếch đại vùng sáng và có thể làm một vùng lõi (Core) nhỏ trở nên đẹp. Nhưng nếu bạn dựa hoàn toàn vào bloom, khi Post-process thay đổi hoặc Scene Exposure khác đi, effect sẽ đổi cảm giác mạnh. Vì vậy, hãy tắt bloom tạm thời khi debug Shape. Nếu không có bloom mà hiệu ứng không đọc được gì, bạn đang mượn bloom quá nhiều.",
                 "Một cách nghĩ thực dụng là chia emissive thành ba lớp: core đọc shape, edge đọc transition, accent đọc motion hoặc energy. Mỗi lớp nên có vai trò. Nếu một layer không có vai trò, nó có thể chỉ là noise trang trí."
             ]),
            ("DepthFade và soft particles",
             [
                 "DepthFade giải quyết một lỗi rất cụ thể: translucent card cắt thẳng vào Geometry. Smoke, fire, splash, dust và magic cloud thường cần nó. Khi không có DepthFade, giao tuyến với sàn hoặc tường nhìn như một mặt phẳng bị lộ. Khi có DepthFade, opacity giảm dần gần Geometry, tạo cảm giác mềm hơn.",
                 "Nhưng DepthFade không phải thuốc chữa mọi thứ. Nếu card quá lớn, alpha quá rộng, sorting sai hoặc smoke texture không có shape tốt, DepthFade chỉ che được một phần. Nó cũng có cost và phụ thuộc depth buffer. Bạn cần test nó với camera gần, xa, góc thấp và scene có nhiều Geometry.",
                 "Một lỗi production phổ biến là để fade distance quá lớn. Smoke bắt đầu biến mất trước khi chạm sàn, làm effect bị mỏng. Fade distance quá nhỏ thì giao tuyến vẫn cứng. Giá trị đúng phụ thuộc scale scene, kích thước card và mục tiêu thị giác."
             ]),
        ],
        "notes": [
            ("Cảnh báo production", "Đừng tăng emissive để sửa mọi feedback. Nhiều khi vấn đề nằm ở mask, Hierarchy, Blend Mode hoặc background contrast.", True),
        ],
        "case": [
            "Một Flare Additive đẹp trên nền đen có thể biến mất trên nền trời sáng. Nếu bạn chỉ tăng Intensity, nó sáng hơn trong một số cảnh nhưng vẫn thiếu shape ở cảnh khác. Vấn đề không phải chỉ là lượng sáng. Vấn đề là flare không có đủ Hierarchy để tự đọc trên nhiều background.",
            "Giải pháp có thể là thay falloff, thêm vùng lõi (Core) rõ hơn, giảm vầng sáng ngoài (Outer Glow), đổi Blend Mode cho một lớp phụ, hoặc kiểm soát bloom. Đây là lý do Blend Mode phải được xem như quyết định render chứ không phải lựa chọn màu sắc."
        ],
        "code": [
            "opacityRaw  = TextureAlpha * ParticleColor.a",
            "opacitySoft = DepthFade(opacityRaw, FadeDistance)",
            "emissive    = Core * CoreIntensity + Edge * EdgeIntensity",
            "final       = emissive * opacitySoft"
        ],
    },
    {
        "title": "Production stories - mobile, stylized, iGaming và AAA",
        "image": ("v11_imagegen_vfx_Atlas.png", "Các nhóm effect khác nhau cần quyết định khác nhau. Một công thức không phục vụ mọi style."),
        "sections": [
            ("Mobile không tha thứ cho sự hào nhoáng không cần thiết",
             [
                 "Mobile VFX có một sự thật hơi lạnh lùng: nếu người chơi không thấy chi tiết, GPU vẫn có thể phải trả tiền cho chi tiết đó. Một noise layer nhỏ trong fire, một distortion nhẹ trong Win Effect, một card Alpha rộng hơn silhouette thật, tất cả đều có thể cộng lại thành frame-time xấu khi effect xuất hiện nhiều lần.",
                 "Điều này không có nghĩa mobile phải xấu. Nhiều game mobile stylized rất đẹp vì chúng chọn đúng thứ để giữ. Chúng giữ silhouette rõ, color sạch, timing tốt, sprite hoặc flipbook được author cẩn thận. Chúng không cố mô phỏng mọi thứ bằng procedural Shader chạy ở Runtime nếu một Atlas nhỏ làm tốt hơn.",
                 "Với iGaming, bài toán còn rõ hơn. Win effect phải sáng, vui, dễ đọc, lặp đi lặp lại và không làm UI khó nhìn. Rất nhiều sparkle, shine sweep hoặc coin burst nên được bake một phần. Runtime chỉ cần tint, scale, fade và chọn frame. Đó không phải là lười. Đó là production."
             ]),
            ("Stylized không có nghĩa là đơn giản",
             [
                 "Stylized VFX thường bị hiểu nhầm là ít thật hơn nên dễ hơn. Thực tế, stylized đòi hỏi quyết định rõ hơn. Anime Slash cần hard shape, taper, smear, impact timing và color accent. Một Portal stylized cần ring đọc rõ trước, swirl sau, particle accent cuối cùng. Một shield stylized cần rim mạnh nhưng không được thành quả cầu nhựa phát sáng.",
                 "Điểm khó của stylized là mọi thứ lộ. Noise sai lộ. Timing sai lộ. Shape yếu lộ. Nếu realistic VFX đôi khi có thể dựa vào complexity để che lỗi, stylized buộc bạn thiết kế shape thật chắc. Vì vậy, nhiều khi stylized shader cần ít node hơn nhưng nhiều art direction hơn.",
                 "Một cách làm tốt là dựng effect ở Grayscale trước. Nếu Slash, Portal hoặc Shield không đọc được khi chỉ có trắng đen, màu sắc sẽ không cứu nó. Sau đó mới thêm palette, emissive, bloom và particle accent."
             ]),
            ("AAA readability và bài học khi scene quá phức tạp",
             [
                 "AAA không có nghĩa là mọi Shader đều phức tạp. AAA thường có nghĩa là Effect phải sống trong Scene phức tạp: Lighting thay đổi, Camera di chuyển, Motion Blur, Post-process, Gameplay đông, nhiều Asset cùng lúc. Trong môi trường đó, readability là kỹ năng sống còn.",
                 "Một projectile đẹp trong turntable có thể biến mất trong combat. Một impact đẹp trong black background có thể cháy trắng trong level tuyết. Một smoke đẹp trong Material Preview có thể cắt sàn khi đặt vào map thật. Đây là lý do production review phải diễn ra trong context thật càng sớm càng tốt.",
                 "Technical VFX Artist ở môi trường này cần biết khi nào vấn đề nằm ở material và khi nào không. Có lúc shader đúng nhưng lighting sai. Có lúc texture tốt nhưng import setting sai. Có lúc material đẹp nhưng Niagara timing làm nó yếu. Có lúc effect đẹp nhưng gameplay không đọc được. Khả năng phân biệt các trường hợp này mới là giá trị thật."
             ]),
        ],
        "notes": [
            ("Ghi nhớ", "Production không thưởng cho shader phức tạp nếu một giải pháp đơn giản đạt visual tốt hơn, rẻ hơn và dễ tinh chỉnh hơn.", False),
        ],
        "case": [
            "Trong một game stylized, slash đẹp thường đến từ shape và timing hơn là simulation. Một mesh arc có taper tốt, một texture hand-painted sạch và một curve scale đúng nhịp có thể thuyết phục hơn một shader procedural nhiều noise.",
            "Trong một game realistic, cùng slash đó có thể cần sparks, debris, lighting interaction và smoke nhỏ. Cùng một loại effect, hai style dẫn tới hai quyết định khác nhau. Đây là lý do học recipe không đủ; bạn phải học cách đọc art direction."
        ],
        "code": [
            "HighSpec:  flipbook + depth fade + distortion + extra particles",
            "LowSpec:   packed Atlas + tint + scale + alpha fade",
            "UI/iGame:  baked sequence + deterministic timing + minimal shader"
        ],
    },
    {
        "title": "Production wisdom từ cộng đồng RealTimeVFX",
        "image": ("v9_shader_complexity_overdraw.png", "Nhiều lỗi production không nằm ở node riêng lẻ, mà nằm ở cách dữ liệu, renderer và ngân sách frame-time gặp nhau."),
        "sections": [
            ("Đọc forum như đọc nhật ký lỗi của ngành",
             [
                 "RealTimeVFX không nên được đọc như một kho công thức để copy. Giá trị lớn hơn của nó nằm ở những câu hỏi lặp đi lặp lại: tại sao Shader Complexity trắng xóa dù Material đơn giản, tại sao SubUV chạy sai frame, tại sao VAT bị seam hoặc nhiễu, tại sao portfolio đẹp nhưng thiếu breakdown, tại sao một Custom Node HLSL gọn hơn nhưng khó debug hơn. Những câu hỏi đó cho thấy production thật thường hỏng ở ranh giới giữa nhiều hệ thống, không phải ở một node duy nhất.",
                 "Với người học Material VFX, cách đọc tốt là biến mỗi thread thành một Failure Pattern. Câu hỏi đầu tiên không phải là 'giải pháp cuối là gì', mà là 'dữ liệu nào đang bị hiểu sai'. Khi cộng đồng nói về Overdraw, thực chất họ đang nói về diện tích pixel và số lớp Translucent chồng nhau. Khi họ nói về SubUV, họ đang nói về hợp đồng giữa renderer, Atlas và Material. Khi họ nói về VAT, họ đang nói về Texture như dữ liệu số, không còn là hình ảnh trang trí.",
                 "Điểm quan trọng là đừng thần thánh hóa bất kỳ câu trả lời nào. Mỗi project có engine version, platform, content style và budget khác nhau. Nhưng nếu nhiều artist độc lập gặp cùng một nhóm lỗi, nhóm lỗi đó xứng đáng đi vào sách như một phản xạ debug."
             ]),
            ("Shader Complexity là bản đồ nhiệt, không phải bản án",
             [
                 "Một pattern rất rõ trên RealTimeVFX là người mới dễ sợ màu đỏ trong Shader Complexity View. Nỗi sợ này có lý, nhưng nếu hiểu sai, nó dẫn tới tối ưu nhầm chỗ. Shader Complexity không chỉ nói Material có bao nhiêu instruction; nó còn phơi ra chuyện một pixel bị xử lý nhiều lần vì Overdraw. Một shader tương đối đơn giản vẫn có thể đắt nếu card Alpha quá rộng, số particle quá nhiều, hoặc nhiều lớp Translucent nằm cùng một vùng màn hình.",
                 "Vì vậy, khi nhìn một effect đắt, hãy tách câu hỏi thành ba lớp. Lớp thứ nhất là diện tích: card có quá lớn so với silhouette thật không. Lớp thứ hai là chồng lớp: có bao nhiêu particle hoặc sprite phủ lên cùng một pixel. Lớp thứ ba mới là shader: mỗi lần pixel chạy, nó phải sample bao nhiêu Texture và tính bao nhiêu phép toán. Thứ tự này giúp bạn tránh cảnh xóa vài Multiply trong Graph nhưng vẫn để nguyên một đám smoke card chồng lên nhau.",
                 "Trong portfolio hoặc review, Shader Complexity cũng cần dùng đúng mục đích. Nó là công cụ kiểm tra và giải thích, không phải thứ bắt buộc phải show trong mọi reel. Điều có giá trị hơn là bạn chứng minh được mình biết vì sao effect tốn, đã giảm chi phí ở đâu, và tradeoff nào được chọn để giữ visual."
             ]),
            ("SubUV, Flipbook và VAT đều là bài học về Texture như dữ liệu",
             [
                 "Các thread về SubUV và flipbook thường xoay quanh một lỗi cốt lõi: renderer chọn frame theo một quy ước, còn Material lại xử lý UV theo một quy ước khác. Khi bạn muốn đặt gradient, border hoặc radial fade trong từng frame, bạn cần local UV bên trong frame hiện tại. Nếu bạn dùng nhầm Atlas UV, hiệu ứng sẽ nhảy, trượt hoặc chạy xuyên qua nhiều ô. Đây là lý do NumberGrid là công cụ debug tốt hơn một Texture lửa đẹp.",
                 "Vertex Animation Texture đẩy bài học này xa hơn. VAT không phải Texture theo nghĩa mỹ thuật; nó là bảng dữ liệu chuyển động. Compression, Mip Map, Full Precision UV, kích thước Texture, số frame và số vertex đều có thể làm dữ liệu bị sai. Một Texture thường có thể chịu nén nhẹ mà mắt vẫn chấp nhận. Một Texture chứa position hoặc normal data thì không khoan dung như vậy, vì một pixel sai có thể thành một vertex sai.",
                 "Bài học chung là: khi Texture được dùng như dữ liệu, hãy đối xử với nó như dữ liệu. Kiểm tra import setting, tắt Mip Map nếu cần, tránh compression phá giá trị, đảm bảo UV sample đúng pixel, và luôn có debug view. Nếu không, bạn sẽ sửa màu hoặc shader trong khi lỗi thật nằm ở cách dữ liệu được đóng gói."
             ]),
            ("Breakdown tốt cho thấy tư duy, không chỉ khoe kết quả",
             [
                 "Các thread portfolio trên RealTimeVFX có một lời nhắc rất thực dụng: VFX không nên chỉ được show như một clip đẹp trên nền đen. Người xem tuyển dụng hoặc người review cần thấy context, timing, camera, texture flats, emitter breakdown, shader notes và những quyết định production. Một effect đẹp nhưng không có breakdown khiến người xem khó biết bạn thật sự làm gì và hiểu tới đâu.",
                 "Với sách này, điều đó chuyển thành nguyên tắc học: mỗi bài thực hành nên để lại một breakdown nhỏ. Không cần dài. Chỉ cần đủ để trả lời: visual goal là gì, dữ liệu chính đến từ đâu, Shape/Mask/Motion/Color được tạo bằng gì, phần nào điều khiển từ Niagara hoặc Material Instance, phần nào có rủi ro performance, và nếu làm lại cho mobile thì sẽ bỏ lớp nào trước.",
                 "Breakdown cũng là cách chống lại học vẹt. Khi bạn phải giải thích vì sao dùng DepthFade, vì sao cần Particle Color, vì sao crop Alpha giúp giảm Overdraw, bạn buộc phải hiểu effect như một hệ thống. Đây là kỹ năng mà một reel đẹp nhưng im lặng không chứng minh được."
             ]),
            ("Custom Node HLSL: gọn hơn không đồng nghĩa rõ hơn",
             [
                 "RealTimeVFX có nhiều ví dụ artist chuyển một node setup lớn thành Custom Node HLSL để làm blur, light shaft, ripple hoặc sampling loop. Đây là hướng đúng khi Graph bắt đầu lặp và khó đọc. Nhưng HLSL chỉ tốt hơn nếu nó làm contract rõ hơn: input là gì, output là gì, range là bao nhiêu, Texture Sampler dùng wrap hay clamp, loop count có cố định không, và debug intermediate ở đâu.",
                 "Một Custom Node xấu có thể làm Material trông gọn nhưng biến toàn bộ logic thành hộp đen. Artist khác không biết sửa Parameter nào. Lead không biết lỗi nằm ở Mask, UV hay sampling. Programmer nhìn vào cũng ngại maintain vì code không nói rõ ý đồ visual. Vì vậy, code shader trong VFX nên nhỏ, có tên biến theo nghĩa thị giác, có clamp/saturate khi cần, và có đường xuất Grayscale để kiểm tra.",
                 "Khi bạn đã hiểu HLSL/GLSL, node Graph và code không còn đối lập. Node Graph giúp team art đọc nhanh. HLSL giúp diễn đạt vòng lặp, sampling nhiều bước hoặc công thức lặp lại gọn hơn. Người master không chọn phe; họ chọn hình thức nào làm ý tưởng rõ nhất và ít rủi ro production nhất."
             ]),
        ],
        "notes": [
            ("Ghi nhớ", "Một thread hay không chỉ cho bạn câu trả lời. Nó cho bạn triệu chứng, giả định sai, cách debug, và tradeoff production. Hãy ghi lại cả bốn thứ đó.", False),
        ],
        "case": [
            "Nếu Shader Complexity trắng ở một cụm smoke, đừng mở Material Graph trước. Hãy nhìn Alpha card, số particle chồng nhau, khoảng cách camera và diện tích màn hình. Nếu vùng phủ quá lớn, tối ưu instruction chưa giải quyết đúng gốc.",
            "Nếu flipbook Fire/Smoke bị nhảy hoặc gradient chạy sai, hãy thay Texture bằng NumberGrid. Nếu NumberGrid sai, sửa SubUV/frame contract trước. Nếu NumberGrid đúng nhưng visual xấu, lúc đó mới quay lại Mask, Color và Distortion.",
            "Nếu VAT bị nhiễu, seam hoặc vertex nhảy, hãy coi Texture như dữ liệu số: kiểm tra compression, mip, precision, UV, texture size và mapping vertex-frame trước khi chỉnh material mỹ thuật."
        ],
        "code": [
            "RTVFX research lens:",
            "Symptom   -> What does the artist see?",
            "Contract  -> Which systems exchange data?",
            "Debug     -> What view isolates the problem?",
            "Tradeoff  -> What gets cheaper, what gets worse?"
        ],
    },
    {
        "title": "Case study - đọc một Material mẫu trong Unreal",
        "image": ("v11_realtime_contract_simple.png", "Khi đọc Material mẫu, hãy tìm hợp đồng dữ liệu giữa renderer, particle data và Material Graph."),
        "sections": [
            ("Bắt đầu từ bối cảnh trước khi nhìn node",
             [
                 "Trong một project Unreal, bạn thường sẽ gặp những Material có sẵn do người khác để lại: một Material tạo flare, một Material dùng cho mesh additive, một Material chạy flipbook lửa, hoặc một utility Material để kiểm tra SubUV. Ở đây, ta dùng các Material mẫu trong VFXSamplePack làm ví dụ thực hành. Người đọc không cần biết trước bộ sample này; chỉ cần hiểu rằng đây là các asset nhỏ được tạo để minh họa những vấn đề production rất thường gặp.",
                 "Trước khi đuổi theo dây nối trong Graph, hãy tự hỏi Material này đang phục vụ renderer nào. Nó được gắn vào Sprite Renderer của Niagara, mesh renderer, ribbon, hay một card phẳng? Dữ liệu điều khiển đến từ Texture, Particle Color, Dynamic Parameter, Material Instance hay tọa độ trong world? Nếu bỏ qua bối cảnh này, bạn có thể hiểu đúng từng node nhưng vẫn hiểu sai mục đích của Material.",
                 "Một cách đọc chắc hơn là chia Graph thành các làn dữ liệu. Làn UV trả lời câu hỏi Texture được đọc bằng TexCoord gốc, Panner, Rotator, Polar UV hay SubUV. Làn Mask trả lời vùng nào quyết định Opacity, dải biên (Edge Band), Distortion hoặc Soft Fade. Làn Color/Emissive trả lời màu đến từ Texture, Parameter, Particle Color hay Gradient. Làn Interaction kiểm tra DepthFade, CameraFade, SceneDepth, Fresnel, WPO hoặc Dynamic Parameter. Làn Cost xem Material tốn ở diện tích pixel, Overdraw, số lượt đọc Texture hay Instruction."
             ]),
            ("Ví dụ 1: Material flare additive M_Flare",
             [
                 "M_Flare có thể được đọc như một bài học nhỏ về Additive. Mục tiêu visual của nó không phải là 'làm một đốm sáng thật mạnh' mà là tạo một vùng sáng có Shape rõ, có Falloff điều khiển được, và có thể được Niagara hoặc Material Instance chỉnh cường độ khi cần.",
                 "Khi mở loại Material này, hãy tìm bốn phần cụ thể. Phần Shape thường đến từ Texture Alpha hoặc Radial Mask. Phần Color thường đến từ Tint nhân với Particle Color. Phần cường độ thường là một tham số như Intensity hoặc Emissive Strength. Phần Fade có thể đến từ Particle Alpha, Camera Fade hoặc DepthFade. Nếu một trong bốn phần này bị thiếu, flare sẽ khó kiểm soát trong Scene thật.",
                 "Bài test thực tế rất đơn giản: tắt Bloom, giảm Intensity, đặt flare lên nền tối rồi nền sáng. Nếu trên nền tối nó đẹp nhưng trên nền sáng biến mất, vấn đề không chỉ là thiếu sáng. Có thể Shape quá mềm, vùng lõi (Core) không đủ rõ, Falloff quá rộng, hoặc Additive đang phụ thuộc quá nhiều vào Background. Lúc đó tăng Emissive chưa chắc đúng; sửa Mask và Hierarchy thường hiệu quả hơn."
             ]),
            ("Ví dụ 2: Material SubUV và kiểm tra Atlas bằng NumberGrid",
             [
                 "Một nhóm lỗi khác thường xuất hiện ở flipbook lửa, khói hoặc impact là Material và renderer không thống nhất về frame. Renderer chọn một ô trong Texture Sheet, nhưng Material lại dùng UV như thể nó đang đọc toàn bộ Atlas. Kết quả là radial fade, border, distortion hoặc mask theo từng frame chạy sai vị trí.",
                 "Vì vậy, một Texture debug như NumberGrid rất hữu ích. Nó không cần đẹp. Nó chỉ cần đánh số rõ từng ô trong Atlas. Nếu frame đang hiển thị sai số, lỗi nằm ở cách chọn frame hoặc thông số SubUV. Nếu số đúng nhưng gradient chạy xuyên qua nhiều ô, lỗi nằm ở local UV. Nếu frame đúng nhưng bị giật, hãy kiểm tra Frame Rate, Frame Interpolation hoặc curve điều khiển frame trong Niagara.",
                 "Bài học production ở đây là: Texture Sheet không chỉ là hình ảnh, nó là dữ liệu đóng gói. Renderer chọn ô, Material đọc ô, còn Niagara điều khiển thời điểm đổi ô. Ba phần này phải dùng cùng quy ước. Khi một hiệu ứng Fire/Smoke bị sai nhịp hoặc méo mask, hãy kiểm tra Atlas bằng NumberGrid trước khi chỉnh màu."
             ]),
        ],
        "notes": [
            ("Ghi nhớ", "Khi đọc một Material có sẵn, hãy viết Breakdown bằng sáu dòng: renderer đang dùng, Shape, UV, Mask, Color/Emissive, Interaction/Cost. Nếu không viết được, bạn chưa thật sự đọc xong.", False),
        ],
        "case": [
            "Nếu đọc một flare additive, đừng sửa Parameter ngay. Hãy ghi rõ Texture nào tạo Shape, Particle Color đi vào đâu, Intensity nằm trước hay sau Mask, có Camera Fade hay không, và Material có đang phụ thuộc quá nhiều vào Bloom không.",
            "Nếu đọc một flipbook Fire/Smoke, hãy thay Texture đẹp bằng NumberGrid trước. NumberGrid cho bạn biết renderer đang chọn đúng frame chưa, local UV có đúng chưa, và lỗi nằm ở Atlas, Niagara hay Material."
        ],
        "code": [
            "Material reading lanes:",
            "1. Shape / silhouette",
            "2. UV / motion",
            "3. Mask / opacity / edge",
            "4. Color / emissive",
            "5. Interaction / renderer data",
            "6. Cost / platform risk"
        ],
    },
    {
        "title": "Lộ trình luyện tập - từ hiểu tới phản xạ",
        "image": ("v9_space_vectors.png", "Space transform là một trong những phần nên luyện chậm, vì nó xuất hiện ở nhiều shader nâng cao."),
        "sections": [
            ("Đọc không đủ, phải dựng lại",
             [
                 "Shader giống ngoại ngữ ở một điểm rất rõ: đọc hiểu chưa đủ. Bạn có thể hiểu SmoothStep khi nhìn công thức, nhưng nếu chưa tự dùng nó để làm dissolve edge, ring, foam hoặc shockwave, bạn chưa có phản xạ. Bạn có thể hiểu Particle Color trên giấy, nhưng nếu chưa từng quên nhân ParticleColor.a rồi tự debug vì Niagara fade không tác dụng, bài học đó chưa bám vào tay.",
                 "Vì vậy, lộ trình học Material nên có output nhỏ liên tục. Mỗi output phải đủ nhỏ để hoàn thành, nhưng đủ thật để có lỗi. Một ngày dựng Mask Grayscale. Một ngày làm radial ring. Một ngày làm fire card baseline. Một ngày đọc M_Flare. Một ngày viết HLSL EdgeBand. Những bài nhỏ này tạo vốn từ vựng thực chiến.",
                 "Đừng đợi tới khi hiểu hết GPU mới làm effect. Cũng đừng chỉ copy recipe mà không hiểu cost. Học tốt nhất là xen kẽ: đọc một khái niệm, dựng một bản tối giản, phá nó, sửa nó, rồi viết lại bằng lời của mình."
             ]),
            ("Một lộ trình 8 tuần thực dụng",
             [
                 "Tuần đầu tiên nên dành cho Grayscale. Bạn dựng Mask từ Gradient, Noise, Radial Distance, UV.x, UV.y và Texture Alpha. Mỗi Mask phải có Threshold, Softness và Debug View. Tuần thứ hai chuyển sang shaping: Power, SmoothStep, remap, contrast, dải biên (Edge Band). Đừng thêm màu quá sớm.",
                 "Tuần ba học UV Motion: Panner, Rotator, Polar, Distortion và Flow. Tuần bốn học Blend Mode, Emissive, Opacity, DepthFade và CameraFade. Tuần năm học Particle Color, Dynamic Parameter, SubUV và NumberGrid. Tuần sáu học Fresnel, Space Transform, WPO và Shield Hit. Tuần bảy học Optimization: Overdraw, lượt đọc Texture, Channel Packing và Mobile Variant. Tuần tám đọc lại một Material mẫu trong project thật và viết Breakdown.",
                 "Điều quan trọng không phải hoàn thành lộ trình hoàn hảo. Điều quan trọng là mỗi tuần có artifact. Một Material nhỏ, một screenshot Grayscale, một ghi chú Debug, một câu giải thích. Sau vài tháng, bạn sẽ có thư viện tư duy của riêng mình."
             ]),
            ("Tự kiểm tra bằng câu hỏi production",
             [
                 "Sau mỗi effect, hãy hỏi vài câu khó. Nếu người khác mở Material này, họ có hiểu parameter không? Nếu effect spawn 100 lần, nó có còn ổn không? Nếu background sáng, Shape có đọc được không? Nếu tắt bloom, visual có sụp không? Nếu chuyển sang mobile, bạn sẽ bỏ layer nào trước? Nếu lead nói organic hơn, bạn sẽ sửa lane nào?",
                 "Những câu hỏi này giúp bạn rời khỏi tư duy làm cho xong. Bạn bắt đầu nghĩ như người chịu trách nhiệm với effect trong production. Đây là ranh giới giữa artist biết node và Technical VFX Artist biết hệ thống.",
                 "Không có con đường tắt để master Shader VFX. Nhưng có con đường rõ: nhìn bằng dữ liệu, học toán theo visual, hiểu cơ bắp render, viết được shader khi cần, và luôn kiểm tra effect trong Scene thật."
             ]),
        ],
        "notes": [
            ("Ghi nhớ", "Mỗi Material học được nên để lại ba thứ: một screenshot Debug, một câu giải thích visual, và một ghi chú Cost.", False),
        ],
        "case": [
            "Một bài tập tốt không cần hoành tráng. Một Radial Mask có Power và SmoothStep đủ để bạn học falloff. Một Dissolve Grayscale đủ để bạn học Threshold. Một SubUV NumberGrid đủ để bạn hiểu Atlas. Nếu bài tập quá lớn, bạn sẽ nhớ cảm giác hoàn thành nhưng quên mất khái niệm.",
            "Hãy giữ mỗi bài nhỏ tới mức bạn có thể làm lại từ đầu trong 30 phút. Khi bạn làm lại được, giải thích được và phá được nó, khái niệm đã bắt đầu thành phản xạ."
        ],
        "code": [
            "Weekly loop:",
            "Read one concept",
            "Build one tiny material",
            "Break it intentionally",
            "Fix it",
            "Explain it in six lanes"
        ],
    },
]


def chapter(doc, idx, data):
    h1(doc, data["title"], idx)
    chapter_image = data.get("image")
    for section_idx, (title, paras) in enumerate(data["sections"]):
        img = chapter_image if section_idx == 0 else None
        grid_block(
            doc,
            title,
            paras,
            margin_title="Ghi chú lề",
            margin_items=margin_notes(title, paras),
            image=img,
        )
    for title, text, warning in data.get("notes", []):
        grid_block(
            doc,
            title,
            [text],
            margin_title="Dùng khi",
            margin_items=["Cần nhớ nhanh khi đang đọc lại chương.", "Không cắt ngang mạch chính của bài."],
        )
    if data.get("case"):
        grid_block(
            doc,
            "Case đọc chậm",
            data["case"],
            margin_title="Production lens",
            margin_items=["Triệu chứng", "Nguyên nhân dữ liệu", "Debug view", "Tradeoff"],
        )
    if data.get("code"):
        grid_block(
            doc,
            "Sketch logic",
            ["Đọc block này như một bản đồ tư duy tối giản. Nó không thay thế Material Graph hoặc HLSL thật, nhưng giúp bạn nhớ thứ tự dữ liệu chạy qua effect."],
            margin_title="Code lề",
            margin_items=["Giữ tên biến theo ý đồ visual.", "Tách Shape, Mask, Color và Cost."],
            code_lines=data["code"],
        )


def closing(doc):
    h1(doc, "Kết luận - viết lại bằng chính mắt của bạn")
    body(doc, "Một người học Material thường muốn đi nhanh tới kết quả cuối cùng. Điều đó tự nhiên, vì VFX là nghề của hình ảnh. Nhưng càng làm lâu, bạn càng thấy tốc độ thật không đến từ việc nối node nhanh hơn. Tốc độ thật đến từ việc nhìn đúng vấn đề ngay từ đầu. Trước khi thêm node, hãy nhìn effect như dữ liệu. Trước khi tăng emissive, hãy xem Hierarchy. Trước khi gọi một effect là xấu, hãy hỏi nó đang hỏng ở shape, mask, motion, color, interaction hay cost.")
    body(doc, "Một cuốn sách Material VFX tốt không thể chỉ là danh sách công thức. Công thức cần thiết, nhưng công thức phải nằm trong câu chuyện của visual và production. Người học cần biết vì sao một trick tồn tại, nó giải quyết vấn đề gì, nó hỏng trong trường hợp nào, và khi nào nên bỏ nó để dùng một giải pháp rẻ hơn.")
    body(doc, "Con đường học Shader VFX nên được xây bằng nhiều bài nhỏ nhưng thật. Dựng một Mask Grayscale. Làm một dải biên (Edge Band). Đọc một Atlas bằng NumberGrid. Tạo một Flare Additive rồi kiểm tra nó trên nền sáng. Làm một Fire Card có DepthFade. Viết một hàm HLSL nhỏ rồi so với Node Graph. Mỗi bài tập như vậy để lại một phần phản xạ. Khi phản xạ đủ nhiều, bạn không chỉ biết làm một hiệu ứng cụ thể; bạn biết cách nghĩ khi gặp một hiệu ứng mới.")


def references(doc):
    h1(doc, "Nguồn tham khảo và hướng đọc tiếp")
    body(doc, "Các nguồn dưới đây được dùng như định hướng học, không phải để trích dài. Khi đọc tiếp, bạn nên đọc theo vấn đề đang gặp trong production thay vì đọc dàn trải tất cả cùng lúc.")
    refs = [
        "The Book of Shaders - cách tiếp cận shader qua pixel, tọa độ, Shaping Function, pattern và noise.",
        "Unity URP Shader/VFX e-book - hướng đọc Shader Graph, URP, VFX Graph, Render Graph, mobile và stylized rendering.",
        "Unreal Engine Material documentation - Material Editor, parameter, Blend Mode, Custom Material Expressions và HLSL.",
        "Unreal Niagara documentation - renderer, Particle Color, Dynamic Parameter, SubUV và dữ liệu runtime.",
        "Real-Time Rendering - nền tảng Graphics Pipeline, rasterization, texture, shading và GPU thinking.",
        "RealTimeVFX - VFX Optimization: overdraw, draw calls, GPU particles, pooling, culling và scalability trong Niagara.",
        "RealTimeVFX - Shader Complexity discussions: xem Shader Complexity như bản đồ cảnh báo, không phải số đo duy nhất của performance.",
        "RealTimeVFX - Vertex Animation Texture threads: Texture data cần đúng compression, mip, precision, UV và kích thước đóng gói.",
        "RealTimeVFX - Niagara/SubUV/Flipbook threads: phân biệt mesh flipbook, texture SubUV, frame index và local UV.",
        "RealTimeVFX - Portfolio/breakdown discussions: show process, shader notes, particle setup và performance thinking thay vì chỉ show final clip.",
        "VFXSamplePack / ForgePilot Knowledge - nguồn thực hành trong project hiện tại."
    ]
    for ref in refs:
        bullet(doc, ref)


def build():
    doc = setup()
    cover(doc)
    preface(doc)
    for i, ch in enumerate(CHAPTERS, 1):
        chapter(doc, i, ch)
    closing(doc)
    references(doc)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()





