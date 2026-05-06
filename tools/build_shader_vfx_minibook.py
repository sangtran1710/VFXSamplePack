from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path.home() / "Downloads" / "Realtime_Shader_VFX_Mini_Book_Dense.docx"

INK = RGBColor(32, 32, 32)
MUTED = RGBColor(90, 90, 90)
BLUE = RGBColor(22, 71, 112)
TEAL = RGBColor(0, 112, 112)
RED = RGBColor(150, 45, 45)
GOLD = RGBColor(145, 94, 10)
LIGHT_BLUE = "EAF3F8"
LIGHT_TEAL = "E9F5F4"
LIGHT_GOLD = "FFF3D6"
LIGHT_GRAY = "F1F3F5"


def font(run, size=10, bold=False, italic=False, color=INK, name="Arial"):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def shade(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def border_bottom(paragraph, color="D9E2EA", size="10"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)


def para(doc, text="", size=10, color=INK, bold=False, italic=False, before=0, after=4, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.08
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        font(r, size=size, bold=bold, italic=italic, color=color)
    return p


def h1(doc, number, title):
    p = para(doc, f"{number}. {title}", size=18, color=BLUE, bold=True, before=4, after=8)
    border_bottom(p)
    return p


def h2(doc, title):
    return para(doc, title, size=12.5, color=TEAL, bold=True, before=9, after=4)


def body(doc, text):
    return para(doc, text, size=10.2, after=5)


def bullet(doc, text, level=0, size=9.6):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.58 + 0.35 * level)
    p.paragraph_format.first_line_indent = Cm(-0.28)
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run("- " + text)
    font(r, size=size)
    return p


def numbered(doc, index, text, size=9.6):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.72)
    p.paragraph_format.first_line_indent = Cm(-0.44)
    p.paragraph_format.space_after = Pt(2.5)
    r = p.add_run(f"{index}. {text}")
    font(r, size=size)
    return p


def term(doc, name, meaning, use):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.38)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.02
    r = p.add_run(name + ": ")
    font(r, size=9.5, bold=True, color=BLUE)
    r = p.add_run(meaning)
    font(r, size=9.5)
    r = p.add_run(" Use: ")
    font(r, size=9.5, bold=True, color=TEAL)
    r = p.add_run(use)
    font(r, size=9.5)


def callout(doc, title, lines, fill=LIGHT_BLUE, color=BLUE):
    p = para(doc, title, size=10.4, color=color, bold=True, before=7, after=2)
    p.paragraph_format.left_indent = Cm(0.25)
    p.paragraph_format.right_indent = Cm(0.25)
    shade(p, fill)
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.55)
        p.paragraph_format.right_indent = Cm(0.25)
        p.paragraph_format.space_after = Pt(1.5)
        shade(p, fill)
        r = p.add_run("- " + line)
        font(r, size=9.4)
    para(doc, "", after=1)


def codebox(doc, lines, title=None):
    if title:
        para(doc, title, size=9.5, color=MUTED, italic=True, after=1)
    for line in lines:
        p = para(doc, line, size=8.9, after=1, color=INK)
        p.paragraph_format.left_indent = Cm(0.35)
        p.paragraph_format.right_indent = Cm(0.2)
        shade(p, LIGHT_GRAY)
        for r in p.runs:
            font(r, size=8.9, name="Consolas")
    para(doc, "", after=1)


def checklist(doc, items):
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.58)
        p.paragraph_format.first_line_indent = Cm(-0.28)
        p.paragraph_format.space_after = Pt(2.5)
        r = p.add_run("[ ] " + item)
        font(r, size=9.5)


def page(doc):
    doc.add_page_break()


def chapter(doc, number, title, intro=None):
    h1(doc, number, title)
    if intro:
        body(doc, intro)


def cover(doc):
    para(doc, "", before=6)
    p = para(doc, "Realtime Shader VFX", size=25, color=BLUE, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    p.paragraph_format.line_spacing = 1.0
    p = para(doc, "Mini-Book", size=25, color=BLUE, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    p.paragraph_format.line_spacing = 1.0
    para(
        doc,
        "Material thinking for Unreal, Unity, Niagara, VFX Graph, Shader Graph, and engine-neutral technical VFX",
        size=10.5,
        color=INK,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=12,
    )
    callout(
        doc,
        "How to read this book",
        [
            "Read one chapter, then open one real material and identify the same concept inside the graph.",
            "For every recipe, explain the shape, mask, motion, color, interaction, and performance cost.",
            "Use the English sentences as speaking drills for review, interview, and daily production communication.",
        ],
        fill=LIGHT_GOLD,
        color=GOLD,
    )
    para(doc, "Prepared for Technical VFX study", size=9, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=10)


def toc(doc):
    h1(doc, "Contents", "Study Map")
    entries = [
        "1. The Mental Model: From pixels to realtime effects",
        "2. Engine Material Pipeline: Unreal, Unity, and the shared core",
        "3. Blend Modes and Rendering Cost",
        "4. UV Space: The first language of VFX materials",
        "5. Mask Grammar: step, smoothstep, remap, contrast",
        "6. Color, Emissive, Exposure, and Bloom",
        "7. Texture Authoring, Channel Packing, and Flipbooks",
        "8. SubUV and Derived 0-1 Frame UV",
        "9. Distortion, Normals, Refraction, and Scene Depth",
        "10. Coordinate Spaces: Object, world, view, screen, tangent",
        "11. Unreal Material Reading Guide",
        "12. Unity Shader Graph / VFX Graph Reading Guide",
        "13. Optimization for Technical VFX",
        "14. Recipe: Additive Flare",
        "15. Recipe: Dissolve / Reveal",
        "16. Recipe: Stylized Fire Card",
        "17. Recipe: Smoke Flipbook",
        "18. Recipe: Slash / Arc",
        "19. Recipe: Shield Hit Ripple",
        "20. Recipe: Portal / Black Hole",
        "21. Production Checklists",
        "22. English Speaking for Technical VFX",
        "23. Six-Week Study Plan",
        "24. Glossary",
    ]
    for e in entries:
        bullet(doc, e, size=9.8)
    callout(doc, "Promise of the book", [
        "The goal is not to memorize a graph. The goal is to understand the visual grammar deeply enough to rebuild the graph in any engine."
    ], fill=LIGHT_TEAL)
    page(doc)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.55)
    section.bottom_margin = Cm(1.45)
    section.left_margin = Cm(1.75)
    section.right_margin = Cm(1.75)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)

    for s in doc.sections:
        footer = s.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = footer.add_run("Realtime Shader VFX Mini-Book")
        font(r, size=8.5, color=MUTED)

    cover(doc)
    toc(doc)

    chapter(doc, 1, "The Mental Model", "A realtime VFX material is not just a pretty graph. It is a small visual system that turns data into pixels every frame.")
    h2(doc, "The core pipeline")
    codebox(doc, [
        "Texture / math / particle data",
        "-> material / shader",
        "-> renderer: sprite, mesh, ribbon, decal, post-process",
        "-> frame buffer: color, alpha, depth, bloom, distortion",
    ])
    body(doc, "When you see a finished effect, train yourself to separate it into six layers:")
    for i, text in enumerate([
        "Shape: the silhouette or readable graphic idea.",
        "Mask: what is visible, hidden, softened, cut, or eroded.",
        "Motion: how UVs, particles, frames, vertices, or time create life.",
        "Color: value, hue, emissive strength, and contrast hierarchy.",
        "Interaction: depth, camera, collision, hit position, sockets, gameplay parameters.",
        "Cost: overdraw, shader instructions, texture fetches, memory, particles, bounds.",
    ], 1):
        numbered(doc, i, text)
    callout(doc, "Tutor question", ["If an art director says 'make it more powerful', which layer are they really asking you to change: shape, motion, color, or timing?"], fill=LIGHT_GOLD, color=GOLD)

    chapter(doc, 2, "Engine Material Pipeline", "Unreal, Unity, proprietary engines, and custom HLSL all share the same foundation: input data becomes shader code, shader code becomes pixels.")
    h2(doc, "What changes between engines")
    bullet(doc, "Unreal Material Graph gives a strong node-based authoring workflow and compiles to HLSL.")
    bullet(doc, "Unity Shader Graph is similar conceptually, but pipeline details depend on URP/HDRP/custom render pipeline.")
    bullet(doc, "Niagara and VFX Graph are particle/data systems; material/shader is the pixel or surface logic.")
    bullet(doc, "A proprietary engine may expose fewer nodes, but the math vocabulary is still the same.")
    h2(doc, "What stays the same")
    for item in ["UVs", "mask math", "texture sampling", "color blending", "depth", "normals", "view direction", "world position", "performance budgets"]:
        bullet(doc, item)
    callout(doc, "Key idea", ["Master the concepts, not just the UI. If you understand the grammar, Unreal nodes, Unity nodes, and HLSL become different accents of the same language."], fill=LIGHT_TEAL)

    chapter(doc, 3, "Blend Modes and Rendering Cost", "Blend mode determines how your material is written into the frame. For VFX, this choice is both artistic and technical.")
    h2(doc, "Common modes")
    bullet(doc, "Opaque: cheapest and depth-friendly, but not suitable for soft smoke or glow.")
    bullet(doc, "Masked / Alpha Clip: hard cutout, better depth behavior, useful for foliage, stylized cards, decals.")
    bullet(doc, "Translucent: soft alpha, smoke, glass, energy fields. More expensive and can have sorting issues.")
    bullet(doc, "Additive: adds brightness on top of the scene, great for magic, sparks, flares, fire cores.")
    bullet(doc, "Modulate / multiply-like approaches: darken or tint scene depending on engine support.")
    h2(doc, "Production instinct")
    bullet(doc, "If the effect can be masked instead of translucent, test it. Overdraw often dominates VFX cost.")
    bullet(doc, "Additive reads well but can disappear on bright backgrounds. Use shape/value planning.")
    bullet(doc, "Translucent smoke looks soft but can become expensive fast when many large sprites overlap.")
    callout(doc, "English drill", ["This effect uses additive blending for the core and translucent blending for the soft outer smoke, so we need to watch overdraw and readability on bright backgrounds."], fill=LIGHT_BLUE)

    chapter(doc, 4, "UV Space", "UV is the first language of material motion. Most 2D VFX materials are UV manipulation plus masks plus color.")
    h2(doc, "Vocabulary")
    term(doc, "Tiling", "Scale UVs so a texture repeats more or less often.", "noise density, flame detail, scanlines.")
    term(doc, "Panner", "Offset UVs over time.", "scrolling smoke, fire, energy flow.")
    term(doc, "Rotator", "Rotate UVs around a pivot.", "magic circle, portal swirl, rotating glints.")
    term(doc, "Polar UV", "Convert x/y into radius/angle.", "vortex, black hole, slash arc.")
    term(doc, "Flow map", "Use a texture's RG channels as UV direction.", "fluid-like movement without simulation.")
    h2(doc, "Mental formula")
    codebox(doc, [
        "finalUV = baseUV * tiling + time * speed",
        "sample = Texture(finalUV)",
        "mask = process(sample)",
    ])
    callout(doc, "Exercise", ["Open any panning material. Find where time enters the graph, then ask: is time moving UVs, changing a threshold, or changing color?"], fill=LIGHT_GOLD, color=GOLD)

    chapter(doc, 5, "Mask Grammar", "A mask is a control signal. It decides where an effect exists and how strong it is.")
    h2(doc, "Core functions")
    term(doc, "step(edge, x)", "Hard 0/1 cut.", "toon style, hard reveal, pixelated effects.")
    term(doc, "smoothstep(a, b, x)", "Soft transition from 0 to 1.", "soft edge, depth fade, dissolve band.")
    term(doc, "pow(x, n)", "Change contrast curve.", "flare core, fresnel intensity, mask sharpness.")
    term(doc, "saturate(x)", "Clamp to 0-1.", "clean final masks.")
    term(doc, "lerp(a, b, t)", "Blend two values using mask t.", "color ramp, texture blend, dissolve state.")
    h2(doc, "Dissolve as grammar")
    codebox(doc, [
        "baseMask = noise or gradient",
        "visible = smoothstep(threshold, threshold + softness, baseMask)",
        "edge = smoothstep(threshold - width, threshold, baseMask) - visible",
        "emissive = baseColor * visible + edgeColor * edge * intensity",
    ])
    callout(doc, "Feedback mapping", ["'Make the edge softer' usually means widen smoothstep range. 'Make it punchier' often means stronger edge value, faster timing, or more contrast."], fill=LIGHT_TEAL)

    chapter(doc, 6, "Color, Emissive, Exposure, Bloom", "VFX color is not only hue. It is value hierarchy, glow behavior, readability, and exposure.")
    h2(doc, "Value hierarchy")
    bullet(doc, "Core: the brightest and most readable area.")
    bullet(doc, "Edge/rim: defines silhouette and graphic identity.")
    bullet(doc, "Body: medium-value fill, often textured or noisy.")
    bullet(doc, "Falloff: soft fade into the scene.")
    h2(doc, "Emissive notes")
    bullet(doc, "Emissive above 1 can trigger bloom depending on post-process/exposure.")
    bullet(doc, "Auto exposure can make the same emissive value look different in dark vs bright scenes.")
    bullet(doc, "A beautiful glow in preview can become unreadable in-game if value contrast is not planned.")
    bullet(doc, "Color ramps are powerful because they separate mask logic from art direction.")
    callout(doc, "English drill", ["The core is intentionally overbright for bloom, but the outer glow is kept lower so the shape remains readable on bright backgrounds."], fill=LIGHT_BLUE)

    chapter(doc, 7, "Textures, Packing, and Flipbooks", "Textures are not just images. In VFX they are data: masks, noise, normals, motion vectors, emissive, opacity, occlusion.")
    h2(doc, "Channel packing")
    bullet(doc, "R can store emissive mask or primary noise.")
    bullet(doc, "G can store opacity, secondary breakup, or occlusion.")
    bullet(doc, "B can store another mask, detail noise, or a fake lighting term.")
    bullet(doc, "A is expensive in memory depending on compression, so use it intentionally.")
    h2(doc, "Flipbook logic")
    body(doc, "A flipbook stores many frames in one atlas. The particle system or shader chooses which cell to sample over time.")
    codebox(doc, [
        "columns = 8, rows = 8",
        "frameIndex = floor(normalizedAge * totalFrames)",
        "cellUV = localUV / float2(columns, rows) + frameOffset",
    ])
    callout(doc, "Production note", ["AE, Houdini, EmberGen, Unreal Niagara Baker, or proprietary tools can all create source flipbooks. The engine material decides how that flipbook lives in-game."], fill=LIGHT_GOLD, color=GOLD)

    chapter(doc, 8, "SubUV and Derived 0-1 Frame UV", "SubUV is one of the most important concepts for fire, smoke, explosion, splash, and many stylized sprite effects.")
    h2(doc, "The problem")
    body(doc, "When Sprite Renderer selects a frame inside an atlas, the real texture UV may be only a small rectangle of the full texture. But material math often expects local 0-1 UV.")
    h2(doc, "Why derive 0-1 UV")
    bullet(doc, "Radial masks need center 0.5,0.5 inside the current frame, not center of the full atlas.")
    bullet(doc, "Edge fade should wrap around the current frame, not the entire sheet.")
    bullet(doc, "Distortion should be stable per frame instead of drifting because of atlas position.")
    bullet(doc, "Debug textures like NumberGrid help you see which frame is being sampled.")
    callout(doc, "Sentence to remember", ["Sprite Renderer chooses which atlas cell to use; derived 0-1 UV lets the material treat that cell like its own independent texture."], fill=LIGHT_TEAL)

    chapter(doc, 9, "Distortion, Normals, Refraction, Scene Depth", "Distortion is powerful because it changes not only the effect, but the scene behind it.")
    h2(doc, "Common sources")
    bullet(doc, "Noise texture: cheap organic wobble.")
    bullet(doc, "Normal map: structured refraction direction.")
    bullet(doc, "Flow map: guided motion for heat, magic, water, smoke.")
    bullet(doc, "Scene depth: intersection and soft particle behavior.")
    h2(doc, "Risks")
    bullet(doc, "Screen-space distortion can artifact near silhouettes and depth discontinuities.")
    bullet(doc, "Refraction/translucency cost varies by engine and render pipeline.")
    bullet(doc, "Too much distortion can reduce readability even if it looks expensive.")
    callout(doc, "Technical note", ["Distortion should usually be strongest near a core, edge, or shock front. If it affects the entire screen evenly, it often feels like a post-process bug."], fill=LIGHT_BLUE)

    chapter(doc, 10, "Coordinate Spaces", "Many shader bugs are not math bugs. They are space bugs.")
    h2(doc, "Spaces")
    bullet(doc, "UV space: 0-1 texture coordinates. Best for sprites, masks, flipbooks.")
    bullet(doc, "Object/local space: stable relative to the mesh. Best for mesh gradients and local dissolve.")
    bullet(doc, "World space: stable in the scene. Best for world-aligned noise or global height effects.")
    bullet(doc, "View space: relative to camera. Best for facing masks and view-driven effects.")
    bullet(doc, "Screen space: pixel/screen coordinates. Best for post-process and scene distortion.")
    h2(doc, "Debug method")
    numbered(doc, 1, "Visualize the value as color.")
    numbered(doc, 2, "Move the camera. If it changes, it may be view/screen dependent.")
    numbered(doc, 3, "Move the object. If it swims, your space may not be object-local.")
    numbered(doc, 4, "Scale the object. If pattern density changes unexpectedly, check coordinate normalization.")

    chapter(doc, 11, "Unreal Material Reading Guide", "When reading a complex Unreal material, start from the outputs and work backward.")
    h2(doc, "Output-first reading")
    bullet(doc, "Emissive Color: where glow and visible VFX color usually end up.")
    bullet(doc, "Opacity / Opacity Mask: where visibility is controlled.")
    bullet(doc, "World Position Offset: where mesh deformation or vertex animation happens.")
    bullet(doc, "Refraction / Pixel Normal Offset: where screen or scene distortion may happen.")
    bullet(doc, "Material Attributes: if used, unpack mentally into the same output categories.")
    h2(doc, "Unreal-specific nodes to recognize")
    bullet(doc, "Particle Color / Dynamic Parameter: data from Niagara/Cascade into material.")
    bullet(doc, "DepthFade / PixelDepth / SceneDepth: soft particles and intersections.")
    bullet(doc, "Camera Vector / Fresnel: view-dependent rim effects.")
    bullet(doc, "Object Position / Actor Position / World Position: spatial masks and gradients.")
    bullet(doc, "Custom node: HLSL escape hatch. Read inputs first, then math.")
    callout(doc, "VFXSamplePack drill", ["Open M_Flare, M_Mesh_Add, and M_FlamesCard. For each one, write one sentence for UV, one for mask, one for color, one for performance risk."], fill=LIGHT_GOLD, color=GOLD)

    chapter(doc, 12, "Unity Shader Graph / VFX Graph Reading Guide", "Unity uses different UI, but the same shader grammar applies.")
    h2(doc, "Shader Graph mindset")
    bullet(doc, "Master Stack outputs are equivalent to final material channels.")
    bullet(doc, "URP/HDRP differences matter for lighting, transparency, distortion, and depth access.")
    bullet(doc, "Properties and Blackboard values are your material parameters.")
    bullet(doc, "Sub Graphs are reusable vocabulary, similar to material functions.")
    h2(doc, "VFX Graph mindset")
    bullet(doc, "Spawn/update blocks control particle data.")
    bullet(doc, "Output blocks decide how particle data becomes pixels.")
    bullet(doc, "Flipbook, color over life, size over life, and custom attributes bridge particles into shader logic.")
    bullet(doc, "Performance depends on particle count, bounds, update simulation, output mode, and material cost.")
    callout(doc, "Transfer skill", ["If you can explain an effect as data -> mask -> color -> renderer -> cost, you can move between Unreal and Unity much faster."], fill=LIGHT_TEAL)

    chapter(doc, 13, "Optimization for Technical VFX", "Optimization is not a final cleanup step. It is part of the design.")
    h2(doc, "Main cost buckets")
    bullet(doc, "Overdraw: many translucent pixels drawn repeatedly.")
    bullet(doc, "Shader instructions: expensive math, branches, repeated operations.")
    bullet(doc, "Texture fetches: many samples, large textures, bad cache behavior.")
    bullet(doc, "Memory: large flipbooks, alpha channels, compression choices.")
    bullet(doc, "Particle simulation: count, collision, GPU/CPU cost, bounds.")
    bullet(doc, "Render pipeline features: refraction, lighting, shadows, sorting.")
    h2(doc, "Common fixes")
    bullet(doc, "Use smaller cards and tighter alpha where possible.")
    bullet(doc, "Bake procedural detail to texture when reused often.")
    bullet(doc, "Pack channels and reuse samples.")
    bullet(doc, "Split high-end and low-end material versions when needed.")
    bullet(doc, "Cull by distance and reduce spawn count with LOD/scalability.")
    callout(doc, "English drill", ["I moved the procedural noise into a packed texture so the material uses fewer instructions and can run on many particles at once."], fill=LIGHT_BLUE)

    chapter(doc, 14, "Recipe: Additive Flare", "This is the hello-world of realtime VFX materials.")
    h2(doc, "Structure")
    codebox(doc, [
        "flare = Texture(flareUV)",
        "mask = pow(flare.r, exponent)",
        "color = particleColor.rgb * tint * intensity",
        "emissive = color * mask",
        "opacity = mask * particleColor.a * cameraFade",
    ])
    h2(doc, "Controls")
    bullet(doc, "Intensity: brightness and bloom strength.")
    bullet(doc, "Exponent: core sharpness and contrast.")
    bullet(doc, "Desaturation: make texture grayscale before tinting.")
    bullet(doc, "Camera fade: prevent harsh clipping near camera.")
    h2(doc, "Feedback mapping")
    bullet(doc, "More punch: increase intensity or exponent.")
    bullet(doc, "Softer: lower exponent, soften texture, reduce opacity edge.")
    bullet(doc, "Too invisible on bright scenes: add stronger core shape or value contrast.")

    chapter(doc, 15, "Recipe: Dissolve / Reveal", "Dissolve is a mask threshold effect. Once mastered, it becomes burn, teleport, poison, ice melt, digital reveal, or shield break.")
    h2(doc, "Structure")
    codebox(doc, [
        "m = noiseOrGradient",
        "visible = smoothstep(amount, amount + softness, m)",
        "edge = smoothstep(amount - edgeWidth, amount, m) - visible",
        "base = albedoOrEmissive * visible",
        "hotEdge = edgeColor * edge * edgeIntensity",
    ])
    h2(doc, "Advanced additions")
    bullet(doc, "Use object/world gradient to control direction.")
    bullet(doc, "Use two noise scales: big shape + small detail.")
    bullet(doc, "Use particle age or gameplay parameter as dissolve amount.")
    bullet(doc, "Use WPO or vertex offset for crumbling/expanding breakup.")
    callout(doc, "Speaking drill", ["The reveal is controlled by a threshold moving through a noise mask. The edge band is a narrow smoothstep region around that threshold."], fill=LIGHT_BLUE)

    chapter(doc, 16, "Recipe: Stylized Fire Card", "A fire card is usually layered UV motion plus masks plus color ramp.")
    h2(doc, "Structure")
    bullet(doc, "Base flame shape: vertical gradient or hand-authored flame texture.")
    bullet(doc, "Motion: upward panning noise or flow map.")
    bullet(doc, "Breakup: noise threshold controls licking edges.")
    bullet(doc, "Color ramp: black/red/orange/yellow/white or stylized palette.")
    bullet(doc, "Opacity: softer at top, stronger core at bottom or center.")
    h2(doc, "Common controls")
    bullet(doc, "Speed U/V: directional flame motion.")
    bullet(doc, "U/V scale: detail density.")
    bullet(doc, "Edge fade: card boundary softness.")
    bullet(doc, "Distortion scale: heat shimmer or internal turbulence.")
    callout(doc, "Unreal link", ["M_FlamesCard-like materials teach the difference between opacity shape, color intensity, UV speed, and distortion layer."], fill=LIGHT_GOLD, color=GOLD)

    chapter(doc, 17, "Recipe: Smoke Flipbook", "Smoke often looks realistic because the hard simulation is baked into a flipbook, then controlled cheaply in engine.")
    h2(doc, "Structure")
    bullet(doc, "Flipbook texture: density/opacity/emissive/occlusion/normals depending on pack.")
    bullet(doc, "SubUV animation: frame selection over lifetime.")
    bullet(doc, "Particle color: tint and alpha over life.")
    bullet(doc, "Depth fade: soft intersection with ground/objects.")
    bullet(doc, "Lighting fake: baked occlusion, normals, or color ramp.")
    h2(doc, "Production notes")
    bullet(doc, "Large smoke cards are overdraw heavy.")
    bullet(doc, "Soft alpha looks good but can become expensive with many overlapping particles.")
    bullet(doc, "Looping smoke needs crossfade or carefully baked loop frames.")
    bullet(doc, "Texture memory can be high for 8x8 or 8x16 atlases.")

    chapter(doc, 18, "Recipe: Slash / Arc", "Stylized anime slash is usually not a simulation. It is graphic design plus timing.")
    h2(doc, "Structure")
    bullet(doc, "Shape: crescent, arc, smear, ribbon, or mesh strip.")
    bullet(doc, "Mask: gradient along length and thickness.")
    bullet(doc, "Motion: panning streak texture, animated alpha, or mesh/ribbon trail.")
    bullet(doc, "Color: bright core, tinted rim, soft outer glow.")
    bullet(doc, "Timing: anticipation, impact frame, decay.")
    h2(doc, "Shader approaches")
    bullet(doc, "UV gradient along U for head/tail fade.")
    bullet(doc, "V gradient for thickness fade.")
    bullet(doc, "Noise erosion for magical breakup.")
    bullet(doc, "Fresnel or view-facing mesh normals for 3D arc readability.")
    callout(doc, "AE transfer", ["Your AE sense of shape, smear, rhythm, and glow is valuable here. Shader knowledge makes that visual interactive and reusable in game."], fill=LIGHT_TEAL)

    chapter(doc, 19, "Recipe: Shield Hit Ripple", "Shield effects combine view-dependent rim, surface pattern, depth interaction, and event-driven hit data.")
    h2(doc, "Structure")
    codebox(doc, [
        "rim = pow(1 - dot(normal, viewDir), fresnelPower)",
        "hitDistance = distance(worldPos, hitPos)",
        "ripple = smoothstep(radius, radius + width, hitDistance) - smoothstep(radius + width, radius + width * 2, hitDistance)",
        "pattern = hexTexture or noise",
        "emissive = rim + ripple * hitIntensity + pattern * patternStrength",
    ])
    h2(doc, "Data problem")
    bullet(doc, "The shader needs hit position and hit time.")
    bullet(doc, "Data can come from dynamic material parameters, Niagara user parameters, Blueprint, C#, or render target.")
    bullet(doc, "Multiple simultaneous hits require array-like data, render targets, or simplified spawning layers.")
    callout(doc, "Technical VFX question", ["How many hits must the shield support at once, and what is the cheapest data path for that requirement?"], fill=LIGHT_GOLD, color=GOLD)

    chapter(doc, 20, "Recipe: Portal / Black Hole", "Portal and black hole effects are built from radial masks, polar UV, distortion, depth, and strong value hierarchy.")
    h2(doc, "Structure")
    bullet(doc, "Radial mask defines the portal area.")
    bullet(doc, "Polar UV creates circular motion, rings, spiral, or vortex flow.")
    bullet(doc, "Scene distortion pulls or twists the background.")
    bullet(doc, "Rim band provides readable silhouette.")
    bullet(doc, "Interior layers create depth: stars, noise, clouds, rings, particles.")
    h2(doc, "Mental formula")
    codebox(doc, [
        "centered = uv - 0.5",
        "r = length(centered)",
        "angle = atan2(centered.y, centered.x)",
        "polarUV = float2(angle / TWO_PI, r)",
        "swirlUV = polarUV + float2(time * speed + r * twist, 0)",
        "rim = smoothstep(radius - width, radius, r) - smoothstep(radius, radius + width, r)",
    ])
    callout(doc, "Advanced path", ["A custom HLSL version can make polar math cleaner, but the concept is still the same as the node graph."], fill=LIGHT_BLUE)

    chapter(doc, 21, "Production Checklists", "Use checklists to turn taste into repeatable production quality.")
    h2(doc, "Visual checklist")
    checklist(doc, [
        "Does the effect read at gameplay camera distance?",
        "Is the silhouette clear in both bright and dark environments?",
        "Is the timing clear: start, peak, decay?",
        "Is the main value hierarchy visible without relying only on bloom?",
        "Does the effect match the gameplay importance?",
    ])
    h2(doc, "Technical checklist")
    checklist(doc, [
        "Are bounds correct?",
        "Is overdraw acceptable?",
        "Are textures packed and compressed reasonably?",
        "Are expensive features justified: refraction, lighting, shadows, collision?",
        "Do parameters have readable names?",
        "Is there a lower-cost version or scalability path?",
    ])
    h2(doc, "Learning checklist")
    checklist(doc, [
        "Can I explain the shape, mask, motion, color, interaction, and cost?",
        "Can I recreate the effect in another engine?",
        "Can I explain the effect in English in under one minute?",
    ])

    chapter(doc, 22, "English Speaking for Technical VFX", "English speaking is part of technical skill because you need to explain decisions, risks, and feedback.")
    h2(doc, "Review sentences")
    codebox(doc, [
        "The effect is built from three layers: a core emissive mask, a scrolling distortion layer, and a soft outer glow.",
        "The mask is driven by particle age, so the timing can be adjusted without changing the texture.",
        "I baked the expensive noise into a texture because this needs to run on multiple characters at once.",
        "The screen-space distortion looks better, but it may create artifacts near silhouettes.",
        "If we want a softer edge, I can widen the smoothstep range instead of lowering the overall opacity.",
        "The issue seems to come from UV space. The mask is using atlas UVs instead of local frame UVs.",
    ])
    h2(doc, "Feedback vocabulary")
    bullet(doc, "too harsh: edge quá gắt -> soften mask, widen smoothstep.")
    bullet(doc, "muddy: hình đục, thiếu đọc hình -> increase value separation.")
    bullet(doc, "noisy: quá rối -> reduce noise intensity or layer count.")
    bullet(doc, "punchier: cần lực hơn -> timing, core contrast, peak brightness.")
    bullet(doc, "grounded: ăn vào world hơn -> depth fade, contact layer, shadow.")
    bullet(doc, "too expensive: quá nặng -> overdraw, texture fetch, particles, feature cost.")

    chapter(doc, 23, "Six-Week Study Plan", "This plan uses VFXSamplePack as a lab while building engine-neutral shader fluency.")
    h2(doc, "Week 1: Additive and emissive")
    bullet(doc, "Study M_Flare. Write down what controls brightness, opacity, exponent, camera fade.")
    bullet(doc, "Exercise: create 3 flare variants: soft, sharp, colored.")
    h2(doc, "Week 2: Masks and dissolve")
    bullet(doc, "Study noise, smoothstep, pow, threshold.")
    bullet(doc, "Exercise: build dissolve and explain edge band in English.")
    h2(doc, "Week 3: SubUV and flipbook")
    bullet(doc, "Study M_SubUV_Utilities and any 8x8 flipbook material.")
    bullet(doc, "Exercise: explain derived 0-1 UV and debug with NumberGrid.")
    h2(doc, "Week 4: Distortion and depth")
    bullet(doc, "Build heat haze or shockwave with depth fade.")
    h2(doc, "Week 5: Fresnel, spaces, hit data")
    bullet(doc, "Build shield rim and one hit ripple.")
    h2(doc, "Week 6: Advanced mini-project")
    bullet(doc, "Build portal or black hole. Present shape, mask, motion, color, interaction, cost.")

    chapter(doc, 24, "Glossary", "Keep this page open while reading materials. These words are your shader vocabulary.")
    glossary = [
        ("Alpha", "Visibility or opacity signal."),
        ("Atlas", "One texture containing many frames or regions."),
        ("Bloom", "Post-process glow caused by bright values."),
        ("Channel packing", "Storing different data in R/G/B/A channels."),
        ("Depth fade", "Soft fade based on distance to scene depth."),
        ("Emissive", "Self-lit color output, often used for glow."),
        ("Fresnel", "View-angle rim effect based on normal and camera direction."),
        ("Mask", "A grayscale control value."),
        ("Overdraw", "Cost from drawing many transparent pixels over the same screen area."),
        ("Polar UV", "UV based on radius and angle instead of x/y."),
        ("Refraction", "Bending or offsetting the scene behind translucent material."),
        ("SubUV", "Sampling one frame/cell from a texture atlas."),
        ("WPO", "World Position Offset, vertex movement in material."),
    ]
    for name, desc in glossary:
        term(doc, name, desc, "Read material graphs and discuss feedback.")
    callout(doc, "Final note", ["When shader becomes a language, you stop asking 'what node should I copy?' and start asking 'what visual sentence do I want to say?'"], fill=LIGHT_GOLD, color=GOLD)

    doc.core_properties.title = "Realtime Shader VFX Mini-Book"
    doc.core_properties.subject = "Material and shader learning path for Technical VFX"
    doc.core_properties.author = "Codex"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
