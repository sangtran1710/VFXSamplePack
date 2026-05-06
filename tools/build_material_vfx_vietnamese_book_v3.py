from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path.home() / "Downloads" / "Giao_Trinh_Material_VFX_Realtime_V3.docx"

BLACK = RGBColor(18, 18, 18)
DARK = RGBColor(38, 38, 38)
MUTED = RGBColor(100, 100, 100)
LINE = "D8D8D8"
SOFT = "F6F6F6"
NOTE = "F2F2F2"


def set_font(run, size=10.5, bold=False, italic=False, color=BLACK, name="Arial"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def shade(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def border_bottom(paragraph, color=LINE, size="8"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), color)


def p(doc, text="", size=10.7, bold=False, italic=False, color=BLACK, before=0, after=5, align=None):
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(before)
    para.paragraph_format.space_after = Pt(after)
    para.paragraph_format.line_spacing = 1.15
    if align is not None:
        para.alignment = align
    if text:
        run = para.add_run(text)
        set_font(run, size=size, bold=bold, italic=italic, color=color)
    return para


def h_chapter(doc, number, title):
    para = p(doc, f"Chương {number}. {title}", size=18, bold=True, color=BLACK, before=14, after=8)
    border_bottom(para, size="10")


def h1(doc, title):
    para = p(doc, title, size=14.2, bold=True, color=DARK, before=9, after=5)
    border_bottom(para, color="E3E3E3", size="6")


def h2(doc, title):
    p(doc, title, size=11.6, bold=True, color=DARK, before=7, after=3)


def body(doc, text):
    p(doc, text, size=10.75, after=5)


def bullet(doc, text):
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Cm(0.58)
    para.paragraph_format.first_line_indent = Cm(-0.25)
    para.paragraph_format.space_after = Pt(3)
    para.paragraph_format.line_spacing = 1.1
    run = para.add_run("- " + text)
    set_font(run, size=10.25)


def note(doc, title, lines):
    para = p(doc, title, size=10.2, bold=True, color=DARK, before=6, after=2)
    para.paragraph_format.left_indent = Cm(0.25)
    para.paragraph_format.right_indent = Cm(0.2)
    shade(para, NOTE)
    for line in lines:
        para = doc.add_paragraph()
        para.paragraph_format.left_indent = Cm(0.55)
        para.paragraph_format.right_indent = Cm(0.2)
        para.paragraph_format.space_after = Pt(2)
        para.paragraph_format.line_spacing = 1.08
        shade(para, NOTE)
        run = para.add_run("- " + line)
        set_font(run, size=9.8)


def code(doc, lines, title=None):
    if title:
        p(doc, title, size=9.1, italic=True, color=MUTED, after=1)
    for line in lines:
        para = p(doc, line, size=9.2, after=1)
        para.paragraph_format.left_indent = Cm(0.35)
        shade(para, SOFT)
        for run in para.runs:
            set_font(run, size=9.2, name="Consolas", color=BLACK)


def cover(doc):
    p(doc, "", before=2)
    p(doc, "Giáo Trình Material VFX Realtime", size=22, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    p(doc, "Bản V3 - bản đọc chậm, ít màu, tập trung vào cách hiểu", size=12, color=DARK, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    note(doc, "Quyển này viết lại để làm gì", [
        "Ít giống checklist hơn. Mỗi chương bắt đầu từ một tình huống làm nghề, rồi mới đi vào node và công thức.",
        "Body dùng chữ đen. Màu được giảm tối đa để khi in ra giấy vẫn dễ đọc.",
        "Mục tiêu là giúp bạn hình thành phản xạ đọc material, không chỉ nhớ từng node riêng lẻ.",
    ])
    p(doc, "Nguồn tham khảo cách viết: Google Developer Documentation Style Guide, Microsoft writing style, GOV.UK content design, Write the Docs.", size=8.8, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, before=4)


def toc(doc):
    h1(doc, "Mục lục")
    for item in [
        "1. Học material như học một ngôn ngữ visual",
        "2. Material không phải một hình đẹp, nó là một hệ thống nhỏ",
        "3. Bài đọc đầu tiên: nhìn M_Flare như một câu đơn giản",
        "4. UV motion: cách texture bắt đầu sống",
        "5. Mask grammar: nơi effect thật sự được viết",
        "6. Màu, alpha và blend mode: pixel hòa vào thế giới ra sao",
        "7. Particle data và SubUV: khi Niagara bắt đầu nói chuyện với material",
        "8. Depth, Fresnel, WorldPosition và WPO: đưa effect ra khỏi mặt phẳng",
        "9. Sáu recipe production được mổ xẻ chậm",
        "10. Đọc VFXSamplePack mà không copy",
        "11. Tối ưu và feedback: phần làm ra tiền",
        "12. Lộ trình luyện tập có thể làm thật",
    ]:
        bullet(doc, item)
    note(doc, "Cách đọc", [
        "Đọc một chương, rồi mở Unreal làm đúng một bài nhỏ. Đừng đọc hết một lần như đọc tài liệu tham khảo.",
        "Nếu gặp công thức, hãy nối nó ra grayscale trước. Khi grayscale chưa rõ, đừng thêm màu.",
        "Mỗi khi thấy một effect đẹp, hãy tự hỏi: shape, mask, motion, color, interaction và cost nằm ở đâu?",
    ])
    doc.add_page_break()


def chapter_1(doc):
    h_chapter(doc, 1, "Học material như học một ngôn ngữ visual")
    body(doc, "Bạn nói rất đúng khi so sánh shader với tiếng Anh. Nhưng có một điểm cần chỉnh nhẹ: học shader không phải chỉ là học nhiều từ vựng node. Nó giống học cách nói bằng hình ảnh. Node là từ, nhưng effect là câu. Một câu hay không phải vì dùng từ khó; nó hay vì đúng ý, đúng nhịp, đúng ngữ cảnh.")
    body(doc, "Khi một art director nói 'làm mạnh hơn', đó không phải một yêu cầu kỹ thuật rõ ràng. Nó là một câu visual. Bạn phải dịch nó thành hành động: tăng core, giảm outer falloff, làm timing nhanh hơn, đổi scale, tăng contrast mask, hoặc chỉnh camera fade. Đây là lý do Technical VFX Artist được trả tiền cao hơn người chỉ biết kéo node. Họ dịch được cảm giác thành hệ thống.")
    body(doc, "AE cho bạn cảm giác trực tiếp: kéo Glow, Twirl, Noise, Turbulent Displace, rồi keyframe. Engine thì khác. Bạn không chỉ vẽ kết quả. Bạn tạo một luật chạy mỗi frame, trong mọi góc camera, mọi lighting, mọi spawn count, và dưới giới hạn performance. Quyển này sẽ không cố làm bạn thuộc hết node. Nó cố giúp bạn đặt đúng câu hỏi trước khi mở node.")
    note(doc, "Câu hỏi gốc của mọi material VFX", [
        "Pixel này lấy dữ liệu từ đâu: texture, UV, particle, world, camera hay scene depth?",
        "Dữ liệu đó đang tạo shape, mask, motion, color, interaction hay cost?",
        "Nếu feedback đến, tôi chỉnh layer nào trước?",
    ])
    code(doc, [
        "Visual feedback -> shader translation",
        "'Mạnh hơn'       -> core? contrast? scale? timing? camera?",
        "'Giả quá'        -> motion quá đều? mask quá sạch? thiếu interaction?",
        "'Nặng quá'       -> overdraw? particle count? texture fetch? feature đắt?",
    ], "Một cách dịch feedback")
    body(doc, "Bạn có thể xem material như một đoạn văn. Có chủ ngữ là source data. Có động từ là node math. Có ngữ pháp là mask và UV. Có giọng điệu là màu và blend mode. Có ngữ cảnh là scene, camera và gameplay. Nếu chỉ nhìn từng node riêng, bạn sẽ thấy rối. Nếu nhìn như một câu, graph bắt đầu có nghĩa.")


def chapter_2(doc):
    h_chapter(doc, 2, "Material là một hệ thống nhỏ")
    body(doc, "Một material VFX không sống một mình. Nó nằm giữa texture, mesh, Niagara, renderer, camera, exposure và performance budget. Đây là điều làm nó khác với một layer AE đã render xong. Khi effect hỏng, lỗi có thể không nằm ở graph. Có thể texture import sai sRGB. Có thể Niagara không gửi Particle Color. Có thể sprite renderer SubUV sai grid. Có thể blend mode đúng trong preview nhưng sai trong level thật.")
    body(doc, "Vì vậy, trước khi sửa node, hãy đọc pipeline. Texture là dữ liệu. Shader là chương trình xử lý. Material là asset bạn author. Material Instance là nơi bạn chỉnh parameter mà không phá graph. Renderer là nơi material được vẽ lên sprite, mesh, ribbon hoặc decal. Niagara là nơi gửi lifetime, size, color, alpha, frame và custom data.")
    code(doc, [
        "Texture / Mesh / Particle Data",
        "        -> Material graph",
        "        -> Compiled shader",
        "        -> Renderer draws sprite / mesh / ribbon",
        "        -> Scene: depth, exposure, camera, overdraw",
    ], "Pipeline data-to-pixel")
    body(doc, "Một cách học tốt là luôn tách trách nhiệm. Texture chịu trách nhiệm chứa mask hay frame. Material chịu trách nhiệm reshape, tint, fade, interact. Niagara chịu trách nhiệm timing, spawn, size, color over life. Renderer chịu trách nhiệm cách material gắn lên geometry. Nếu bạn nhồi mọi thứ vào material, graph sẽ nặng và khó chỉnh. Nếu bạn bake mọi thứ vào texture, effect sẽ thiếu interaction.")
    note(doc, "Bài tập 20 phút", [
        "Tạo một sprite material chỉ có Texture * ParticleColor.",
        "Trong Niagara, đổi Color over Life và Alpha over Life.",
        "Nếu material không đổi màu hoặc không fade, hãy tìm lỗi trong hợp đồng giữa Niagara và material.",
    ])
    body(doc, "Đây là bài tập nhỏ nhưng rất quan trọng. Nhiều material production fail chỉ vì thiếu một dây nối đơn giản từ Particle Color alpha vào opacity hoặc emissive fade. Người mới thường đi tìm node phức tạp; người có kinh nghiệm kiểm hợp đồng dữ liệu trước.")


def chapter_3(doc):
    h_chapter(doc, 3, "Bài đọc đầu tiên: M_Flare")
    body(doc, "Nếu phải chọn một material để bắt đầu, flare là bài rất tốt. Nó đơn giản, nhưng chứa gần đủ tư duy VFX: texture shape, alpha power, emissive intensity, particle color, additive fade và camera fade. Một flare tệ là một đốm trắng lớn. Một flare tốt có core, falloff, timing và fade sạch.")
    body(doc, "Hãy tưởng tượng bạn mở M_Flare. Đừng đọc từng node ngay. Trước tiên hỏi: texture nào tạo shape? Alpha đi đâu? Particle Color ảnh hưởng màu hay alpha? Emissive có nhân với intensity nào? Nếu blend mode là Additive, fade có thật sự đi vào Emissive không? Có camera fade để flare không đập vào màn hình không?")
    code(doc, [
        "Texture RGB/A",
        "  -> alpha shaping: pow(alpha, AlphaPower)",
        "  -> color: TextureRGB * ParticleColor.rgb * Intensity",
        "  -> fade: ParticleColor.a * CameraFade",
        "  -> output: Emissive = color * alpha * fade",
    ], "M_Flare đọc như một câu")
    body(doc, "Câu 'TextureRGB nhân ParticleColor rồi nhân Intensity' nghe rất đơn giản. Nhưng nó là nền của rất nhiều sprite VFX. Nếu ParticleColor.a không scale Emissive trong Additive, Niagara alpha curve có thể không fade như bạn tưởng. Nếu AlphaPower quá cao, flare có core rất gắt nhưng mất outer glow. Nếu Intensity quá cao, mọi thứ thành trắng bệt.")
    note(doc, "Cách debug", [
        "Nối alpha texture ra grayscale. Đừng nhìn màu trước.",
        "Tắt bloom hoặc giảm intensity để xem shape thật.",
        "Trong Niagara, kéo alpha từ 1 về 0. Nếu flare không biến mất, fade path sai.",
        "Test trên nền sáng. Additive thường đẹp trên nền tối nhưng mất đọc trên nền sáng.",
    ])
    body(doc, "Bài học của M_Flare không phải là 'flare dùng những node nào'. Bài học là một material nhỏ vẫn phải có hierarchy. Core, outer glow, tint, lifetime fade và camera fade đều có vai trò. Khi bạn hiểu được flare, bạn đã hiểu một câu đơn giản của ngôn ngữ shader.")


def chapter_4(doc):
    h_chapter(doc, 4, "UV motion: texture bắt đầu sống")
    body(doc, "UV là câu hỏi bạn gửi cho texture: pixel này nên đọc chỗ nào trên ảnh? Khi UV đứng yên, texture đứng yên. Khi UV cộng thêm Time * Speed, texture bắt đầu chảy. Đó là Panner. Nó là một trong những thủ thuật rẻ nhất để tạo chuyển động.")
    body(doc, "Nhưng panner đơn rất dễ bị lộ. Nếu một noise texture chạy đều từ dưới lên, mắt sẽ thấy nó như một tấm giấy trượt. Motion tốt thường có ít nhất hai lớp: một lớp lớn chậm tạo hướng chính, một lớp nhỏ nhanh tạo chi tiết. Bạn có thể nhân hai lớp, lerp chúng, hoặc dùng lớp nhỏ để distort lớp lớn.")
    code(doc, [
        "uvLarge = UV * 2 + Time * float2(0, 0.2)",
        "uvDetail = UV * 8 - Time * float2(0.1, 0.7)",
        "mask = NoiseLarge(uvLarge) * NoiseDetail(uvDetail)",
    ], "Hai lớp panner")
    body(doc, "Trước khi dùng noise đẹp, hãy dùng NumberGrid hoặc checker. Checker cho bạn biết UV đang đi hướng nào, scale ra sao, có xoay đúng pivot không. Đây là thói quen nhỏ nhưng cứu rất nhiều thời gian. Nếu dùng noise ngay, bạn có thể không nhận ra UV đang sai hướng.")
    note(doc, "Bài tập", [
        "Tạo material sample NumberGrid.",
        "Thêm Tiling, Panner, Rotator.",
        "Đổi pivot Rotator từ 0.5,0.5 sang 0,0 để thấy lỗi lệch tâm.",
        "Sau đó mới thay NumberGrid bằng noise.",
    ])
    body(doc, "Portal, black hole, magic ring cần motion vòng. Khi đó panner ngang dọc không đủ. Bạn cần radial hoặc polar UV. Polar UV biến một texture phẳng thành không gian có trục angle và radius. Noise có thể chạy quanh tâm thay vì trượt ngang. Đây là lúc Atan2 và Length bắt đầu có ý nghĩa visual.")
    code(doc, [
        "p = UV - 0.5",
        "r = length(p)",
        "a = atan2(p.y, p.x) / TWO_PI",
        "polarUV = float2(a + Time * Spin, r)",
    ], "Polar UV cơ bản")


def chapter_5(doc):
    h_chapter(doc, 5, "Mask grammar: nơi effect được viết")
    body(doc, "Nếu chỉ được học một thứ trong material VFX, hãy học mask. Màu làm effect đẹp hơn, nhưng mask quyết định effect có đọc được hay không. Mask nói pixel nào hiện, pixel nào mất, vùng nào sáng, vùng nào mềm, vùng nào bị ăn mòn. Một mask tốt ở grayscale thường sẽ thành effect tốt khi thêm màu. Một mask xấu không nên được che bằng bloom.")
    body(doc, "Dissolve là ví dụ rõ nhất. Dissolve không phải fade alpha đều. Nó là một threshold chạy qua một field. Field có thể là noise, height map, vertex color, world aligned texture hoặc texture vẽ tay. Amount tăng lên, threshold di chuyển, một phần bề mặt biến mất. Edge band là vùng sát threshold, nơi bạn thêm lửa, băng, magic hoặc energy.")
    code(doc, [
        "visible = smoothstep(Amount, Amount + Softness, Mask)",
        "edge = smoothstep(Amount - EdgeWidth, Amount, Mask)",
        "     - smoothstep(Amount, Amount + Softness, Mask)",
        "emissive = Base * visible + EdgeColor * edge * EdgeIntensity",
    ], "Dissolve có visible và edge riêng")
    body(doc, "Điểm quan trọng là visible và edge phải tách nhau. Nếu art director nói 'viền mỏng hơn nhưng sáng hơn', bạn cần EdgeWidth và EdgeIntensity riêng. Nếu mọi thứ bị trộn vào một alpha duy nhất, feedback nhỏ cũng trở thành sửa graph.")
    note(doc, "Cách học mask", [
        "Luôn nối mask ra Emissive grayscale trước.",
        "Tạo ba giá trị Amount: thấp, giữa, cao.",
        "Tạo ba giá trị Softness: sắc, vừa, mềm.",
        "Chỉ thêm màu khi grayscale đã đọc rõ.",
    ])
    body(doc, "Channel packing cũng thuộc mask grammar. Một texture R/G/B/A có thể chứa bốn loại dữ liệu: shape, breakup, edge, alpha. Đó không phải ảnh màu để ngắm. Nó là bộ nhớ dữ liệu. Khi dùng packed texture, hãy debug từng kênh. Đừng đoán alpha luôn là opacity.")


def chapter_6(doc):
    h_chapter(doc, 6, "Màu, alpha và blend mode")
    body(doc, "Emissive không có nghĩa là kéo sáng thật mạnh. Emissive tốt có hierarchy: core sáng nhất, body thấp hơn, edge hoặc detail vừa đủ. Nếu mọi layer cùng sáng, mắt không biết nhìn đâu. Nếu intensity quá cao, bloom biến mọi thứ thành một đám trắng.")
    body(doc, "Blend mode là luật pixel hòa vào background. Additive cộng sáng, hợp flare, spark, energy. Translucent pha alpha, hợp smoke, mist, portal interior, fire mềm. Masked cắt cứng, hợp surface dissolve hoặc cutout rõ. Một texture giống nhau nhưng đổi blend mode sẽ thành visual khác hẳn.")
    code(doc, [
        "Additive:    scene + emissive",
        "Translucent: lerp(scene, material, opacity)",
        "Masked:      pixel sống nếu opacityMask > clipValue",
    ], "Ba luật hòa pixel")
    body(doc, "Lỗi rất hay gặp: dùng Additive nhưng chỉ nối alpha vào Opacity, rồi thắc mắc vì sao particle không fade. Với Additive, fade thường phải scale Emissive. Với Translucent, Opacity là đường chính. Với Masked, Opacity Mask và clip threshold quyết định pixel sống chết.")
    note(doc, "Bài tập", [
        "Dùng cùng một radial texture tạo ba material: Additive, Translucent, Masked.",
        "Đặt cả ba trên nền sáng và nền tối.",
        "Ghi lại: bản nào mất đọc, bản nào cắt cứng, bản nào overdraw nhiều.",
    ])
    body(doc, "Khi màu bị bẩn, thường là do bạn nhân nhiều texture RGB với nhau hoặc để texture color đánh nhau với color ramp. Cách sạch hơn là dùng grayscale mask để điều khiển color parameter. Texture quyết định vùng, parameter quyết định hue.")


def chapter_7(doc):
    h_chapter(doc, 7, "Particle data và SubUV")
    body(doc, "Material bắt đầu thật sự realtime khi nó nhận data từ Niagara. Particle Color là dây nối đầu tiên. RGB tint màu. Alpha fade opacity hoặc emissive. Nếu Color over Life trong Niagara không có tác dụng, đừng tìm node lạ trước. Hãy kiểm material có dùng Particle Color không.")
    code(doc, [
        "emissive = TextureRGB * ParticleColor.rgb * Intensity",
        "opacity  = TextureA * ParticleColor.a",
        "additiveEmissive = emissive * ParticleColor.a",
    ], "Particle Color contract")
    body(doc, "Dynamic Parameter là bốn kênh float từ Niagara vào material. Nó quý vì ít. Hãy đặt tên ổn định: R có thể là EmissivePower, G là FadeDistance, B là CameraFadeLength, A là AlphaPower. Nếu mỗi material dùng kênh khác nhau, project sẽ khó maintain.")
    body(doc, "SubUV là một bài khác. Với flipbook, texture atlas chứa nhiều frame. Niagara sprite renderer quyết định frame nào đang được vẽ. Material dùng ParticleSubUV để đọc frame đó. Nếu grid sai, material đúng vẫn ra sai. Vì vậy NumberGrid là bạn tốt nhất của bạn khi học SubUV.")
    code(doc, [
        "ParticleSubUV.rgb -> texture frame color",
        "ParticleSubUV.a   -> texture frame alpha",
        "ParticleColor     -> tint/fade từ Niagara",
        "DepthFade         -> soft intersection",
    ], "Smoke/fire flipbook baseline")
    body(doc, "Derived 0-1 UV là UV local trong frame hiện tại. Nó hữu ích khi bạn muốn tạo fade, radial mask hoặc distortion bên trong từng frame mà không sample sang frame khác. Atlas UV là toàn sheet; derived UV là căn phòng nhỏ bên trong frame.")
    note(doc, "Bài tập", [
        "Dùng NumberGrid 3x3 làm flipbook.",
        "Bật SubUV trong Niagara và kiểm thứ tự frame.",
        "Tạo fade từ trái sang phải trong từng frame bằng local 0-1 UV.",
    ])


def chapter_8(doc):
    h_chapter(doc, 8, "Depth, Fresnel, WorldPosition và WPO")
    body(doc, "Một sprite smoke cắt qua mặt đất sẽ lộ cạnh cứng. DepthFade làm mềm giao tuyến đó bằng cách so pixel particle với scene depth. Đây là một trong những node khiến particle bớt giống card phẳng. Nhưng DepthFade không sửa mask xấu; nó chỉ sửa giao tuyến.")
    body(doc, "Fresnel dùng góc giữa normal và view direction để tạo rim. Shield, hologram, forcefield, rim glow đều thường bắt đầu từ Fresnel. Nhưng Fresnel đơn dễ sạch quá. Hãy nhân nó với noise, hit ripple hoặc opacity control để có vật liệu sống hơn.")
    code(doc, [
        "facing = dot(normalize(NormalWS), normalize(ViewDirWS))",
        "rim = pow(1 - saturate(facing), RimPower)",
    ], "Fresnel thủ công")
    body(doc, "WorldPosition đưa material ra khỏi UV. Nó cho biết pixel đang ở đâu trong thế giới. Bạn có thể làm height dissolve bằng WorldPosition.z, hoặc hit ripple bằng distance(WorldPosition, HitPosition). Đây là điểm shader bắt đầu nói chuyện với gameplay.")
    code(doc, [
        "d = distance(WorldPos, HitPos)",
        "ring = smoothstep(Radius - Width, Radius, d)",
        "     - smoothstep(Radius, Radius + Width, d)",
        "WPO = VertexNormalWS * ring * ImpactAmount",
    ], "Hit ripple + WPO")
    note(doc, "Lỗi cần nhớ", [
        "HitPosition và WorldPosition phải cùng space.",
        "WPO đẩy vertex, không đẩy pixel. Mesh ít vertex sẽ biến dạng thô.",
        "WPO lớn cần bounds đủ lớn, nếu không mesh có thể bị culling.",
    ])


def recipe(doc, title, paragraphs, graph, debug, exercise):
    h1(doc, title)
    for text in paragraphs:
        body(doc, text)
    code(doc, graph, "Graph đọc theo lane")
    note(doc, "Debug", debug)
    note(doc, "Bài tập", exercise)


def chapter_9(doc):
    h_chapter(doc, 9, "Sáu recipe production được mổ xẻ chậm")
    recipe(doc, "Flare/glow", [
        "Flare là câu đơn giản nhất: texture tạo shape, alpha tạo falloff, Particle Color tạo tint/fade, intensity tạo lực sáng. Nhưng đơn giản không có nghĩa dễ. Flare rất dễ cháy trắng, mất đọc trên nền sáng, hoặc không fade đúng với Niagara.",
        "Khi làm flare, đừng bắt đầu bằng bloom. Hãy tắt bloom, nhìn alpha, chỉnh core trước. Sau đó mới tăng emissive vừa đủ.",
    ], [
        "UV -> sample flare texture",
        "Alpha -> pow(alpha, AlphaPower)",
        "RGB -> * ParticleColor.rgb * Intensity",
        "Fade -> * ParticleColor.a * CameraFade",
        "Output -> Emissive",
    ], [
        "Nối alpha ra grayscale.",
        "Kéo Niagara alpha về 0 để kiểm fade.",
        "Test trên background sáng.",
    ], [
        "Tạo ba instance: spark nhỏ, impact flash, magic bloom.",
    ])
    recipe(doc, "Dissolve/reveal", [
        "Dissolve production là threshold chạy qua một field, không phải opacity fade đều. Field có thể là noise, height, vertex color hoặc world gradient. Edge band là phần ăn tiền vì nó cho cảm giác lửa, băng, magic hoặc corrosion.",
        "Nếu muốn feedback dễ sửa, Amount, Softness, EdgeWidth và EdgeIntensity phải độc lập.",
    ], [
        "Source mask -> SmoothStep Amount",
        "Source mask -> Edge band around Amount",
        "Visible -> Opacity/OpacityMask",
        "Edge -> Emissive accent",
    ], [
        "Tách source, visible và edge thành ba grayscale.",
        "Animate Amount chậm.",
        "Test trên mesh UV tốt và UV xấu.",
    ], [
        "Làm một bản cháy lửa và một bản đóng băng từ cùng master.",
    ])
    recipe(doc, "Fire card", [
        "Fire card procedural dạy bạn lửa được cấu trúc như thế nào: gốc dày, ngọn mỏng, motion đi lên nhưng không đều, core nóng và edge tối hơn. Nếu mọi thứ chỉ là một panner, lửa sẽ như giấy trượt.",
        "Một bản tốt có vertical gradient giữ shape, hai noise panner tạo breakup, color ramp tạo core và edge.",
    ], [
        "UV.y -> vertical gradient",
        "Noise large + Noise detail -> breakup",
        "Gradient * breakup -> alpha",
        "Alpha/core -> color ramp -> emissive",
    ], [
        "Nối vertical gradient riêng.",
        "Tắt từng noise layer.",
        "Test card trên nền sáng.",
    ], [
        "Tạo torch, magic flame và wall flame bằng cùng graph.",
    ])
    recipe(doc, "Smoke flipbook", [
        "Smoke thường nên để flipbook gánh motion phức tạp. Material chỉ nên shape alpha, tint, depth fade và giữ contract với Niagara. Smoke nặng chủ yếu vì overdraw, không chỉ vì shader graph.",
    ], [
        "ParticleSubUV RGB/A",
        "* ParticleColor RGB/A",
        "Alpha -> pow(alpha, AlphaPower)",
        "Opacity -> DepthFade",
    ], [
        "Dùng NumberGrid trước atlas thật.",
        "Xem shader complexity với spawn count thật.",
        "Tighten alpha để giảm vùng trong suốt vẫn vẽ.",
    ], [
        "Tạo smoke mỏng, smoke dày và ambient smoke từ cùng material.",
    ])
    recipe(doc, "Portal ring", [
        "Portal phải đọc là vòng trước khi đọc là noise. Ring band đến từ radial distance. Noise flow có thể chạy trong polar UV để tạo cảm giác xoáy quanh tâm. Interior nên yếu hơn edge để không biến portal thành đĩa sáng.",
    ], [
        "UV -> radial distance",
        "Radius/Width -> ring band",
        "Polar UV -> moving noise",
        "Ring * noise -> emissive edge",
    ], [
        "Tắt interior, chỉ xem ring.",
        "Kiểm polar seam.",
        "Giữ Radius/Width chỉnh được độc lập với noise.",
    ], [
        "Tạo clean sci-fi portal và noisy magic portal bằng cùng graph.",
    ])
    recipe(doc, "Forcefield", [
        "Forcefield là nơi material trở thành interactive. Fresnel tạo shell, noise tạo surface pattern, WorldPosition distance tạo hit ripple, WPO nhỏ tạo phản hồi vật lý giả. Nếu opacity body quá cao, shield thành quả bóng đặc.",
    ], [
        "Fresnel -> rim opacity",
        "Noise -> pattern breakup",
        "distance(WorldPos, HitPos) -> hit ring",
        "VertexNormalWS * hit ring -> WPO",
    ], [
        "Debug Fresnel, noise và hit ring riêng.",
        "Kiểm HitPosition có cùng world space không.",
        "Giữ WPO nhỏ và kiểm bounds.",
    ], [
        "Tạo shield có hit point parameter, sau đó nối với Blueprint/Niagara.",
    ])


def chapter_10(doc):
    h_chapter(doc, 10, "Đọc VFXSamplePack mà không copy")
    body(doc, "Khi mở material có sẵn, mục tiêu không phải copy node. Mục tiêu là hiểu vì sao node ở đó. Hãy viết lại graph bằng lời. Nếu không viết được, bạn chưa thật sự đọc xong.")
    for title, text in [
        ("M_Flare", "Đọc như bài về additive sprite: texture shape, alpha power, Particle Color, emissive intensity, camera fade. Bài rebuild: tạo bản tối giản chỉ giữ năm control chính."),
        ("M_Mesh_Add", "Đọc như bài về mesh additive: mesh cung cấp shape/normal/world position; material tạo energy trên bề mặt. Tìm WPO, gradient, noise mask và lý do dùng mesh thay vì sprite."),
        ("M_FlamesCard", "Đọc như bài về card flame: UV speed, vertical fade, edge fade, distortion, emissive ramp. Tách motion khỏi color, tách silhouette khỏi detail."),
        ("M_SubUV_Utilities", "Đọc như bài về atlas logic: NumberGrid, SubUV, frame UV, derived 0-1 UV. Đây là material để học cách nhìn frame bằng mắt."),
    ]:
        h1(doc, title)
        body(doc, text)
        note(doc, "Cách đọc", [
            "Tìm output trước: Emissive, Opacity, WPO.",
            "Lần ngược về source data.",
            "Tách lane: UV, Mask, Color, Fade, Interaction, Cost.",
            "Rebuild bản nhỏ từ trí nhớ.",
        ])
    body(doc, "Một bài đọc material tốt kết thúc bằng một bản rebuild nhỏ. Bản rebuild không cần giống 100%. Nó phải chứng minh bạn hiểu câu chính. Khi bạn tự rebuild được, graph gốc không còn là bí mật.")


def chapter_11(doc):
    h_chapter(doc, 11, "Tối ưu và feedback")
    body(doc, "Optimization không bắt đầu bằng việc xóa vài node. Với VFX, overdraw và particle count thường đáng sợ hơn vài phép Multiply. Một smoke material đơn giản vẫn nặng nếu 200 sprite lớn phủ màn hình. Một graph có thêm vài math có thể ổn nếu effect nhỏ và hiếm.")
    note(doc, "Thứ tự kiểm cost", [
        "Screen area và overdraw.",
        "Particle count, lifetime, spawn rate.",
        "Texture size, texture fetch, alpha vùng trống.",
        "Feature đắt: refraction, lit translucency, scene depth.",
        "Shader instructions.",
        "Bounds và sorting.",
    ])
    body(doc, "Feedback cũng cần dịch đúng. 'Mạnh hơn' không mặc định là tăng intensity. 'Giả quá' không mặc định là thêm noise. 'Nặng quá' không mặc định là xóa node. Hãy hỏi feedback đó thuộc layer nào.")
    for title, actions in [
        ("Mạnh hơn", ["Tăng core trước, không tăng toàn bộ glow.", "Tăng contrast mask nếu effect bị mờ.", "Tăng scale hoặc timing nếu effect đọc quá nhỏ/quá nhanh."]),
        ("Giả quá", ["Giảm motion quá đều.", "Thêm variation theo particle.", "Thêm interaction như depth fade hoặc hit response."]),
        ("Không đọc được", ["Quay về silhouette grayscale.", "Giảm detail nhỏ đang che shape lớn.", "Test ở gameplay distance."]),
        ("Nặng quá", ["Xem overdraw trước.", "Tighten alpha hoặc giảm card size.", "Giảm spawn count/lifetime nếu particle system là nguyên nhân."]),
    ]:
        h1(doc, title)
        for action in actions:
            bullet(doc, action)
    body(doc, "Điều làm bạn thành Technical VFX Artist mạnh là khả năng nói chuyện với cả hai phía: artist và engine. Bạn nghe được feedback cảm giác, rồi trả lời bằng graph, parameter và performance trade-off.")


def chapter_12(doc):
    h_chapter(doc, 12, "Lộ trình luyện tập")
    body(doc, "Đừng học shader bằng cách hỏi AI mỗi khi kẹt rồi dán node. Hãy dùng AI như gia sư: hỏi vì sao, yêu cầu giải thích từng layer, rồi tự rebuild. Bạn cần muscle memory. Dưới đây là lộ trình thực tế.")
    for week, text in [
        ("Tuần 1: Mask", "Radial, linear, noise threshold, edge band. Chỉ grayscale."),
        ("Tuần 2: UV motion", "Panner, two-layer noise, rotator, polar UV. Dùng NumberGrid để debug."),
        ("Tuần 3: Particle contract", "Particle Color, Dynamic Parameter, SubUV, CameraDepthFade, DepthFade."),
        ("Tuần 4: Interaction", "Fresnel, WorldPosition, hit ripple, WPO, bounds."),
        ("Tuần 5: Recipe", "Flare, dissolve, fire, smoke, portal, shield, beam."),
        ("Tuần 6: Rebuild", "Chọn một material VFXSamplePack, đọc lane, rebuild bản nhỏ, so với bản gốc."),
    ]:
        h1(doc, week)
        body(doc, text)
    h1(doc, "Nhật ký 30 ngày")
    days = [
        "Ngày 1: radial mask.",
        "Ngày 2: ring band.",
        "Ngày 3: alpha power.",
        "Ngày 4: OneMinus reveal/destroy.",
        "Ngày 5: noise dissolve.",
        "Ngày 6: dissolve edge.",
        "Ngày 7: rebuild dissolve.",
        "Ngày 8: panner một lớp.",
        "Ngày 9: hai panner.",
        "Ngày 10: rotator pivot.",
        "Ngày 11: polar UV.",
        "Ngày 12: portal ring.",
        "Ngày 13: Particle Color fade.",
        "Ngày 14: Dynamic Parameter.",
        "Ngày 15: CameraDepthFade.",
        "Ngày 16: DepthFade.",
        "Ngày 17: NumberGrid SubUV.",
        "Ngày 18: smoke flipbook.",
        "Ngày 19: fire card.",
        "Ngày 20: Fresnel shield.",
        "Ngày 21: hit ripple.",
        "Ngày 22: WPO pulse.",
        "Ngày 23: height dissolve.",
        "Ngày 24: channel packing.",
        "Ngày 25: texture import check.",
        "Ngày 26: optimize overdraw.",
        "Ngày 27: optimize texture samples.",
        "Ngày 28: read M_Flare.",
        "Ngày 29: read M_FlamesCard.",
        "Ngày 30: rebuild một effect cũ của bạn.",
    ]
    for d in days:
        bullet(doc, d)
    note(doc, "Kết luận", [
        "Node là từ. Mask, UV và color là ngữ pháp. Recipe là câu mẫu. Feedback production là bài nói thật.",
        "Bạn master khi có thể nhìn visual phức tạp và tự hỏi đúng: dữ liệu nào, mask nào, motion nào, output nào, cost nào.",
    ])


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.35)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10.7)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Giáo Trình Material VFX Realtime V3")
    set_font(run, size=7.8, color=MUTED)

    cover(doc)
    toc(doc)
    chapter_1(doc)
    chapter_2(doc)
    chapter_3(doc)
    chapter_4(doc)
    chapter_5(doc)
    chapter_6(doc)
    chapter_7(doc)
    chapter_8(doc)
    chapter_9(doc)
    chapter_10(doc)
    chapter_11(doc)
    chapter_12(doc)

    doc.core_properties.title = "Giáo Trình Material VFX Realtime V3"
    doc.core_properties.subject = "Vietnamese shader VFX learning guide"
    doc.core_properties.author = "Codex"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
